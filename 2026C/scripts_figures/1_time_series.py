"""读取附件 1，绘制问题一的电价及负载、光伏预测功率日内折线图。

时间沿用附件标签：00:10 至 24:00，每 10 分钟一个点；0:00+1 转为 24:00。
输出目录：2026C/outputs/figures，包含 300 dpi PNG 和 SVG 矢量图。
"""

from datetime import time
from math import isfinite
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import MultipleLocator
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATA_FILE = PROJECT_DIR / "CUMCM 2026 C题" / "附件" / "附件1.xlsx"
FIGURE_DIR = PROJECT_DIR / "outputs" / "figures"
COLUMNS = ("时间", "电价", "小区负载", "光伏发电预测功率")


def time_to_minutes(value: time | str) -> int:
    """兼容附件中混合存储的 Excel 时间和字符串时间。"""
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"时间应精确到分钟：{value!r}")
        return value.hour * 60 + value.minute
    if isinstance(value, str):
        text = value.strip()
        if text == "0:00+1":
            return 24 * 60
        hour, minute = map(int, text.split(":"))
        if 0 <= hour < 24 and 0 <= minute < 60:
            return hour * 60 + minute
    raise ValueError(f"无法识别附件中的时间：{value!r}")


def load_daily_data() -> tuple[list[float], list[list[float]]]:
    """读取数据并核对表头、144 个时间点、间隔及数值有效性。"""
    workbook = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows or rows[0] != COLUMNS:
        raise ValueError(f"附件 1 的表头应为：{COLUMNS}")

    minutes = [time_to_minutes(row[0]) for row in rows[1:]]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError("时间应依次为 00:10 至 24:00，共 144 个点，间隔 10 分钟。")

    series = []
    for column, name in enumerate(COLUMNS[1:], start=1):
        values = []
        for excel_row, row in enumerate(rows[1:], start=2):
            value = row[column]
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not isfinite(value)
                or value < 0
            ):
                raise ValueError(f"第 {excel_row} 行的{name}应为有限非负数：{value!r}")
            values.append(float(value))
        series.append(values)

    return [minute / 60 for minute in minutes], series


def configure_style() -> None:
    """自动选择已安装的中文字体，避免中文显示为方框。"""
    candidates = (
        "Noto Sans CJK SC", "Source Han Sans SC", "Microsoft YaHei",
        "SimHei", "PingFang SC", "Heiti SC", "Arial Unicode MS",
    )
    installed = {font.name for font in font_manager.fontManager.ttflist}
    chinese_font = next((name for name in candidates if name in installed), None)
    if chinese_font is None:
        raise RuntimeError("未找到中文字体，请安装 Noto Sans CJK SC、微软雅黑或黑体后重试。")

    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": [chinese_font, "DejaVu Sans"],
        "axes.unicode_minus": False,
        "axes.titlesize": 16,
        "axes.labelsize": 12,
        "axes.labelpad": 10,
        "xtick.labelsize": 10,
        "ytick.labelsize": 10,
        "legend.fontsize": 11,
        "lines.linewidth": 2.2,
        "savefig.dpi": 300,
        "svg.fonttype": "path",  # SVG 中保留字形，跨设备查看无需另装字体。
    })


def create_axes(title: str, ylabel: str):
    """统一两张图的尺寸、时间轴和版式。"""
    fig, ax = plt.subplots(figsize=(11, 5.3))
    fig.subplots_adjust(left=0.095, right=0.96, bottom=0.20, top=0.87)
    ax.set_title(title, loc="left", pad=16, fontweight="semibold")
    ax.set_xlabel("时间")
    ax.set_ylabel(ylabel)
    ax.set_xlim(0, 24)
    ticks = list(range(0, 25, 2))
    ax.set_xticks(ticks, [f"{hour:02d}:00" for hour in ticks])
    ax.xaxis.set_minor_locator(MultipleLocator(1))
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="#DEE4EB", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#9AA6B2")
    ax.tick_params(axis="both", which="both", color="#9AA6B2")
    fig.text(
        0.095, 0.04,
        "数据来源：附件1.xlsx；原始时间点 00:10-24:00，间隔 10 分钟。",
        fontsize=9, color="#5D6975",
    )
    return fig, ax


def save_figure(fig, name: str) -> None:
    for extension in ("png", "svg"):
        output_file = FIGURE_DIR / f"{name}.{extension}"
        fig.savefig(output_file, facecolor="white")
        print(f"已保存：{output_file}")
    plt.close(fig)


def main() -> None:
    hours, (prices, load, photovoltaic) = load_daily_data()
    configure_style()
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)

    fig, ax = create_axes("问题一：电价随时间的变化", "电价（元/kWh）")
    ax.plot(hours, prices, color="#2667A6")
    ax.set_ylim(0, max(prices) * 1.12)
    ax.yaxis.set_major_locator(MultipleLocator(0.2))
    save_figure(fig, "1_电价时间折线图")

    fig, ax = create_axes(
        "问题一：小区负载与光伏发电预测功率", "功率（kW）"
    )
    ax.plot(hours, load, color="#2667A6", label="小区负载")
    ax.plot(
        hours, photovoltaic, color="#D97924", linestyle="--",
        label="光伏发电预测功率",
    )
    ax.set_ylim(0, max(max(load), max(photovoltaic)) * 1.15)
    ax.yaxis.set_major_locator(MultipleLocator(2000))
    ax.legend(loc="upper left", frameon=False, ncol=2)
    save_figure(fig, "1_小区负载与光伏发电预测功率时间折线图")

    print(f"绘图完成：{len(hours)} 个原始时间点，2 张图，各保存 PNG 和 SVG。")


if __name__ == "__main__":
    main()
