"""求解问题一的确定性日内储能经济调度模型。

模型依据 docs/Q1.md：将一天离散为 144 个 10 分钟时段，使用混合整数
线性规划决定外网购电量、储能充放电量、弃光量和储电量。所有计算保留完整
精度，通过约束校验后再将结果写入 outputs/tables。
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import time
from math import isfinite
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATA_FILE = PROJECT_DIR / "CUMCM 2026 C题" / "附件" / "附件1.xlsx"
OUTPUT_DIR = PROJECT_DIR / "outputs" / "tables"

N_PERIODS = 144
DELTA_T = 1 / 6
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
P_CHARGE_MAX = 5000.0
P_DISCHARGE_MAX = 5000.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
CHARGE_LIMIT = P_CHARGE_MAX * DELTA_T
DISCHARGE_LIMIT = P_DISCHARGE_MAX * DELTA_T
INPUT_COLUMNS = ("时间", "电价", "小区负载", "光伏发电预测功率")

BALANCE_TOLERANCE = 1e-5
BOUND_TOLERANCE = 1e-5
INTEGER_TOLERANCE = 1e-6
DISPLAY_DECIMALS = 6


@dataclass(frozen=True)
class DailyData:
    """附件 1 中完成时间对齐后的单日输入。"""

    interval_labels: list[str]
    price: np.ndarray
    load_power: np.ndarray
    photovoltaic_power: np.ndarray
    load_energy: np.ndarray
    photovoltaic_energy: np.ndarray


@dataclass(frozen=True)
class DispatchResult:
    """混合整数规划的最优调度结果。"""

    grid_purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    curtailment: np.ndarray
    mode: np.ndarray
    stored_energy: np.ndarray
    objective_value: float
    solver_message: str


@dataclass(frozen=True)
class ModelChecks:
    """用于判断求解结果是否可用于输出的校验指标。"""

    max_balance_residual: float
    max_state_residual: float
    max_simultaneous_energy: float
    max_mode_integrality_error: float
    min_stored_energy: float
    max_stored_energy: float
    terminal_energy_error: float
    total_energy_residual: float
    no_storage_cost: float
    optimized_cost: float


def time_to_minutes(value: time | str) -> int:
    """将附件中的 Excel 时间或字符串时间转换为当天累计分钟。"""
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"时间应精确到分钟：{value!r}")
        return value.hour * 60 + value.minute

    if isinstance(value, str):
        text = value.strip()
        if text == "0:00+1":
            return 24 * 60
        try:
            hour, minute = map(int, text.split(":"))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"无法识别附件中的时间：{value!r}") from exc
        if 0 <= hour < 24 and 0 <= minute < 60:
            return hour * 60 + minute

    raise ValueError(f"无法识别附件中的时间：{value!r}")


def format_clock(minutes: int) -> str:
    """将 0 至 1440 分钟格式化为明确的当天时刻。"""
    if not 0 <= minutes <= 24 * 60:
        raise ValueError(f"分钟数超出当天范围：{minutes}")
    if minutes == 24 * 60:
        return "24:00"
    hour, minute = divmod(minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def validate_number(value: object, row_number: int, column_name: str) -> float:
    """检查附件中的模型输入是否为有限非负数。"""
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not isfinite(value)
        or value < 0
    ):
        raise ValueError(
            f"附件 1 第 {row_number} 行的{column_name}应为有限非负数：{value!r}"
        )
    return float(value)


def load_daily_data() -> DailyData:
    """只读加载附件 1，并按区间结束时刻假设构造 144 个调度时段。"""
    workbook = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("附件 1 缺少工作表 Sheet1。")
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows or tuple(rows[0]) != INPUT_COLUMNS:
        raise ValueError(f"附件 1 的表头应为：{INPUT_COLUMNS}")
    if len(rows) != N_PERIODS + 1:
        raise ValueError(
            f"附件 1 应包含 {N_PERIODS} 条数据，实际为 {len(rows) - 1} 条。"
        )

    end_minutes = [time_to_minutes(row[0]) for row in rows[1:]]
    expected_minutes = list(range(10, 24 * 60 + 1, 10))
    if end_minutes != expected_minutes:
        raise ValueError("时间应从 00:10 到 24:00 连续排列，间隔为 10 分钟。")

    price = np.array(
        [validate_number(row[1], i, "电价") for i, row in enumerate(rows[1:], 2)]
    )
    load_power = np.array(
        [
            validate_number(row[2], i, "小区负载")
            for i, row in enumerate(rows[1:], 2)
        ]
    )
    photovoltaic_power = np.array(
        [
            validate_number(row[3], i, "光伏发电预测功率")
            for i, row in enumerate(rows[1:], 2)
        ]
    )

    interval_labels = [
        f"{format_clock(end - 10)}-{format_clock(end)}" for end in end_minutes
    ]
    return DailyData(
        interval_labels=interval_labels,
        price=price,
        load_power=load_power,
        photovoltaic_power=photovoltaic_power,
        load_energy=load_power * DELTA_T,
        photovoltaic_energy=photovoltaic_power * DELTA_T,
    )


def solve_dispatch(data: DailyData) -> DispatchResult:
    """建立并求解 docs/Q1.md 中的 MILP 模型。"""
    offset_grid = 0
    offset_charge = N_PERIODS
    offset_discharge = 2 * N_PERIODS
    offset_curtailment = 3 * N_PERIODS
    offset_mode = 4 * N_PERIODS
    offset_energy = 5 * N_PERIODS
    variable_count = 6 * N_PERIODS + 1

    objective = np.zeros(variable_count)
    objective[offset_grid : offset_grid + N_PERIODS] = data.price

    lower_bounds = np.zeros(variable_count)
    upper_bounds = np.full(variable_count, np.inf)
    upper_bounds[offset_charge : offset_charge + N_PERIODS] = CHARGE_LIMIT
    upper_bounds[offset_discharge : offset_discharge + N_PERIODS] = (
        DISCHARGE_LIMIT
    )
    upper_bounds[offset_curtailment : offset_curtailment + N_PERIODS] = (
        data.photovoltaic_energy
    )
    upper_bounds[offset_mode : offset_mode + N_PERIODS] = 1.0
    lower_bounds[offset_energy:] = E_MIN
    upper_bounds[offset_energy:] = E_MAX
    lower_bounds[offset_energy] = E_INITIAL
    upper_bounds[offset_energy] = E_INITIAL
    lower_bounds[offset_energy + N_PERIODS] = E_INITIAL
    upper_bounds[offset_energy + N_PERIODS] = E_INITIAL

    integrality = np.zeros(variable_count, dtype=int)
    integrality[offset_mode : offset_mode + N_PERIODS] = 1

    constraint_count = 4 * N_PERIODS
    matrix = lil_matrix((constraint_count, variable_count), dtype=float)
    constraint_lower = np.full(constraint_count, -np.inf)
    constraint_upper = np.full(constraint_count, np.inf)

    for t in range(N_PERIODS):
        balance_row = t
        matrix[balance_row, offset_grid + t] = 1.0
        matrix[balance_row, offset_charge + t] = -1.0
        matrix[balance_row, offset_discharge + t] = 1.0
        matrix[balance_row, offset_curtailment + t] = -1.0
        balance_rhs = data.load_energy[t] - data.photovoltaic_energy[t]
        constraint_lower[balance_row] = balance_rhs
        constraint_upper[balance_row] = balance_rhs

        state_row = N_PERIODS + t
        matrix[state_row, offset_charge + t] = -ETA_CHARGE
        matrix[state_row, offset_discharge + t] = 1.0 / ETA_DISCHARGE
        matrix[state_row, offset_energy + t] = -1.0
        matrix[state_row, offset_energy + t + 1] = 1.0
        constraint_lower[state_row] = 0.0
        constraint_upper[state_row] = 0.0

        charge_mode_row = 2 * N_PERIODS + t
        matrix[charge_mode_row, offset_charge + t] = 1.0
        matrix[charge_mode_row, offset_mode + t] = -CHARGE_LIMIT
        constraint_upper[charge_mode_row] = 0.0

        discharge_mode_row = 3 * N_PERIODS + t
        matrix[discharge_mode_row, offset_discharge + t] = 1.0
        matrix[discharge_mode_row, offset_mode + t] = DISCHARGE_LIMIT
        constraint_upper[discharge_mode_row] = DISCHARGE_LIMIT

    result = milp(
        c=objective,
        integrality=integrality,
        bounds=Bounds(lower_bounds, upper_bounds),
        constraints=LinearConstraint(
            matrix.tocsr(), constraint_lower, constraint_upper
        ),
        options={"disp": False, "mip_rel_gap": 1e-9},
    )

    if not result.success or result.x is None:
        raise RuntimeError(
            f"MILP 求解失败，状态码 {result.status}：{result.message}"
        )

    solution = result.x
    return DispatchResult(
        grid_purchase=solution[offset_grid : offset_grid + N_PERIODS],
        charge=solution[offset_charge : offset_charge + N_PERIODS],
        discharge=solution[offset_discharge : offset_discharge + N_PERIODS],
        curtailment=solution[
            offset_curtailment : offset_curtailment + N_PERIODS
        ],
        mode=solution[offset_mode : offset_mode + N_PERIODS],
        stored_energy=solution[offset_energy:],
        objective_value=float(result.fun),
        solver_message=str(result.message),
    )


def check_solution(data: DailyData, result: DispatchResult) -> ModelChecks:
    """按 Q1.md 的约束和守恒关系独立检查最优解。"""
    balance_residual = (
        result.grid_purchase
        + data.photovoltaic_energy
        - result.curtailment
        + result.discharge
        - data.load_energy
        - result.charge
    )
    state_residual = (
        result.stored_energy[1:]
        - result.stored_energy[:-1]
        - ETA_CHARGE * result.charge
        + result.discharge / ETA_DISCHARGE
    )
    storage_loss = np.sum(
        (1 - ETA_CHARGE) * result.charge
        + (1 / ETA_DISCHARGE - 1) * result.discharge
    )
    total_energy_residual = (
        np.sum(result.grid_purchase)
        + np.sum(data.photovoltaic_energy - result.curtailment)
        - np.sum(data.load_energy)
        - storage_loss
    )
    baseline_purchase = np.maximum(
        data.load_energy - data.photovoltaic_energy, 0.0
    )
    no_storage_cost = float(np.dot(data.price, baseline_purchase))

    checks = ModelChecks(
        max_balance_residual=float(np.max(np.abs(balance_residual))),
        max_state_residual=float(np.max(np.abs(state_residual))),
        max_simultaneous_energy=float(
            np.max(np.minimum(result.charge, result.discharge))
        ),
        max_mode_integrality_error=float(
            np.max(np.abs(result.mode - np.rint(result.mode)))
        ),
        min_stored_energy=float(np.min(result.stored_energy)),
        max_stored_energy=float(np.max(result.stored_energy)),
        terminal_energy_error=float(abs(result.stored_energy[-1] - E_INITIAL)),
        total_energy_residual=float(abs(total_energy_residual)),
        no_storage_cost=no_storage_cost,
        optimized_cost=result.objective_value,
    )

    failures: list[str] = []
    if checks.max_balance_residual > BALANCE_TOLERANCE:
        failures.append("时段电能平衡残差超限")
    if checks.max_state_residual > BALANCE_TOLERANCE:
        failures.append("储能状态转移残差超限")
    if checks.max_simultaneous_energy > BOUND_TOLERANCE:
        failures.append("存在同时充放电")
    if checks.max_mode_integrality_error > INTEGER_TOLERANCE:
        failures.append("储能模式变量不是整数")
    if checks.min_stored_energy < E_MIN - BOUND_TOLERANCE:
        failures.append("储电量低于安全下限")
    if checks.max_stored_energy > E_MAX + BOUND_TOLERANCE:
        failures.append("储电量高于安全上限")
    if checks.terminal_energy_error > BALANCE_TOLERANCE:
        failures.append("24:00 储电量未回归初值")
    if checks.total_energy_residual > BALANCE_TOLERANCE:
        failures.append("全天能量守恒残差超限")
    if checks.optimized_cost > checks.no_storage_cost + BALANCE_TOLERANCE:
        failures.append("最优费用高于不使用储能的基准费用")
    if np.min(result.grid_purchase) < -BOUND_TOLERANCE:
        failures.append("购电量出现负值")
    if np.min(result.curtailment) < -BOUND_TOLERANCE:
        failures.append("弃光量出现负值")
    if np.max(result.curtailment - data.photovoltaic_energy) > BOUND_TOLERANCE:
        failures.append("弃光量超过可用光伏电量")
    if np.max(result.charge) > CHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("充电量超过时段上限")
    if np.max(result.discharge) > DISCHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("放电量超过时段上限")

    if failures:
        raise RuntimeError("模型结果校验失败：" + "；".join(failures))
    return checks


def display_number(value: float) -> str:
    """统一结果精度，并清除数值求解产生的极小正负误差。"""
    if abs(value) < 0.5 * 10 ** (-DISPLAY_DECIMALS):
        value = 0.0
    return f"{value:.{DISPLAY_DECIMALS}f}"


def write_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    """以 Excel 可直接识别中文的 UTF-8 BOM 编码写出 CSV。"""
    with path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow(header)
        writer.writerows(rows)


def write_full_strategy(data: DailyData, result: DispatchResult) -> Path:
    """输出每个 10 分钟时段的完整调度变量。"""
    path = OUTPUT_DIR / "1_完整计划购电策略.csv"
    rows: list[list[object]] = []
    for t in range(N_PERIODS):
        rows.append(
            [
                t + 1,
                data.interval_labels[t],
                display_number(data.price[t]),
                display_number(data.load_power[t]),
                display_number(data.photovoltaic_power[t]),
                display_number(data.load_energy[t]),
                display_number(data.photovoltaic_energy[t]),
                display_number(result.grid_purchase[t]),
                display_number(result.charge[t]),
                display_number(result.discharge[t]),
                display_number(result.curtailment[t]),
                display_number(result.stored_energy[t]),
                display_number(result.stored_energy[t + 1]),
                display_number(data.price[t] * result.grid_purchase[t]),
            ]
        )

    write_csv(
        path,
        [
            "序号",
            "时间段",
            "电价（元/kWh）",
            "小区负载功率（kW）",
            "光伏预测功率（kW）",
            "负载电量（kWh）",
            "可用光伏电量（kWh）",
            "计划购电量（kWh）",
            "充电量（kWh）",
            "放电量（kWh）",
            "弃光量（kWh）",
            "时段初储电量（kWh）",
            "时段末储电量（kWh）",
            "时段购电费（元）",
        ],
        rows,
    )
    return path


def write_purchase_summary(data: DailyData, result: DispatchResult) -> Path:
    """输出题目表 1 指定时段以及全天购电结果。"""
    path = OUTPUT_DIR / "1_指定时段购电量及全天结果.csv"
    requested_starts = (10, 12, 14, 16, 18, 20)
    rows: list[list[object]] = []
    for hour in requested_starts:
        index = hour * 6
        rows.append(
            [
                "指定时段购电量",
                data.interval_labels[index],
                display_number(result.grid_purchase[index]),
                "kWh",
            ]
        )
    rows.extend(
        [
            [
                "全天购电量",
                "00:00-24:00",
                display_number(float(np.sum(result.grid_purchase))),
                "kWh",
            ],
            [
                "全天购电费",
                "00:00-24:00",
                display_number(result.objective_value),
                "元",
            ],
        ]
    )
    write_csv(path, ["指标", "时间段", "数值", "单位"], rows)
    return path


def write_storage_summary(data: DailyData, result: DispatchResult) -> Path:
    """输出题目表 2 所需的 4 小时充放电汇总和边界储电量。"""
    path = OUTPUT_DIR / "1_分时段充放电量及储电量.csv"
    rows: list[list[object]] = []
    periods_per_block = 4 * 6
    for block in range(6):
        start = block * periods_per_block
        stop = (block + 1) * periods_per_block
        rows.append(
            [
                "分时段充放电量",
                f"{format_clock(block * 240)}-{format_clock((block + 1) * 240)}",
                display_number(float(np.sum(result.charge[start:stop]))),
                display_number(float(np.sum(result.discharge[start:stop]))),
                "",
            ]
        )
    rows.extend(
        [
            [
                "边界储电量",
                "00:00",
                "",
                "",
                display_number(result.stored_energy[0]),
            ],
            [
                "边界储电量",
                "24:00",
                "",
                "",
                display_number(result.stored_energy[-1]),
            ],
        ]
    )
    write_csv(
        path,
        ["指标", "时间段或时刻", "充电量（kWh）", "放电量（kWh）", "储电量（kWh）"],
        rows,
    )
    return path


def write_model_checks(result: DispatchResult, checks: ModelChecks) -> Path:
    """输出求解状态、成本对比和主要约束残差。"""
    path = OUTPUT_DIR / "1_模型校验.csv"
    rows = [
        ["求解状态", result.solver_message, "", "通过"],
        [
            "最大时段电能平衡残差",
            display_number(checks.max_balance_residual),
            "kWh",
            "通过",
        ],
        [
            "最大状态转移残差",
            display_number(checks.max_state_residual),
            "kWh",
            "通过",
        ],
        [
            "最大同时充放电量",
            display_number(checks.max_simultaneous_energy),
            "kWh",
            "通过",
        ],
        [
            "模式变量最大整数误差",
            display_number(checks.max_mode_integrality_error),
            "",
            "通过",
        ],
        [
            "最小储电量",
            display_number(checks.min_stored_energy),
            "kWh",
            "通过",
        ],
        [
            "最大储电量",
            display_number(checks.max_stored_energy),
            "kWh",
            "通过",
        ],
        [
            "日末储电量误差",
            display_number(checks.terminal_energy_error),
            "kWh",
            "通过",
        ],
        [
            "全天能量守恒残差",
            display_number(checks.total_energy_residual),
            "kWh",
            "通过",
        ],
        [
            "不使用储能的基准购电费",
            display_number(checks.no_storage_cost),
            "元",
            "参考",
        ],
        [
            "优化后购电费",
            display_number(checks.optimized_cost),
            "元",
            "通过",
        ],
        [
            "相对基准节省费用",
            display_number(checks.no_storage_cost - checks.optimized_cost),
            "元",
            "通过",
        ],
    ]
    write_csv(path, ["校验项", "数值或状态", "单位", "结论"], rows)
    return path


def main() -> None:
    """读取数据、求解、校验并输出问题一结果。"""
    data = load_daily_data()
    result = solve_dispatch(data)
    checks = check_solution(data, result)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_files = [
        write_full_strategy(data, result),
        write_purchase_summary(data, result),
        write_storage_summary(data, result),
        write_model_checks(result, checks),
    ]

    print(f"求解状态：{result.solver_message}")
    print(f"全天购电量：{np.sum(result.grid_purchase):.6f} kWh")
    print(f"全天购电费：{result.objective_value:.6f} 元")
    print(f"基准购电费：{checks.no_storage_cost:.6f} 元")
    print(f"节省费用：{checks.no_storage_cost - checks.optimized_cost:.6f} 元")
    print(f"最大电能平衡残差：{checks.max_balance_residual:.3e} kWh")
    print(f"全天能量守恒残差：{checks.total_energy_residual:.3e} kWh")
    for output_file in output_files:
        print(f"已保存：{output_file}")


if __name__ == "__main__":
    main()
