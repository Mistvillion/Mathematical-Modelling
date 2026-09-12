"""问题二全知视角（完美预见）参考模型：单一全时段 MILP 的全局最优解。

依据 docs/Q2全知视角.md 实现，与 scripts_tables/2_rolling_optimization.py 的
因果滚动模型互为参照。全知视角在制定计划时即可使用附件 2 中 2 月 1 日至
12 月 31 日的全部实际负载和实际光伏，因此：

1. 预测没有误差，0:00 确定的计划购电量就是实际购电量，不需要紧急购电，
   弃光量仍由优化决定；
2. 模型把 334 天共 48096 个 10 分钟时段一次性放入同一个混合整数线性规划，
   跨日储电量连续，所以结果是完美预见下的全局最优解，而不是逐日贪心解；
3. 该最优费用是问题二因果滚动模型实际费用的理论下界，只能作为论文参考，
   不能作为附件 5 的正式结果提交。

求解前先解 LP 松弛：若 LP 最优解满足充放电互补条件 c_t d_t = 0，则把它按
模式修复后即为 MILP 最优解，可在不做分支定界的情况下认证全局最优；否则
自动转入 MILP 分支定界，并同时报告 LP 下界、MILP 对偶界和相对间隙。

在 conda 环境 2026C 中运行：
    python scripts_tables/2_omniscient_optimization.py

输出：outputs/tables/2_全知视角*.csv 和 outputs/results/result2_全知视角.xlsx。
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
import tempfile
import time as timing
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time
from math import isfinite
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import coo_matrix


PROJECT_DIR = Path(__file__).resolve().parents[1]
ATTACHMENT_DIR = PROJECT_DIR / "CUMCM 2026 C题" / "附件"
PRICE_FILE = ATTACHMENT_DIR / "附件1.xlsx"
ACTUAL_DATA_FILE = ATTACHMENT_DIR / "附件2.xlsx"
RESULT_TEMPLATE_FILE = ATTACHMENT_DIR / "附件5" / "result2.xlsx"
OUTPUT_TABLE_DIR = PROJECT_DIR / "outputs" / "tables"
OUTPUT_RESULT_DIR = PROJECT_DIR / "outputs" / "results"
RESULT_FILE = OUTPUT_RESULT_DIR / "result2_全知视角.xlsx"
CAUSAL_CHECK_FILE = OUTPUT_TABLE_DIR / "2_模型校验.csv"

DETAIL_FILE = OUTPUT_TABLE_DIR / "2_全知视角调度明细.csv"
DAILY_FILE = OUTPUT_TABLE_DIR / "2_全知视角逐日费用.csv"
CHECK_FILE = OUTPUT_TABLE_DIR / "2_全知视角模型校验.csv"
PURCHASE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_全知视角指定日期购电量及全天结果.csv"
STORAGE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_全知视角指定日期充放电量及储电量.csv"
EMERGENCY_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_全知视角指定日期紧急购电量.csv"
COMPARISON_FILE = OUTPUT_TABLE_DIR / "2_全知视角与因果模型对比.csv"

N_PERIODS = 144
DELTA_T = 1 / 6
YEAR_START = date(2025, 1, 1)
OFFICIAL_START = date(2025, 2, 1)
YEAR_END = date(2025, 12, 31)
OFFICIAL_DAYS = 334

E_RATED = 12000.0
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
E_DAY_MIN = 5400.0
E_DAY_MAX = 6600.0
P_CHARGE_MAX = 5000.0
P_DISCHARGE_MAX = 5000.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
CHARGE_LIMIT = P_CHARGE_MAX * DELTA_T
DISCHARGE_LIMIT = P_DISCHARGE_MAX * DELTA_T
EMERGENCY_PRICE_MULTIPLIER = 5.0
PHOTOVOLTAIC_NOISE_THRESHOLD = 1.0

DEFAULT_MIP_REL_GAP = 1e-6
DEFAULT_TIME_LIMIT_SECONDS = 7200.0
COMPLEMENTARITY_TOLERANCE = 1e-6

BALANCE_TOLERANCE = 1e-5
BOUND_TOLERANCE = 1e-5
INTEGER_TOLERANCE = 1e-6
TERMINAL_TOLERANCE = 1e-4
OBJECTIVE_TOLERANCE = 1e-4
DISPLAY_DECIMALS = 6
DISPLAY_ZERO_TOLERANCE = 0.5 * 10 ** (-DISPLAY_DECIMALS)

ATTACHMENT1_COLUMNS = ("时间", "电价", "小区负载", "光伏发电预测功率")
ACTUAL_SHEET_NAMES = ("小区负载", "光伏发电实际功率")
RESULT_SHEET_NAMES = ("计划购电量", "充放电量", "紧急购电量")
STORAGE_BLOCK_LABELS = (
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
)
SPECIFIED_DATES = (
    date(2025, 3, 20),
    date(2025, 6, 21),
    date(2025, 9, 23),
    date(2025, 12, 21),
)
SPECIFIED_HOURS = (10, 12, 14, 16, 18, 20)


@dataclass(frozen=True)
class ModelData:
    """完成校验、清洗和时间对齐后的全部读取数据。"""

    dates: tuple[date, ...]
    interval_labels: tuple[str, ...]
    price: np.ndarray
    actual_load_power: np.ndarray
    actual_photovoltaic_power: np.ndarray
    actual_load_energy: np.ndarray
    actual_photovoltaic_energy: np.ndarray


@dataclass(frozen=True)
class OmniscientProblem:
    """正式期全时段单一 MILP 的完整结构化描述。"""

    day_start: int
    day_count: int
    period_count: int
    initial_energy: float
    price: np.ndarray
    load_energy: np.ndarray
    photovoltaic_energy: np.ndarray
    objective: np.ndarray
    lower_bounds: np.ndarray
    upper_bounds: np.ndarray
    integrality: np.ndarray
    constraint_matrix: object
    constraint_lower: np.ndarray
    constraint_upper: np.ndarray
    day_end_energy_indices: np.ndarray
    offset_grid: int
    offset_charge: int
    offset_discharge: int
    offset_curtailment: int
    offset_mode: int
    offset_energy: int


@dataclass(frozen=True)
class OmniscientSolution:
    """全知视角的购电与储能全局最优解及求解证书。"""

    certificate: str
    is_certified: bool
    grid_purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    curtailment: np.ndarray
    mode: np.ndarray
    stored_energy: np.ndarray
    objective_value: float
    lp_value: float | None
    lp_overlap: float | None
    milp_dual_bound: float | None
    mip_gap: float | None
    solver_message: str
    solver_status: int
    lp_message: str | None
    lp_seconds: float | None
    milp_seconds: float | None


@dataclass(frozen=True)
class StageResult:
    """一次 LP 或 MILP 求解的结果摘要。"""

    success: bool
    status: int
    message: str
    objective_value: float | None
    dual_bound: float | None
    mip_gap: float | None
    seconds: float
    solution: np.ndarray | None


@dataclass(frozen=True)
class CheckMetrics:
    """全知视角解的独立约束校验指标。"""

    max_balance_residual: float
    max_state_residual: float
    max_simultaneous_overlap: float
    max_mode_integer_error: float
    min_stored_energy: float
    max_stored_energy: float
    min_terminal_energy: float
    max_terminal_energy: float
    max_terminal_overrun: float
    max_power_limit_overrun: float
    max_curtailment_overrun: float
    objective_recompute_error: float
    total_energy_residual: float


def normalize_excel_date(value: object, location: str) -> date:
    """把 Excel 日期单元格统一转换为不含时分秒的 date。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"{location}应为 Excel 日期，实际为：{value!r}")


def time_to_end_minutes(value: time | str) -> int:
    """把附件的区间结束时刻转换为当天累计分钟。"""
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"时间应精确到分钟：{value!r}")
        return value.hour * 60 + value.minute

    if isinstance(value, str):
        text_value = value.strip()
        if text_value == "0:00+1":
            return 24 * 60
        try:
            hour, minute = map(int, text_value.split(":"))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"无法识别附件中的时间：{value!r}") from exc
        if 0 <= hour < 24 and 0 <= minute < 60:
            return hour * 60 + minute

    raise ValueError(f"无法识别附件中的时间：{value!r}")


def format_clock(minutes: int) -> str:
    """将 0 至 1440 分钟格式化为明确的当天时刻。"""
    if not 0 <= minutes <= 24 * 60:
        raise ValueError(f"分钟数超出当天范围：{minutes}")
    if minutes == 24 * 60:
        return "24:00"
    hour, minute = divmod(minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def validate_nonnegative_number(value: object, location: str) -> float:
    """检查输入是否为有限非负实数。"""
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not isfinite(value)
        or value < 0
    ):
        raise ValueError(f"{location}应为有限非负数，实际为：{value!r}")
    return float(value)


def validate_positive_number(value: object, location: str) -> float:
    """检查输入是否为有限正数；全知模型要求电价严格为正。"""
    number = validate_nonnegative_number(value, location)
    if number <= 0:
        raise ValueError(f"{location}应为严格正数，实际为：{value!r}")
    return number


def display_number(value: float) -> str:
    """按展示精度格式化数值，并清除极小的正负误差。"""
    if abs(value) < DISPLAY_ZERO_TOLERANCE:
        value = 0.0
    return f"{value:.{DISPLAY_DECIMALS}f}"


def excel_number(value: float) -> float:
    """仅在写出 Excel 时取六位小数并消除负零。"""
    return float(display_number(float(value)))


def write_csv(
    path: Path,
    header: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> Path:
    """用 Excel 可直接识别的 UTF-8 BOM 编码写出 CSV。"""
    with path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.writer(csv_file, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    return path


def read_price_curve() -> np.ndarray:
    """只读加载附件 1 的重复日内电价，并校验其严格为正。"""
    workbook = load_workbook(PRICE_FILE, read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("附件 1 缺少工作表 Sheet1。")
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows or tuple(rows[0]) != ATTACHMENT1_COLUMNS:
        raise ValueError(f"附件 1 的表头应为：{ATTACHMENT1_COLUMNS}")
    if len(rows) != N_PERIODS + 1:
        raise ValueError(
            f"附件 1 应包含 {N_PERIODS} 条数据，实际为 {len(rows) - 1} 条。"
        )
    end_minutes = [time_to_end_minutes(row[0]) for row in rows[1:]]
    if end_minutes != list(range(10, 24 * 60 + 1, 10)):
        raise ValueError("附件 1 的时间应从 0:10 到 0:00+1，每 10 分钟连续排列。")
    return np.array(
        [
            validate_positive_number(row[1], f"附件 1 第 {row_index} 行电价")
            for row_index, row in enumerate(rows[1:], start=2)
        ],
        dtype=float,
    )


def read_actual_sheet(
    worksheet_name: str,
    rows: list[tuple[object, ...]],
) -> tuple[tuple[date, ...], tuple[int, ...], np.ndarray]:
    """校验并读取附件 2 的一个 365 天功率工作表。"""
    if not rows:
        raise ValueError(f"附件 2 的“{worksheet_name}”工作表为空。")
    if len(rows) != 366:
        raise ValueError(
            f"附件 2 的“{worksheet_name}”应包含表头和 365 天数据，"
            f"实际共有 {len(rows)} 行。"
        )
    if len(rows[0]) != N_PERIODS + 1 or rows[0][0] != "日期\\时间":
        raise ValueError(
            f"附件 2 的“{worksheet_name}”首列应为“日期\\时间”，"
            f"并包含 {N_PERIODS} 个时段。"
        )
    end_minutes = tuple(time_to_end_minutes(value) for value in rows[0][1:])
    if end_minutes != tuple(range(10, 24 * 60 + 1, 10)):
        raise ValueError(
            f"附件 2 的“{worksheet_name}”时间应从 0:10 到 0:00+1 连续排列。"
        )

    dates = tuple(
        normalize_excel_date(row[0], f"“{worksheet_name}”第 {row_index} 行日期")
        for row_index, row in enumerate(rows[1:], start=2)
    )
    expected_dates = tuple(
        date.fromordinal(YEAR_START.toordinal() + day_offset)
        for day_offset in range(365)
    )
    if dates != expected_dates:
        raise ValueError(
            f"附件 2 的“{worksheet_name}”日期必须从 {YEAR_START} 连续到 {YEAR_END}。"
        )

    power = np.empty((365, N_PERIODS), dtype=float)
    for day_index, row in enumerate(rows[1:]):
        if len(row) != N_PERIODS + 1:
            raise ValueError(
                f"附件 2 的“{worksheet_name}”第 {day_index + 2} 行列数不正确。"
            )
        for period_index, value in enumerate(row[1:]):
            power[day_index, period_index] = validate_nonnegative_number(
                value,
                f"“{worksheet_name}”第 {day_index + 2} 行第 {period_index + 2} 列",
            )
    return dates, end_minutes, power


def load_model_data() -> ModelData:
    """加载附件 1 电价和附件 2 实际负载、实际光伏。"""
    price = read_price_curve()

    workbook = load_workbook(ACTUAL_DATA_FILE, read_only=True, data_only=True)
    try:
        missing_sheets = set(ACTUAL_SHEET_NAMES).difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "附件 2 缺少工作表：" + "、".join(sorted(missing_sheets))
            )
        load_rows = list(workbook[ACTUAL_SHEET_NAMES[0]].iter_rows(values_only=True))
        photovoltaic_rows = list(
            workbook[ACTUAL_SHEET_NAMES[1]].iter_rows(values_only=True)
        )
    finally:
        workbook.close()

    load_dates, load_end_minutes, load_power = read_actual_sheet(
        ACTUAL_SHEET_NAMES[0], load_rows
    )
    photovoltaic_dates, photovoltaic_end_minutes, photovoltaic_power = (
        read_actual_sheet(ACTUAL_SHEET_NAMES[1], photovoltaic_rows)
    )
    if load_dates != photovoltaic_dates:
        raise ValueError("附件 2 的负载和光伏工作表日期不一致。")
    if load_end_minutes != photovoltaic_end_minutes:
        raise ValueError("附件 2 的负载和光伏工作表时间表头不一致。")

    # 与 Q2 相同的固定阈值清洗：只把近零正值归零，不填补白天真实零出力。
    photovoltaic_power = photovoltaic_power.copy()
    photovoltaic_power[photovoltaic_power < PHOTOVOLTAIC_NOISE_THRESHOLD] = 0.0
    interval_labels = tuple(
        f"{format_clock(end_minute - 10)}-{format_clock(end_minute)}"
        for end_minute in load_end_minutes
    )
    return ModelData(
        dates=load_dates,
        interval_labels=interval_labels,
        price=price,
        actual_load_power=load_power,
        actual_photovoltaic_power=photovoltaic_power,
        actual_load_energy=load_power * DELTA_T,
        actual_photovoltaic_energy=photovoltaic_power * DELTA_T,
    )


def horizon_start_index(data: ModelData) -> int:
    """定位正式期首日，允许附件 2 起始日期发生变化。"""
    try:
        day_start = data.dates.index(OFFICIAL_START)
    except ValueError as exc:
        raise ValueError(f"附件 2 缺少正式期首日 {OFFICIAL_START}。") from exc
    if day_start + OFFICIAL_DAYS > len(data.dates):
        raise ValueError("附件 2 的日期不足以覆盖 2 月 1 日至 12 月 31 日。")
    if data.dates[day_start + OFFICIAL_DAYS - 1] != YEAR_END:
        raise ValueError(f"正式期最后一天应为 {YEAR_END}。")
    return day_start


def validate_initial_energy(initial_energy: float) -> float:
    """初始储电量必须是有限值并落在设备安全范围内。"""
    if not isfinite(initial_energy):
        raise ValueError("初始储电量必须是有限数值。")
    if not E_MIN - BOUND_TOLERANCE <= initial_energy <= E_MAX + BOUND_TOLERANCE:
        raise ValueError(
            f"初始储电量 {initial_energy:.6f} kWh 超出设备安全范围"
            f"[{E_MIN:g}, {E_MAX:g}] kWh。"
        )
    return float(initial_energy)


def build_problem(data: ModelData, initial_energy: float) -> OmniscientProblem:
    """构造 2 月 1 日至 12 月 31 日全时段的单一 MILP 数据结构。

    变量顺序为（g, c, d, r, u, e），其中 u 为充放电模式二元变量，e 含 T+1 个
    状态边界。日末硬区间直接写成 e 在每天末边界上的上下界，不额外增加约束行。
    """
    day_start = horizon_start_index(data)
    day_count = OFFICIAL_DAYS
    period_count = day_count * N_PERIODS
    initial_energy = validate_initial_energy(initial_energy)

    load = np.asarray(
        data.actual_load_energy[day_start : day_start + day_count], dtype=float
    ).reshape(-1)
    photovoltaic = np.asarray(
        data.actual_photovoltaic_energy[day_start : day_start + day_count],
        dtype=float,
    ).reshape(-1)
    price = np.tile(np.asarray(data.price, dtype=float), day_count)
    if not (
        np.isfinite(load).all()
        and np.isfinite(photovoltaic).all()
        and np.isfinite(price).all()
    ):
        raise ValueError("全时段负载、光伏或电价存在非有限值。")

    offset_grid = 0
    offset_charge = period_count
    offset_discharge = 2 * period_count
    offset_curtailment = 3 * period_count
    offset_mode = 4 * period_count
    offset_energy = 5 * period_count
    variable_count = 6 * period_count + 1

    objective = np.zeros(variable_count)
    objective[offset_grid : offset_grid + period_count] = price

    lower_bounds = np.zeros(variable_count)
    upper_bounds = np.full(variable_count, np.inf)
    upper_bounds[offset_charge : offset_charge + period_count] = CHARGE_LIMIT
    upper_bounds[offset_discharge : offset_discharge + period_count] = (
        DISCHARGE_LIMIT
    )
    upper_bounds[offset_curtailment : offset_curtailment + period_count] = (
        photovoltaic
    )
    upper_bounds[offset_mode : offset_mode + period_count] = 1.0
    lower_bounds[offset_energy : offset_energy + period_count + 1] = E_MIN
    upper_bounds[offset_energy : offset_energy + period_count + 1] = E_MAX
    lower_bounds[offset_energy] = initial_energy
    upper_bounds[offset_energy] = initial_energy

    day_end_energy_indices = np.arange(1, day_count + 1) * N_PERIODS
    lower_bounds[offset_energy + day_end_energy_indices] = E_DAY_MIN
    upper_bounds[offset_energy + day_end_energy_indices] = E_DAY_MAX

    integrality = np.zeros(variable_count, dtype=np.int32)
    integrality[offset_mode : offset_mode + period_count] = 1

    period = np.arange(period_count)
    # 四组约束行：电量平衡（含计划购电、充电、放电、弃光）、储电量转移、
    # 充电模式上界 c <= C*u、放电模式上界 d <= D*(1-u)。
    balance_rows = period
    state_rows = period_count + period
    charge_mode_rows = 2 * period_count + period
    discharge_mode_rows = 3 * period_count + period
    rows = np.concatenate(
        [
            balance_rows,
            balance_rows,
            balance_rows,
            balance_rows,
            state_rows,
            state_rows,
            state_rows,
            state_rows,
            charge_mode_rows,
            charge_mode_rows,
            discharge_mode_rows,
            discharge_mode_rows,
        ]
    )
    columns = np.concatenate(
        [
            offset_grid + period,
            offset_charge + period,
            offset_discharge + period,
            offset_curtailment + period,
            offset_charge + period,
            offset_discharge + period,
            offset_energy + period,
            offset_energy + period + 1,
            offset_charge + period,
            offset_mode + period,
            offset_discharge + period,
            offset_mode + period,
        ]
    )
    values = np.concatenate(
        [
            np.ones(period_count),
            -np.ones(period_count),
            np.ones(period_count),
            -np.ones(period_count),
            -np.full(period_count, ETA_CHARGE),
            np.full(period_count, 1.0 / ETA_DISCHARGE),
            -np.ones(period_count),
            np.ones(period_count),
            np.ones(period_count),
            -np.full(period_count, CHARGE_LIMIT),
            np.ones(period_count),
            np.full(period_count, DISCHARGE_LIMIT),
        ]
    )
    constraint_matrix = coo_matrix(
        (values, (rows, columns)),
        shape=(4 * period_count, variable_count),
    ).tocsc()

    constraint_lower = np.empty(4 * period_count)
    constraint_upper = np.empty(4 * period_count)
    # 电量平衡 g - c + d - r = load - pv；其余为等式转移和两条单向模式上界。
    balance_rhs = load - photovoltaic
    constraint_lower[:period_count] = balance_rhs
    constraint_upper[:period_count] = balance_rhs
    constraint_lower[period_count : 2 * period_count] = 0.0
    constraint_upper[period_count : 2 * period_count] = 0.0
    constraint_lower[2 * period_count : 3 * period_count] = -np.inf
    constraint_upper[2 * period_count : 3 * period_count] = 0.0
    constraint_lower[3 * period_count : 4 * period_count] = -np.inf
    constraint_upper[3 * period_count : 4 * period_count] = DISCHARGE_LIMIT

    return OmniscientProblem(
        day_start=day_start,
        day_count=day_count,
        period_count=period_count,
        initial_energy=initial_energy,
        price=price,
        load_energy=load,
        photovoltaic_energy=photovoltaic,
        objective=objective,
        lower_bounds=lower_bounds,
        upper_bounds=upper_bounds,
        integrality=integrality,
        constraint_matrix=constraint_matrix,
        constraint_lower=constraint_lower,
        constraint_upper=constraint_upper,
        day_end_energy_indices=day_end_energy_indices,
        offset_grid=offset_grid,
        offset_charge=offset_charge,
        offset_discharge=offset_discharge,
        offset_curtailment=offset_curtailment,
        offset_mode=offset_mode,
        offset_energy=offset_energy,
    )


def run_stage(
    problem: OmniscientProblem,
    integrality: np.ndarray | None,
    options: dict[str, object],
) -> StageResult:
    """调用 HiGHS 求解一次 LP 或 MILP，并保留可行解与界的信息。"""
    start_time = timing.perf_counter()
    result = milp(
        c=problem.objective,
        integrality=integrality,
        bounds=Bounds(problem.lower_bounds, problem.upper_bounds),
        constraints=LinearConstraint(
            problem.constraint_matrix,
            problem.constraint_lower,
            problem.constraint_upper,
        ),
        options=options,
    )
    seconds = timing.perf_counter() - start_time

    solution = None
    if getattr(result, "x", None) is not None:
        candidate = np.asarray(result.x, dtype=float)
        if candidate.shape == (len(problem.objective),) and np.isfinite(candidate).all():
            solution = candidate
    objective_value = None
    if solution is not None and getattr(result, "fun", None) is not None:
        candidate_value = float(result.fun)
        if isfinite(candidate_value):
            objective_value = candidate_value
    dual_bound = getattr(result, "mip_dual_bound", None)
    mip_gap = getattr(result, "mip_gap", None)
    return StageResult(
        success=bool(result.success),
        status=int(result.status),
        message=str(result.message),
        objective_value=objective_value,
        dual_bound=float(dual_bound) if dual_bound is not None and isfinite(float(dual_bound)) else None,
        mip_gap=float(mip_gap) if mip_gap is not None and isfinite(float(mip_gap)) else None,
        seconds=seconds,
        solution=solution,
    )


def solve_options(
    time_limit: float,
    mip_rel_gap: float | None,
    solver_log: bool,
) -> dict[str, object]:
    """组装 HiGHS 选项；0 秒表示不设置时间上限。"""
    options: dict[str, object] = {"presolve": True, "disp": solver_log}
    if time_limit > 0:
        options["time_limit"] = float(time_limit)
    if mip_rel_gap is not None:
        options["mip_rel_gap"] = float(mip_rel_gap)
    return options


def split_solution(
    problem: OmniscientProblem,
    vector: np.ndarray,
) -> tuple[
    np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray
]:
    """按变量顺序拆分求解向量并修复充放电模式。"""
    period_count = problem.period_count
    grid_purchase = vector[problem.offset_grid : problem.offset_grid + period_count]
    charge = vector[problem.offset_charge : problem.offset_charge + period_count]
    discharge = vector[
        problem.offset_discharge : problem.offset_discharge + period_count
    ]
    curtailment = vector[
        problem.offset_curtailment : problem.offset_curtailment + period_count
    ]
    mode = vector[problem.offset_mode : problem.offset_mode + period_count]
    stored_energy = vector[
        problem.offset_energy : problem.offset_energy + period_count + 1
    ]
    # LP 松弛解可能给出小数模式；只要互补条件成立，按充电量重新取 0-1 即可，
    # 因为 c <= C*u 与 d <= D*(1-u) 在这种取整下仍然成立。
    repaired_mode = (charge > COMPLEMENTARITY_TOLERANCE).astype(float)
    return (
        grid_purchase,
        charge,
        discharge,
        curtailment,
        mode,
        stored_energy,
        repaired_mode,
    )


def maximum_simultaneous_overlap(
    charge: np.ndarray,
    discharge: np.ndarray,
) -> float:
    """同时充放电重叠量 kappa = max_t min(c_t, d_t)。"""
    return float(np.max(np.minimum(charge, discharge))) if len(charge) else 0.0


def solve_omniscient(
    problem: OmniscientProblem,
    lp_first: bool,
    mip_rel_gap: float,
    time_limit: float,
    solver_log: bool,
) -> OmniscientSolution:
    """先尝试 LP 互补性认证，必要时转入 MILP 分支定界。"""
    lp_stage: StageResult | None = None
    lp_overlap: float | None = None
    if lp_first:
        print("开始求解 LP 松弛……", flush=True)
        lp_stage = run_stage(
            problem,
            integrality=None,
            options=solve_options(time_limit, None, solver_log),
        )
        print(
            f"LP 松弛完成：状态 {lp_stage.status}，耗时 {lp_stage.seconds:.2f} s，"
            f"目标值 {lp_stage.objective_value}",
            flush=True,
        )
        if lp_stage.solution is not None:
            candidate = split_solution(problem, lp_stage.solution)
            lp_overlap = maximum_simultaneous_overlap(candidate[1], candidate[2])
            print(f"LP 松弛解的同时充放电最大重叠量：{lp_overlap:.12g} kWh", flush=True)
            if lp_stage.success and lp_overlap <= COMPLEMENTARITY_TOLERANCE:
                solution = OmniscientSolution(
                    certificate="LP 松弛互补性认证",
                    is_certified=True,
                    grid_purchase=candidate[0],
                    charge=candidate[1],
                    discharge=candidate[2],
                    curtailment=candidate[3],
                    mode=candidate[6],
                    stored_energy=candidate[5],
                    objective_value=float(lp_stage.objective_value),
                    lp_value=float(lp_stage.objective_value),
                    lp_overlap=lp_overlap,
                    milp_dual_bound=None,
                    mip_gap=0.0,
                    solver_message=lp_stage.message,
                    solver_status=lp_stage.status,
                    lp_message=lp_stage.message,
                    lp_seconds=lp_stage.seconds,
                    milp_seconds=None,
                )
                print("LP 松弛解的充放电互补条件成立，已认证为 MILP 全局最优解。", flush=True)
                return solution
            print("LP 松弛解的充放电互补条件不成立，转入 MILP 分支定界。", flush=True)

    print("开始求解全时段 MILP……", flush=True)
    milp_stage = run_stage(
        problem,
        integrality=problem.integrality,
        options=solve_options(time_limit, mip_rel_gap, solver_log),
    )
    if milp_stage.solution is None:
        raise RuntimeError(
            f"全时段 MILP 未返回可行解：状态 {milp_stage.status}，"
            f"信息 {milp_stage.message}。"
        )
    candidate = split_solution(problem, milp_stage.solution)
    gap = milp_stage.mip_gap if milp_stage.mip_gap is not None else 0.0
    if milp_stage.success:
        certificate = f"MILP 分支定界（相对间隙 {gap:.3g}）"
        print(
            f"MILP 已在相对间隙 {gap:.3g} 内证明全局最优，耗时 "
            f"{milp_stage.seconds:.2f} s。",
            flush=True,
        )
    else:
        # 达到时间上限时仍保留可行的上界解，并在校验表中如实标注未证明最优。
        certificate = (
            f"MILP 可行上界解（未证明最优；状态 {milp_stage.status}，"
            f"相对间隙 {gap:.3g}）"
        )
        print(
            f"MILP 在时间上限内未证明最优：状态 {milp_stage.status}，"
            f"相对间隙 {gap:.3g}；保存可行上界解供参考。",
            flush=True,
        )
    lp_value = (
        float(lp_stage.objective_value)
        if lp_stage is not None and lp_stage.objective_value is not None
        else None
    )
    return OmniscientSolution(
        certificate=certificate,
        is_certified=bool(milp_stage.success),
        grid_purchase=candidate[0],
        charge=candidate[1],
        discharge=candidate[2],
        curtailment=candidate[3],
        mode=candidate[6],
        stored_energy=candidate[5],
        objective_value=float(milp_stage.objective_value),
        lp_value=lp_value,
        lp_overlap=lp_overlap,
        milp_dual_bound=milp_stage.dual_bound,
        mip_gap=gap,
        solver_message=milp_stage.message,
        solver_status=milp_stage.status,
        lp_message=lp_stage.message if lp_stage is not None else None,
        lp_seconds=lp_stage.seconds if lp_stage is not None else None,
        milp_seconds=milp_stage.seconds,
    )


def validate_solution(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> CheckMetrics:
    """独立校验全知视角解的全部物理约束、终端区间和费用一致性。"""
    period_count = problem.period_count
    failures: list[str] = []
    for name in ("grid_purchase", "charge", "discharge", "curtailment", "mode"):
        array = getattr(solution, name)
        if array.shape != (period_count,) or not np.isfinite(array).all():
            raise RuntimeError(f"{name} 维度或数值无效。")
    if solution.stored_energy.shape != (period_count + 1,) or not np.isfinite(
        solution.stored_energy
    ).all():
        raise RuntimeError("储电量维度或数值无效。")

    balance_residual = (
        solution.grid_purchase
        + problem.photovoltaic_energy
        + solution.discharge
        - solution.curtailment
        - problem.load_energy
        - solution.charge
    )
    state_residual = (
        solution.stored_energy[1:]
        - solution.stored_energy[:-1]
        - ETA_CHARGE * solution.charge
        + solution.discharge / ETA_DISCHARGE
    )
    overlap = np.minimum(solution.charge, solution.discharge)
    mode_integer_error = float(
        np.max(np.abs(solution.mode - np.rint(solution.mode)))
    )
    terminal_energy = solution.stored_energy[problem.day_end_energy_indices]
    terminal_overrun = np.maximum(
        np.maximum(E_DAY_MIN - terminal_energy, terminal_energy - E_DAY_MAX), 0.0
    )
    power_overrun = float(
        max(
            0.0,
            np.max(solution.charge - CHARGE_LIMIT),
            np.max(solution.discharge - DISCHARGE_LIMIT),
            np.max(-solution.charge),
            np.max(-solution.discharge),
            np.max(-solution.grid_purchase),
        )
    )
    curtailment_overrun = float(
        max(0.0, np.max(solution.curtailment - problem.photovoltaic_energy))
    )
    objective_recompute = float(problem.price @ solution.grid_purchase)
    objective_error = abs(objective_recompute - solution.objective_value)
    total_energy_residual = float(
        np.sum(
            solution.grid_purchase
            + problem.photovoltaic_energy
            - solution.curtailment
            - problem.load_energy
            - (1 - ETA_CHARGE) * solution.charge
            - (1 / ETA_DISCHARGE - 1) * solution.discharge
        )
        - (solution.stored_energy[-1] - problem.initial_energy)
    )

    if np.max(np.abs(balance_residual)) > BALANCE_TOLERANCE:
        failures.append("电量平衡残差超限")
    if np.max(np.abs(state_residual)) > BALANCE_TOLERANCE:
        failures.append("储电量状态转移残差超限")
    if float(np.max(overlap)) > COMPLEMENTARITY_TOLERANCE:
        failures.append("存在同时充放电")
    if mode_integer_error > INTEGER_TOLERANCE:
        failures.append("充放电模式不是二元值")
    if np.max(solution.charge - CHARGE_LIMIT * solution.mode) > BOUND_TOLERANCE:
        failures.append("充电量违反模式约束")
    if np.max(solution.discharge - DISCHARGE_LIMIT * (1 - solution.mode)) > BOUND_TOLERANCE:
        failures.append("放电量违反模式约束")
    if power_overrun > BOUND_TOLERANCE:
        failures.append("充放电或购电量越界")
    if curtailment_overrun > BOUND_TOLERANCE:
        failures.append("弃光量超过可用光伏电量")
    if (
        np.min(solution.stored_energy) < E_MIN - BOUND_TOLERANCE
        or np.max(solution.stored_energy) > E_MAX + BOUND_TOLERANCE
    ):
        failures.append("储电量超出设备安全范围")
    if abs(solution.stored_energy[0] - problem.initial_energy) > BOUND_TOLERANCE:
        failures.append("初始储电量与设定值不一致")
    if float(np.max(terminal_overrun)) > TERMINAL_TOLERANCE:
        failures.append("存在日末储电量超出 [5400, 6600] kWh 区间")
    if objective_error > OBJECTIVE_TOLERANCE + 1e-9 * abs(objective_recompute):
        failures.append("目标函数重算不一致")
    if abs(total_energy_residual) > period_count * BALANCE_TOLERANCE:
        failures.append("全时段能量守恒残差超限")
    if failures:
        raise RuntimeError("全知视角解校验失败：" + "；".join(failures))

    return CheckMetrics(
        max_balance_residual=float(np.max(np.abs(balance_residual))),
        max_state_residual=float(np.max(np.abs(state_residual))),
        max_simultaneous_overlap=float(np.max(overlap)),
        max_mode_integer_error=mode_integer_error,
        min_stored_energy=float(np.min(solution.stored_energy)),
        max_stored_energy=float(np.max(solution.stored_energy)),
        min_terminal_energy=float(np.min(terminal_energy)),
        max_terminal_energy=float(np.max(terminal_energy)),
        max_terminal_overrun=float(np.max(terminal_overrun)),
        max_power_limit_overrun=power_overrun,
        max_curtailment_overrun=curtailment_overrun,
        objective_recompute_error=objective_error,
        total_energy_residual=total_energy_residual,
    )


def daily_slices(problem: OmniscientProblem) -> tuple[slice, ...]:
    """返回每天在全时段数组中的切片。"""
    return tuple(
        slice(day * N_PERIODS, (day + 1) * N_PERIODS)
        for day in range(problem.day_count)
    )


def write_detail_csv(
    data: ModelData,
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> Path:
    """写出正式期 334 天逐 10 分钟的购电、充放电、弃光和储电量明细。"""
    header = (
        "日期",
        "时间段",
        "电价（元/kWh）",
        "实际负载电量（kWh）",
        "实际光伏电量（kWh）",
        "净负荷（kWh）",
        "计划购电量（kWh）",
        "充电量（kWh）",
        "放电量（kWh）",
        "运行模式",
        "弃光量（kWh）",
        "未利用计划购电量（kWh）",
        "紧急购电量（kWh）",
        "时段初储电量（kWh）",
        "时段末储电量（kWh）",
        "时段末SOC（%）",
        "时段购电费（元）",
        "累计购电费（元）",
    )

    def rows() -> Iterable[Sequence[object]]:
        cumulative_cost = 0.0
        for day_offset, day_slice in enumerate(daily_slices(problem)):
            day = data.dates[problem.day_start + day_offset]
            for period_index, interval_label in enumerate(data.interval_labels):
                index = day_slice.start + period_index
                period_cost = (
                    problem.price[index] * solution.grid_purchase[index]
                )
                cumulative_cost += period_cost
                yield (
                    day.isoformat(),
                    interval_label,
                    display_number(problem.price[index]),
                    display_number(problem.load_energy[index]),
                    display_number(problem.photovoltaic_energy[index]),
                    display_number(
                        problem.load_energy[index] - problem.photovoltaic_energy[index]
                    ),
                    display_number(solution.grid_purchase[index]),
                    display_number(solution.charge[index]),
                    display_number(solution.discharge[index]),
                    display_number(solution.mode[index]),
                    display_number(solution.curtailment[index]),
                    display_number(0.0),
                    display_number(0.0),
                    display_number(solution.stored_energy[index]),
                    display_number(solution.stored_energy[index + 1]),
                    display_number(
                        100.0 * solution.stored_energy[index + 1] / E_RATED
                    ),
                    display_number(period_cost),
                    display_number(cumulative_cost),
                )

    return write_csv(DETAIL_FILE, header, rows())


def write_daily_csv(
    data: ModelData,
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> Path:
    """写出正式期每天的购电量、费用、储能动作与日末区间状态。"""
    header = (
        "日期",
        "计划购电量（kWh）",
        "购电费（元）",
        "充电量（kWh）",
        "放电量（kWh）",
        "弃光量（kWh）",
        "紧急购电量（kWh）",
        "紧急购电费（元）",
        "总购电费（元）",
        "累计总购电费（元）",
        "0:00储电量（kWh）",
        "24:00储电量（kWh）",
        "日末储电量距下限（kWh）",
        "日末储电量距上限（kWh）",
        "日末区间越界量（kWh）",
    )
    day_energy = solution.stored_energy[::N_PERIODS]

    def rows() -> Iterable[Sequence[object]]:
        cumulative_cost = 0.0
        for day_offset, day_slice in enumerate(daily_slices(problem)):
            day = data.dates[problem.day_start + day_offset]
            purchase = float(np.sum(solution.grid_purchase[day_slice]))
            cost = float(
                problem.price[day_slice] @ solution.grid_purchase[day_slice]
            )
            cumulative_cost += cost
            end_energy = float(day_energy[day_offset + 1])
            overrun = float(
                max(E_DAY_MIN - end_energy, end_energy - E_DAY_MAX, 0.0)
            )
            yield (
                day.isoformat(),
                display_number(purchase),
                display_number(cost),
                display_number(float(np.sum(solution.charge[day_slice]))),
                display_number(float(np.sum(solution.discharge[day_slice]))),
                display_number(float(np.sum(solution.curtailment[day_slice]))),
                display_number(0.0),
                display_number(0.0),
                display_number(cost),
                display_number(cumulative_cost),
                display_number(float(day_energy[day_offset])),
                display_number(end_energy),
                display_number(max(end_energy - E_DAY_MIN, 0.0)),
                display_number(max(E_DAY_MAX - end_energy, 0.0)),
                display_number(overrun),
            )

    return write_csv(DAILY_FILE, header, rows())


def read_causal_check_metrics() -> dict[str, float]:
    """读取因果滚动模型的 2_模型校验.csv，用于参照对比；缺失时返回空字典。"""
    if not CAUSAL_CHECK_FILE.exists():
        return {}
    metrics: dict[str, float] = {}
    with CAUSAL_CHECK_FILE.open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.reader(csv_file)
        if next(reader, None) is None:
            return {}
        for row in reader:
            if len(row) < 3:
                continue
            try:
                metrics[row[1]] = float(row[2])
            except ValueError:
                continue
    return metrics


def summarize_storage_cost_flow(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> dict[str, float]:
    """汇总全知视角的总量、损耗与费用指标。"""
    cycle_loss = float(
        np.sum(
            (1 - ETA_CHARGE) * solution.charge
            + (1 / ETA_DISCHARGE - 1) * solution.discharge
        )
    )
    return {
        "计划购电总量": float(np.sum(solution.grid_purchase)),
        "充电总量": float(np.sum(solution.charge)),
        "放电总量": float(np.sum(solution.discharge)),
        "弃光总量": float(np.sum(solution.curtailment)),
        "未利用计划购电总量": 0.0,
        "紧急购电总量": 0.0,
        "储能循环损耗": cycle_loss,
        "总购电费": float(problem.price @ solution.grid_purchase),
        "负载总量": float(np.sum(problem.load_energy)),
        "光伏总量": float(np.sum(problem.photovoltaic_energy)),
    }


def build_check_rows(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
    metrics: CheckMetrics,
    causal_metrics: dict[str, float],
    initial_energy: float,
) -> list[Sequence[object]]:
    """汇总求解证书、约束残差、总量、费用和与因果模型的对比。"""
    totals = summarize_storage_cost_flow(problem, solution)
    total_cost = totals["总购电费"]
    period_count = problem.period_count
    causal_cost = causal_metrics.get("334天总费用")
    gap_rows: list[Sequence[object]] = []
    if causal_cost is not None and causal_cost > 0:
        gap_rows = [
            ("对比", "因果滚动模型 334 天总费用", display_number(causal_cost), "元", "参考"),
            ("对比", "全知视角相对因果模型减少费用", display_number(causal_cost - total_cost), "元", "参考"),
            (
                "对比",
                "全知视角相对因果模型下降比例",
                display_number((causal_cost - total_cost) / causal_cost),
                "比例",
                "参考",
            ),
            (
                "对比",
                "因果滚动模型紧急购电费",
                display_number(causal_metrics.get("紧急购电费", 0.0)),
                "元",
                "参考",
            ),
        ]
    else:
        gap_rows = [
            (
                "对比",
                "因果滚动模型 334 天总费用",
                "缺失",
                "元",
                f"未找到 {CAUSAL_CHECK_FILE.name}，请先运行 2_rolling_optimization.py",
            )
        ]

    rows: list[Sequence[object]] = [
        ("区间", "评价区间", f"{OFFICIAL_START} 至 {YEAR_END}", "", "参考"),
        ("区间", "正式期天数", str(problem.day_count), "天", "通过" if problem.day_count == OFFICIAL_DAYS else "失败：不得提交"),
        ("区间", "10 分钟时段数", str(period_count), "个", "通过" if period_count == OFFICIAL_DAYS * N_PERIODS else "失败：不得提交"),
        ("区间", "初始储电量", display_number(initial_energy), "kWh", "参考"),
        ("区间", "每日末态允许区间", f"[{E_DAY_MIN:g}, {E_DAY_MAX:g}]", "kWh", "参考"),
        ("信息", "预测使用的最晚实际数据日期", YEAR_END.isoformat(), "", "全知视角：与正式期最后一天相同，不满足非前视约束"),
        (
            "信息",
            "紧急购电价格倍数",
            f"{EMERGENCY_PRICE_MULTIPLIER:g}",
            "倍",
            "因果模型费用口径；全知视角不产生紧急购电，该倍数不进入本模型目标函数",
        ),
        ("信息", "紧急购电需求量", display_number(0.0), "kWh", "全知视角：完美预见下计划购电即实际购电"),
        ("求解", "全局最优证书", solution.certificate, "", "参考"),
        ("求解", "求解器状态码", str(solution.solver_status), "", "通过" if solution.is_certified else "参考：仅得到可行上界解"),
        ("求解", "求解器信息", solution.solver_message, "", "参考"),
        ("求解", "LP 松弛最优值", "无" if solution.lp_value is None else display_number(solution.lp_value), "元", "下界"),
        ("求解", "MILP 目标值", display_number(solution.objective_value), "元", "参考"),
        ("求解", "MILP 对偶界", "无" if solution.milp_dual_bound is None else display_number(solution.milp_dual_bound), "元", "下界"),
        (
            "求解",
            "MIP 相对间隙",
            f"{solution.mip_gap:.12g}" if solution.mip_gap is not None else "无",
            "比例",
            "通过" if solution.is_certified else "参考",
        ),
        (
            "求解",
            "是否已证明全时段全局最优",
            "是" if solution.is_certified else "否",
            "",
            "是：LP 互补性认证或 MILP 相对间隙达标；否：仅给出上下界",
        ),
        ("求解", "LP 松弛同时充放电最大重叠量", "无" if solution.lp_overlap is None else f"{solution.lp_overlap:.12g}", "kWh", "参考"),
        ("求解", "LP 求解耗时", "无" if solution.lp_seconds is None else f"{solution.lp_seconds:.3f}", "s", "参考"),
        ("求解", "MILP 求解耗时", "无" if solution.milp_seconds is None else f"{solution.milp_seconds:.3f}", "s", "参考"),
        ("约束", "最大电量平衡残差", f"{metrics.max_balance_residual:.12g}", "kWh", "通过"),
        ("约束", "最大储电量状态转移残差", f"{metrics.max_state_residual:.12g}", "kWh", "通过"),
        ("约束", "最大同时充放电重叠量", f"{metrics.max_simultaneous_overlap:.12g}", "kWh", "通过"),
        ("约束", "充放电模式最大整数残差", f"{metrics.max_mode_integer_error:.12g}", "", "通过"),
        ("约束", "充放电与购电边界最大越界量", f"{metrics.max_power_limit_overrun:.12g}", "kWh", "通过"),
        ("约束", "弃光上限最大越界量", f"{metrics.max_curtailment_overrun:.12g}", "kWh", "通过"),
        ("约束", "最小储电量", display_number(metrics.min_stored_energy), "kWh", "通过"),
        ("约束", "最大储电量", display_number(metrics.max_stored_energy), "kWh", "通过"),
        ("约束", "日末储电量最大值", display_number(metrics.max_terminal_energy), "kWh", "参考"),
        ("约束", "日末储电量最小值", display_number(metrics.min_terminal_energy), "kWh", "参考"),
        ("约束", "日末区间最大越界量", f"{metrics.max_terminal_overrun:.12g}", "kWh", "通过"),
        ("约束", "目标函数重算误差", f"{metrics.objective_recompute_error:.12g}", "元", "通过"),
        ("约束", "全时段能量守恒残差", f"{metrics.total_energy_residual:.12g}", "kWh", "通过"),
        ("数量", "计划购电总量", display_number(totals["计划购电总量"]), "kWh", "参考"),
        ("数量", "实际负载总量", display_number(totals["负载总量"]), "kWh", "参考"),
        ("数量", "可用光伏总量", display_number(totals["光伏总量"]), "kWh", "参考"),
        ("数量", "充电总量", display_number(totals["充电总量"]), "kWh", "参考"),
        ("数量", "放电总量", display_number(totals["放电总量"]), "kWh", "参考"),
        ("数量", "弃光总量", display_number(totals["弃光总量"]), "kWh", "参考"),
        ("数量", "储能循环损耗", display_number(totals["储能循环损耗"]), "kWh", "参考"),
        ("数量", "未利用计划购电总量", display_number(0.0), "kWh", "全知视角恒为 0"),
        ("数量", "紧急购电总量", display_number(0.0), "kWh", "全知视角恒为 0"),
        ("数量", "紧急购电 10 分钟时段数", "0", "个", "全知视角恒为 0"),
        ("费用", "总购电费", display_number(total_cost), "元", "参考"),
        ("费用", "计划购电费", display_number(total_cost), "元", "参考"),
        ("费用", "紧急购电费", display_number(0.0), "元", "全知视角恒为 0"),
        ("费用", "日均购电费", display_number(total_cost / problem.day_count), "元/天", "参考"),
        (
            "费用",
            "单位负载购电成本",
            display_number(total_cost / totals["负载总量"] if totals["负载总量"] else 0.0),
            "元/kWh",
            "参考",
        ),
        (
            "费用",
            "单位净负荷购电成本",
            display_number(
                total_cost / (totals["负载总量"] - totals["光伏总量"] + totals["弃光总量"])
                if totals["负载总量"] - totals["光伏总量"] + totals["弃光总量"] > 0
                else 0.0
            ),
            "元/kWh",
            "参考",
        ),
        *gap_rows,
        (
            "结论",
            "全知视角能否作为问题二理论下界",
            "是",
            "",
            "全知视角最优费用不高于任何满足非前视约束的因果策略实际费用",
        ),
        (
            "结论",
            "全知视角能否作为附件 5 正式结果",
            "否",
            "",
            "全知视角使用未来实际数据，违反题目非前视要求，只作论文参考",
        ),
    ]
    return rows


def write_check_csv(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
    metrics: CheckMetrics,
    causal_metrics: dict[str, float],
    initial_energy: float,
) -> Path:
    """写出全知视角的求解证书、约束校验、总量费用与对比结论。"""
    return write_csv(
        CHECK_FILE,
        ("类别", "校验项", "数值", "单位", "结论"),
        build_check_rows(problem, solution, metrics, causal_metrics, initial_energy),
    )


def specified_day_results(
    data: ModelData,
    problem: OmniscientProblem,
) -> dict[date, int]:
    """建立四个指定日期到全知视角天序号的索引。"""
    mapping: dict[date, int] = {}
    for specified_day in SPECIFIED_DATES:
        day_offset = (specified_day - OFFICIAL_START).days
        if not 0 <= day_offset < problem.day_count:
            raise RuntimeError(f"指定日期 {specified_day} 不在正式期内。")
        if data.dates[problem.day_start + day_offset] != specified_day:
            raise RuntimeError(f"指定日期 {specified_day} 与附件 2 日期不一致。")
        mapping[specified_day] = day_offset
    return mapping


def write_purchase_summary_csv(
    data: ModelData,
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> Path:
    """按表 1 口径汇总四个指定日期的计划购电量和全天结果。"""
    mapping = specified_day_results(data, problem)
    rows: list[Sequence[object]] = []
    for specified_day in SPECIFIED_DATES:
        day_offset = mapping[specified_day]
        day_slice = slice(day_offset * N_PERIODS, (day_offset + 1) * N_PERIODS)
        for hour in SPECIFIED_HOURS:
            period_index = hour * 6
            rows.append(
                (
                    specified_day.isoformat(),
                    "指定时段计划购电量",
                    data.interval_labels[period_index],
                    display_number(solution.grid_purchase[day_slice][period_index]),
                    "kWh",
                )
            )
        rows.extend(
            (
                (
                    specified_day.isoformat(),
                    "全天的计划购电量",
                    "00:00-24:00",
                    display_number(float(np.sum(solution.grid_purchase[day_slice]))),
                    "kWh",
                ),
                (
                    specified_day.isoformat(),
                    "全天购电费",
                    "00:00-24:00",
                    display_number(
                        float(
                            problem.price[day_slice]
                            @ solution.grid_purchase[day_slice]
                        )
                    ),
                    "元",
                ),
            )
        )
    return write_csv(
        PURCHASE_SUMMARY_FILE,
        ("日期", "指标", "时间段", "数值", "单位"),
        rows,
    )


def write_storage_summary_csv(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> Path:
    """按表 2 口径汇总四个指定日期的充放电量和 0:00、24:00 储电量。"""
    rows: list[Sequence[object]] = []
    periods_per_block = 4 * 6
    day_energy = solution.stored_energy[::N_PERIODS]
    for specified_day in SPECIFIED_DATES:
        day_offset = (specified_day - OFFICIAL_START).days
        if not 0 <= day_offset < problem.day_count:
            raise RuntimeError(f"指定日期 {specified_day} 不在正式期内。")
        day_slice = slice(day_offset * N_PERIODS, (day_offset + 1) * N_PERIODS)
        day_charge = solution.charge[day_slice]
        day_discharge = solution.discharge[day_slice]
        for block_index, block_label in enumerate(STORAGE_BLOCK_LABELS):
            start = block_index * periods_per_block
            stop = (block_index + 1) * periods_per_block
            rows.append(
                (
                    specified_day.isoformat(),
                    block_label,
                    display_number(float(np.sum(day_charge[start:stop]))),
                    display_number(float(np.sum(day_discharge[start:stop]))),
                    "",
                    "",
                )
            )
        rows.extend(
            (
                (
                    specified_day.isoformat(),
                    "",
                    "",
                    "",
                    "00:00",
                    display_number(float(day_energy[day_offset])),
                ),
                (
                    specified_day.isoformat(),
                    "",
                    "",
                    "",
                    "24:00",
                    display_number(float(day_energy[day_offset + 1])),
                ),
            )
        )
    return write_csv(
        STORAGE_SUMMARY_FILE,
        (
            "日期",
            "时间段",
            "充电量（kWh）",
            "放电量（kWh）",
            "时刻",
            "储电量（kWh）",
        ),
        rows,
    )


def write_emergency_summary_csv() -> Path:
    """按表 3 口径说明全知视角不存在紧急购电区间。"""
    rows = [
        (specified_day.isoformat(), "无", display_number(0.0))
        for specified_day in SPECIFIED_DATES
    ]
    return write_csv(
        EMERGENCY_SUMMARY_FILE,
        ("日期", "紧急购电时间段", "紧急购电量（kWh）"),
        rows,
    )


def write_comparison_csv(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
    causal_metrics: dict[str, float],
) -> Path:
    """写出全知视角与因果滚动模型的费用、购电和弃光对比。"""
    totals = summarize_storage_cost_flow(problem, solution)
    items = (
        ("计划购电费", totals["总购电费"], "计划购电费", "元"),
        ("紧急购电费", 0.0, "紧急购电费", "元"),
        ("334 天总费用", totals["总购电费"], "334天总费用", "元"),
        ("计划购电总量", totals["计划购电总量"], "计划购电总量", "kWh"),
        ("紧急购电总量", 0.0, "紧急购电总量", "kWh"),
        ("弃光总量", totals["弃光总量"], "弃光总量", "kWh"),
        ("紧急购电 10 分钟时段数", 0.0, "紧急购电 10 分钟时段数", "个"),
    )

    def format_value(value: float) -> str:
        return display_number(value)

    rows: list[Sequence[object]] = []
    for name, omniscient_value, causal_key, unit in items:
        causal_value = causal_metrics.get(causal_key)
        if causal_value is None:
            rows.append(
                (
                    name,
                    format_value(omniscient_value),
                    "缺失",
                    "缺失",
                    "缺失",
                    unit,
                    f"未在 {CAUSAL_CHECK_FILE.name} 中找到“{causal_key}”",
                )
            )
            continue
        difference = causal_value - omniscient_value
        relative = difference / abs(causal_value) if causal_value else 0.0
        rows.append(
            (
                name,
                format_value(omniscient_value),
                format_value(causal_value),
                format_value(difference),
                f"{100.0 * relative:.{DISPLAY_DECIMALS}f}",
                unit,
                "差额为正表示因果滚动模型高于全知视角",
            )
        )
    return write_csv(
        COMPARISON_FILE,
        (
            "指标",
            "全知视角（完美预见）",
            "因果滚动模型",
            "差额（因果−全知）",
            "相对差（%）",
            "单位",
            "备注",
        ),
        rows,
    )


def capture_row_styles(
    worksheet: object,
    row_numbers: Sequence[int],
    maximum_column: int,
) -> tuple[tuple[object, ...], tuple[float | None, ...]]:
    """在扩展模板工作表前保存示例行的单元格样式和行高。"""
    styles = tuple(
        tuple(
            copy(worksheet.cell(row=row_number, column=column_index)._style)
            for column_index in range(1, maximum_column + 1)
        )
        for row_number in row_numbers
    )
    heights = tuple(
        worksheet.row_dimensions[row_number].height for row_number in row_numbers
    )
    return styles, heights


def apply_row_style(
    worksheet: object,
    row_number: int,
    styles: Sequence[object],
    row_height: float | None,
) -> None:
    """把模板示例行样式复制到新结果行。"""
    for column_index, style in enumerate(styles, start=1):
        worksheet.cell(row=row_number, column=column_index)._style = copy(style)
    worksheet.row_dimensions[row_number].height = row_height


def validate_result_template(workbook: object) -> None:
    """在写入前核对附件 5 问题二模板的工作表、表头、日期和时段划分。"""
    missing_sheets = set(RESULT_SHEET_NAMES).difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(
            "result2.xlsx 模板缺少工作表：" + "、".join(sorted(missing_sheets))
        )

    purchase_sheet = workbook["计划购电量"]
    if purchase_sheet.cell(row=1, column=1).value != "日期\\时间":
        raise ValueError("计划购电量工作表 A1 应为“日期\\时间”。")
    interval_headers = [
        purchase_sheet.cell(row=1, column=column_index).value
        for column_index in range(2, N_PERIODS + 2)
    ]
    if any(not isinstance(header, str) or not header.strip() for header in interval_headers):
        raise ValueError("计划购电量工作表必须保留 144 个非空时间段表头。")
    if len(set(interval_headers)) != N_PERIODS:
        raise ValueError("计划购电量工作表的 144 个时间段表头必须互不重复。")

    def minutes(label: str) -> int:
        clock_label, _, day_offset = label.partition("+")
        hour, minute = map(int, clock_label.split(":"))
        if not (0 <= hour <= 24 and 0 <= minute < 60):
            raise ValueError(f"模板时间标签无效：{label}")
        return hour * 60 + minute + 1440 * int(day_offset or 0)

    # 模板表头整体从 0:10 开始，比内部区间起点晚 10 分钟；按序号映射即可，
    # 不能按标签再移动一次数据。
    for index, header in enumerate(interval_headers):
        start_label, stop_label = header.split("-")
        start, stop = minutes(start_label), minutes(stop_label)
        if index == N_PERIODS - 1 and start == 0:
            start = 1440
        if (start, stop) != ((index + 1) * 10, (index + 2) * 10):
            raise ValueError(f"计划购电模板第 {index + 2} 列时间错位：{header}")
    if purchase_sheet.cell(row=1, column=N_PERIODS + 2).value != "全天购电量":
        raise ValueError("计划购电量工作表倒数第二列应为“全天购电量”。")
    if purchase_sheet.cell(row=1, column=N_PERIODS + 3).value != "全天购电费":
        raise ValueError("计划购电量工作表最后一列应为“全天购电费”。")

    template_dates = tuple(
        normalize_excel_date(
            purchase_sheet.cell(row=row_index, column=1).value,
            f"计划购电量工作表 A{row_index}",
        )
        for row_index in range(2, OFFICIAL_DAYS + 2)
    )
    expected_dates = tuple(
        date.fromordinal(OFFICIAL_START.toordinal() + offset)
        for offset in range(OFFICIAL_DAYS)
    )
    if template_dates != expected_dates:
        raise ValueError("计划购电量工作表日期必须从 2025-02-01 连续到 2025-12-31。")

    storage_sheet = workbook["充放电量"]
    storage_headers = tuple(
        storage_sheet.cell(row=1, column=column_index).value
        for column_index in range(1, 7)
    )
    if storage_headers != ("日期", "时间段", "充电量", "放电量", "时刻", "储电量"):
        raise ValueError("充放电量工作表表头与模板约定不一致。")
    template_blocks = tuple(
        storage_sheet.cell(row=row_index, column=2).value
        for row_index in range(2, 8)
    )
    if template_blocks != STORAGE_BLOCK_LABELS:
        raise ValueError("充放电量工作表的六个 4 小时时间段与题目不一致。")

    emergency_headers = tuple(
        workbook["紧急购电量"].cell(row=1, column=column_index).value
        for column_index in range(1, 4)
    )
    if emergency_headers != ("日期", "购电时间段", "购电量"):
        raise ValueError("紧急购电量工作表表头与模板约定不一致。")


def fill_purchase_sheet(
    worksheet: object,
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> None:
    """填写 334 天 × 144 个计划购电量、全天购电量和全天购电费。"""
    for day_offset, day_slice in enumerate(daily_slices(problem)):
        row_index = day_offset + 2
        for period_index, value in enumerate(solution.grid_purchase[day_slice]):
            cell = worksheet.cell(row=row_index, column=period_index + 2)
            cell.value = excel_number(value)
            cell.number_format = "0.000000"
        purchase_cell = worksheet.cell(row=row_index, column=N_PERIODS + 2)
        cost_cell = worksheet.cell(row=row_index, column=N_PERIODS + 3)
        purchase_cell.value = excel_number(
            float(np.sum(solution.grid_purchase[day_slice]))
        )
        cost_cell.value = excel_number(
            float(problem.price[day_slice] @ solution.grid_purchase[day_slice])
        )
        purchase_cell.number_format = "0.000000"
        cost_cell.number_format = "0.000000"


def fill_storage_sheet(
    worksheet: object,
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> None:
    """扩展模板并填写每天六个 4 小时时段的充放电量和边界储电量。"""
    prototype_styles, prototype_heights = capture_row_styles(
        worksheet,
        row_numbers=tuple(range(2, 8)),
        maximum_column=6,
    )
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    periods_per_block = 4 * 6
    day_energy = solution.stored_energy[::N_PERIODS]
    destination_row = 2
    for day_offset, day_slice in enumerate(daily_slices(problem)):
        day = date.fromordinal(OFFICIAL_START.toordinal() + day_offset)
        for block_index, block_label in enumerate(STORAGE_BLOCK_LABELS):
            apply_row_style(
                worksheet,
                destination_row,
                prototype_styles[block_index],
                prototype_heights[block_index],
            )
            if block_index == 0:
                worksheet.cell(row=destination_row, column=1).value = day
                worksheet.cell(row=destination_row, column=1).number_format = "yyyy/m/d"
            worksheet.cell(row=destination_row, column=2).value = block_label
            start = block_index * periods_per_block
            stop = (block_index + 1) * periods_per_block
            charge_cell = worksheet.cell(row=destination_row, column=3)
            discharge_cell = worksheet.cell(row=destination_row, column=4)
            charge_cell.value = excel_number(
                float(np.sum(solution.charge[day_slice][start:stop]))
            )
            discharge_cell.value = excel_number(
                float(np.sum(solution.discharge[day_slice][start:stop]))
            )
            charge_cell.number_format = "0.000000"
            discharge_cell.number_format = "0.000000"
            if block_index == 0:
                worksheet.cell(row=destination_row, column=5).value = time(0, 0)
                worksheet.cell(row=destination_row, column=5).number_format = "h:mm"
                storage_cell = worksheet.cell(row=destination_row, column=6)
                storage_cell.value = excel_number(float(day_energy[day_offset]))
                storage_cell.number_format = "0.000000"
            elif block_index == 1:
                worksheet.cell(row=destination_row, column=5).value = "24:00"
                storage_cell = worksheet.cell(row=destination_row, column=6)
                storage_cell.value = excel_number(float(day_energy[day_offset + 1]))
                storage_cell.number_format = "0.000000"
            destination_row += 1


def fill_emergency_sheet(worksheet: object) -> None:
    """写入一行说明：全知视角不存在紧急购电。"""
    prototype_styles, prototype_heights = capture_row_styles(
        worksheet,
        row_numbers=(2,),
        maximum_column=3,
    )
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    apply_row_style(worksheet, 2, prototype_styles[0], prototype_heights[0])
    worksheet.cell(row=2, column=1).value = OFFICIAL_START
    worksheet.cell(row=2, column=1).number_format = "yyyy/m/d"
    worksheet.cell(row=2, column=2).value = "无（完美预见，全期无紧急购电）"
    purchase_cell = worksheet.cell(row=2, column=3)
    purchase_cell.value = excel_number(0.0)
    purchase_cell.number_format = "0.000000"


def write_result_workbook(
    problem: OmniscientProblem,
    solution: OmniscientSolution,
) -> Path:
    """从附件 5 模板生成 result2_全知视角.xlsx，回读校验后原子替换。"""
    if not RESULT_TEMPLATE_FILE.exists():
        raise FileNotFoundError(f"未找到问题二结果模板：{RESULT_TEMPLATE_FILE}")
    OUTPUT_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(RESULT_TEMPLATE_FILE)
    temporary = None
    try:
        validate_result_template(workbook)
        fill_purchase_sheet(workbook["计划购电量"], problem, solution)
        fill_storage_sheet(workbook["充放电量"], problem, solution)
        fill_emergency_sheet(workbook["紧急购电量"])
        with tempfile.NamedTemporaryFile(
            dir=OUTPUT_RESULT_DIR, suffix=".xlsx", delete=False
        ) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        written = load_workbook(temporary, read_only=True, data_only=True)
        try:
            purchase = written["计划购电量"]
            if purchase.max_row != OFFICIAL_DAYS + 1:
                raise RuntimeError("写出后的计划购电表行数不正确。")
            for row_index, values in enumerate(
                purchase.iter_rows(min_row=2, values_only=True)
            ):
                day_slice = slice(row_index * N_PERIODS, (row_index + 1) * N_PERIODS)
                expected = [
                    excel_number(value) for value in solution.grid_purchase[day_slice]
                ]
                expected += [
                    excel_number(float(np.sum(solution.grid_purchase[day_slice]))),
                    excel_number(
                        float(problem.price[day_slice] @ solution.grid_purchase[day_slice])
                    ),
                ]
                if not np.allclose(
                    np.asarray(values[1:], dtype=float), expected, atol=1e-9, rtol=0.0
                ):
                    raise RuntimeError("写出后的 Excel 购电数据回读校验失败。")
            if written["充放电量"].max_row != 1 + OFFICIAL_DAYS * 6:
                raise RuntimeError("写出后的储能表行数不正确。")
            if written["紧急购电量"].max_row != 2:
                raise RuntimeError("写出后的紧急购电表行数不正确。")
        finally:
            written.close()
        os.replace(temporary, RESULT_FILE)
    finally:
        workbook.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return RESULT_FILE


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    """解析全知视角脚本的命令行参数。"""
    parser = argparse.ArgumentParser(
        description="问题二全知视角（完美预见）全局最优参考模型",
    )
    parser.add_argument(
        "--initial-energy",
        type=float,
        default=E_INITIAL,
        help=(
            "2025-02-01 0:00 的初始储电量（kWh），默认 6000；"
            "与因果滚动模型比较时应取因果模型当天的实际储电量"
        ),
    )
    parser.add_argument(
        "--mip-rel-gap",
        type=float,
        default=DEFAULT_MIP_REL_GAP,
        help=f"转入 MILP 时的相对最优性间隙，默认 {DEFAULT_MIP_REL_GAP:g}",
    )
    parser.add_argument(
        "--time-limit",
        type=float,
        default=DEFAULT_TIME_LIMIT_SECONDS,
        help=f"单次求解的时间上限（秒），默认 {DEFAULT_TIME_LIMIT_SECONDS:g}；0 表示不限制",
    )
    parser.add_argument(
        "--no-lp-first",
        action="store_true",
        help="跳过 LP 松弛互补性认证，直接求解全时段 MILP",
    )
    parser.add_argument(
        "--skip-workbook",
        action="store_true",
        help="不生成 outputs/results/result2_全知视角.xlsx",
    )
    parser.add_argument(
        "--solver-log",
        action="store_true",
        help="打开 HiGHS 求解日志；大模型求解时建议打开以观察进度",
    )
    args = parser.parse_args(argv)
    if not 0 < args.mip_rel_gap <= 1:
        raise ValueError("--mip-rel-gap 必须位于 (0, 1]。")
    if args.time_limit < 0:
        raise ValueError("--time-limit 不能为负。")
    validate_initial_energy(args.initial_energy)
    return args


def main(argv: Sequence[str] | None = None) -> None:
    """读取数据、构造全时段 MILP、求解、校验并输出参考结果。"""
    if Path(sys.prefix).name != "2026C" or not (Path(sys.prefix) / "conda-meta").is_dir():
        raise RuntimeError("项目要求使用 conda 环境 2026C。")
    args = parse_args(argv)
    data = load_model_data()
    problem = build_problem(data, args.initial_energy)
    print(
        f"全知视角：{problem.day_count} 天、{problem.period_count} 个 10 分钟时段、"
        f"{len(problem.objective)} 个变量（其中 {problem.period_count} 个二元变量），"
        f"{problem.constraint_matrix.shape[0]} 条约束。",
        flush=True,
    )
    template = load_workbook(RESULT_TEMPLATE_FILE)
    try:
        validate_result_template(template)
    finally:
        template.close()

    solution = solve_omniscient(
        problem,
        lp_first=not args.no_lp_first,
        mip_rel_gap=args.mip_rel_gap,
        time_limit=args.time_limit,
        solver_log=args.solver_log,
    )
    metrics = validate_solution(problem, solution)
    causal_metrics = read_causal_check_metrics()

    OUTPUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    output_files = [
        write_detail_csv(data, problem, solution),
        write_daily_csv(data, problem, solution),
        write_check_csv(problem, solution, metrics, causal_metrics, args.initial_energy),
        write_purchase_summary_csv(data, problem, solution),
        write_storage_summary_csv(problem, solution),
        write_emergency_summary_csv(),
        write_comparison_csv(problem, solution, causal_metrics),
    ]
    if not args.skip_workbook:
        output_files.append(write_result_workbook(problem, solution))

    totals = summarize_storage_cost_flow(problem, solution)
    print(f"全局最优证书：{solution.certificate}")
    if not solution.is_certified:
        print("注意：本次运行未证明全局最优，输出的是带上下界的可行参考解。")
    print(f"总购电费：{solution.objective_value:.6f} 元")
    print(f"计划购电总量：{totals['计划购电总量']:.6f} kWh")
    print(f"弃光总量：{totals['弃光总量']:.6f} kWh")
    print(f"日末储电量范围：{metrics.min_terminal_energy:.6f} - {metrics.max_terminal_energy:.6f} kWh")
    for output_file in output_files:
        print(f"已保存：{output_file.relative_to(PROJECT_DIR)}")


if __name__ == "__main__":
    main()
