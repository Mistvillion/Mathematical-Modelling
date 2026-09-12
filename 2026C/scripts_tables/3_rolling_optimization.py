"""问题三：净负荷情景下的购电与储能滚动调度。

在 conda 环境 2026C 中运行：
    python scripts_tables/3_rolling_optimization.py

默认完整执行 docs/Q3.md 的四次购电决策主方案及五个信息更新对照，生成
outputs/tables/3_*.csv 和 outputs/results/result3.xlsx。快速检查可运行：
    python scripts_tables/3_rolling_optimization.py --smoke-days 2

实现口径：直接从历史实际净负荷轨迹构造条件情景；购电阶段先做连续模式松弛，
再固定模式求解 LP，并在异常时回退 MILP；目标中计入情景期望紧急购电费。
储能阶段每10分钟更新情景权重，只规划至下一个6小时边界，近端保持10分钟、
远端按30分钟聚合，并执行非前视的共同当前动作。模型参数只按已经结束日期的实际总
费用滚动选择，MAE与覆盖率仅作诊断。调整费用相对0点计划最终结算一次。
"""

from __future__ import annotations

import argparse
import csv
import importlib
import os
import sys
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time
from math import isfinite
from pathlib import Path
from types import SimpleNamespace
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import coo_matrix, vstack


PROJECT_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

# 问题二已经集中实现并验证了附件1/2读取、实际储能滚动LP及模板样式工具。
# 通过 importlib 导入以避免复制两千余行相同的物理模型代码。
q2 = importlib.import_module("2_rolling_optimization")

ATTACHMENT_DIR = PROJECT_DIR / "CUMCM 2026 C题" / "附件"
PHOTOVOLTAIC_FORECAST_FILE = ATTACHMENT_DIR / "附件3.xlsx"
RESULT_TEMPLATE_FILE = ATTACHMENT_DIR / "附件5" / "result3.xlsx"
OUTPUT_TABLE_DIR = PROJECT_DIR / "outputs" / "tables"
OUTPUT_RESULT_DIR = PROJECT_DIR / "outputs" / "results"
RESULT_FILE = OUTPUT_RESULT_DIR / "result3.xlsx"

DECISION_FILE = OUTPUT_TABLE_DIR / "3_预报与购电决策版本.csv"
DETAIL_FILE = OUTPUT_TABLE_DIR / "3_实际调度与费用明细.csv"
CHECK_FILE = OUTPUT_TABLE_DIR / "3_模型校验.csv"
MAPPING_FILE = OUTPUT_TABLE_DIR / "3_输出时段映射.csv"
COMPARISON_FILE = OUTPUT_TABLE_DIR / "3_预报更新策略对照.csv"
PURCHASE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "3_指定日期购电量及全天结果.csv"
STORAGE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "3_指定日期充放电量及储电量.csv"
EMERGENCY_SUMMARY_FILE = OUTPUT_TABLE_DIR / "3_指定日期紧急购电量.csv"

N_PERIODS = q2.N_PERIODS
DELTA_T = q2.DELTA_T
YEAR_START = q2.YEAR_START
OFFICIAL_START = q2.OFFICIAL_START
YEAR_END = q2.YEAR_END
JANUARY_DAYS = q2.JANUARY_DAYS
OFFICIAL_DAYS = q2.OFFICIAL_DAYS

E_MIN = q2.E_MIN
E_MAX = q2.E_MAX
E_INITIAL = q2.E_INITIAL
E_REFERENCE = q2.E_REFERENCE
E_TERMINAL_MIN = q2.E_TERMINAL_MIN
E_TERMINAL_MAX = q2.E_TERMINAL_MAX
ETA_CHARGE = q2.ETA_CHARGE
ETA_DISCHARGE = q2.ETA_DISCHARGE
CHARGE_LIMIT = q2.CHARGE_LIMIT
DISCHARGE_LIMIT = q2.DISCHARGE_LIMIT
EMERGENCY_PRICE_MULTIPLIER = q2.EMERGENCY_PRICE_MULTIPLIER

DIAGNOSTIC_QUANTILE = 0.8
TERMINAL_PENALTY_MULTIPLIER = 1.0
MIP_REL_GAP = q2.MIP_REL_GAP
MIN_TUNING_DAYS = 7
DEFAULT_VALIDATION_WINDOW = 28

BALANCE_TOLERANCE = q2.BALANCE_TOLERANCE
BOUND_TOLERANCE = q2.BOUND_TOLERANCE
INTEGER_TOLERANCE = q2.INTEGER_TOLERANCE
TERMINAL_TOLERANCE = q2.TERMINAL_TOLERANCE
DISPLAY_DECIMALS = q2.DISPLAY_DECIMALS
DISPLAY_ZERO_TOLERANCE = q2.DISPLAY_ZERO_TOLERANCE

RELEASE_HOURS = (0, 6, 12, 18)
RELEASE_PERIODS = (0, 36, 72, 108)
PERIODS_PER_BLOCK = 36
REALTIME_FINE_PERIODS = 6
REALTIME_AGGREGATION_PERIODS = 3
PHOTOVOLTAIC_FORECAST_COLUMNS = tuple(
    ["日期", "预报时刻"] + [f"预报{hour}小时" for hour in range(1, 25)]
)
RESULT_SHEET_NAMES = ("计划购电量", "调整购电量", "充放电量", "紧急购电量")
STORAGE_BLOCK_LABELS = q2.STORAGE_BLOCK_LABELS
SPECIFIED_DATES = q2.SPECIFIED_DATES
SPECIFIED_HOURS = q2.SPECIFIED_HOURS


@dataclass(frozen=True)
class ModelData:
    """附件1、2、3完成校验并对齐后的模型输入。"""

    base: object
    photovoltaic_forecast_power: np.ndarray
    photovoltaic_feature_energy: np.ndarray
    actual_net_load_energy: np.ndarray


@dataclass(frozen=True)
class ScenarioParameters:
    """由历史实际总费用选择的一组净负荷情景参数。"""

    key: str
    decay_days: float
    similarity_scale: float
    scenario_count: int
    sampling_rule: str
    validation_window: int = DEFAULT_VALIDATION_WINDOW


@dataclass(frozen=True)
class NetLoadScenarios:
    """一次边界决策可用的整轨迹净负荷情景及其来源。"""

    decision_index: int
    source_batch_index: int
    net_load_energy: np.ndarray
    weights: np.ndarray
    origin_dates: tuple[date | None, ...]
    diagnostic_quantile_energy: np.ndarray
    photovoltaic_feature_energy: np.ndarray
    photovoltaic_start_power: float
    photovoltaic_start_source: str
    latest_training_date: date | None
    used_prior: bool
    parameters: ScenarioParameters


@dataclass(frozen=True)
class PurchaseDecision:
    """共享购电、情景补救变量和显式紧急购电风险的候选解。"""

    decision_index: int
    start_period: int
    grid_purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    emergency_purchase: np.ndarray
    residual_energy: np.ndarray
    mode: np.ndarray
    stored_energy: np.ndarray
    absolute_adjustment: np.ndarray
    terminal_deviation: np.ndarray
    objective_value: float
    expected_emergency_cost: float
    expected_terminal_penalty: float
    throughput_regularization: float
    solver_message: str
    solver_kind: str
    mip_gap: float


@dataclass(frozen=True)
class RealtimeStep:
    """截短多时间尺度实时模型在一个10分钟时段的动作和诊断信息。"""

    period_index: int
    horizon: int
    boundary_period: int
    fine_block_count: int
    aggregated_block_count: int
    charge: float
    discharge: float
    emergency_purchase: float
    curtailment: float
    unused_planned_purchase: float
    predicted_boundary_energy: float
    charge_capacity: float
    discharge_capacity: float
    solver_message: str


@dataclass(frozen=True)
class ReplayResult:
    """一天逐10分钟实际执行的储能与紧急购电结果。"""

    charge: np.ndarray
    discharge: np.ndarray
    emergency_purchase: np.ndarray
    curtailment: np.ndarray
    unused_planned_purchase: np.ndarray
    stored_energy: np.ndarray
    steps: tuple[RealtimeStep, ...]


@dataclass(frozen=True)
class DailyResult:
    """一个自然日的四次预测版本、购电版本、执行结果与账单。"""


    day: date
    parameters: ScenarioParameters
    parameter_latest_validation_date: date | None
    parameter_validation_days: int
    parameter_validation_cost: float | None
    scenarios: tuple[NetLoadScenarios, ...]
    decisions: tuple[PurchaseDecision | None, ...]
    purchase_versions: np.ndarray
    effective_purchase: np.ndarray
    actual_load_energy: np.ndarray
    actual_photovoltaic_energy: np.ndarray
    replay: ReplayResult
    initial_plan_cost: float
    increase_purchase_cost: float
    decrease_refund: float
    decrease_penalty: float
    emergency_purchase_cost: float
    realized_total_cost: float


@dataclass(frozen=True)
class StrategySpec:
    """对照策略允许接受的光伏批次及允许更新购电的边界。"""

    key: str
    name: str
    forecast_updates: tuple[int, ...]
    purchase_updates: tuple[int, ...]


@dataclass(frozen=True)
class TuningRecord:
    """某候选参数在一个已结束验证日上的完整实际账单。"""

    validation_day: date
    parameters: ScenarioParameters
    realized_total_cost: float


@dataclass(frozen=True)
class StrategyResult:
    """一种信息更新策略的连续因果全年运行结果。"""

    spec: StrategySpec
    warmup_days: tuple[DailyResult, ...]
    official_days: tuple[DailyResult, ...]
    tuning_replay_count: int
    tuning_records: tuple[TuningRecord, ...]

    @property
    def all_days(self) -> tuple[DailyResult, ...]:
        return self.warmup_days + self.official_days


MAIN_STRATEGY = StrategySpec("main", "四次决策主方案", (1, 2, 3), (1, 2, 3))
COMPARISON_STRATEGIES = (
    StrategySpec("zero_only", "仅0点预报", (), ()),
    StrategySpec("forecast_only", "更新预报、购电冻结", (1, 2, 3), ()),
    StrategySpec("single_6", "仅6点日内更新", (1,), (1,)),
    StrategySpec("single_12", "仅12点日内更新", (2,), (2,)),
    StrategySpec("single_18", "仅18点日内更新", (3,), (3,)),
)

PARAMETER_CANDIDATES = (
    ScenarioParameters("短窗集中", 14.0, 0.75, 3, "top_weighted"),
    ScenarioParameters("均衡重采样", 30.0, 1.00, 5, "systematic"),
    ScenarioParameters("长窗宽核", 60.0, 1.50, 7, "systematic"),
)
DEFAULT_PARAMETERS = PARAMETER_CANDIDATES[1]


def validate_nonnegative_array(values: np.ndarray, name: str, shape: tuple[int, ...]) -> None:
    """统一检查模型输入数组的形状、有限性和非负性。"""
    if values.shape != shape or not np.isfinite(values).all() or np.any(values < 0):
        raise ValueError(f"{name}必须是形状{shape}的有限非负数组。")


def validate_finite_array(values: np.ndarray, name: str, shape: tuple[int, ...]) -> None:
    """净负荷允许为负，但形状和有限性仍必须满足要求。"""
    if values.shape != shape or not np.isfinite(values).all():
        raise ValueError(f"{name}必须是形状{shape}的有限数组。")


def parse_forecast_date(value: object, location: str) -> date:
    """解析附件3中仅每天首行填写的文本日期。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                pass
    raise ValueError(f"{location}不是有效日期：{value!r}")


def parse_release_hour(value: object, location: str) -> int:
    """把附件3的预报时刻转换为0、6、12、18。"""
    if isinstance(value, time):
        if value.minute or value.second or value.microsecond:
            raise ValueError(f"{location}必须是整点：{value!r}")
        hour = value.hour
    elif isinstance(value, str):
        pieces = value.strip().split(":")
        if len(pieces) != 2 or pieces[1] != "00":
            raise ValueError(f"{location}格式应为H:00：{value!r}")
        hour = int(pieces[0])
    else:
        raise ValueError(f"{location}不是有效时刻：{value!r}")
    if hour not in RELEASE_HOURS:
        raise ValueError(f"{location}应为0:00、6:00、12:00或18:00。")
    return hour


def read_photovoltaic_forecasts() -> np.ndarray:
    """校验并读取附件3的365天、每天4批、每批24个功率节点。"""
    workbook = load_workbook(PHOTOVOLTAIC_FORECAST_FILE, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("附件3应且仅应包含工作表Sheet1。")
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) != 365 * 4 + 1:
        raise ValueError(f"附件3应包含1460批预报，实际为{len(rows) - 1}批。")
    if tuple(rows[0]) != PHOTOVOLTAIC_FORECAST_COLUMNS:
        raise ValueError("附件3表头与‘日期、预报时刻、预报1至24小时’约定不一致。")

    expected_dates = tuple(date.fromordinal(YEAR_START.toordinal() + i) for i in range(365))
    forecast = np.empty((365, 4, 24), dtype=float)
    current_day: date | None = None
    for row_offset, row in enumerate(rows[1:]):
        excel_row = row_offset + 2
        day_index, release_index = divmod(row_offset, 4)
        if len(row) != 26:
            raise ValueError(f"附件3第{excel_row}行列数不为26。")
        if row[0] not in (None, ""):
            current_day = parse_forecast_date(row[0], f"附件3A{excel_row}")
        if current_day != expected_dates[day_index]:
            raise ValueError(f"附件3第{excel_row}行日期分组或顺序错误。")
        hour = parse_release_hour(row[1], f"附件3B{excel_row}")
        if hour != RELEASE_HOURS[release_index]:
            raise ValueError(f"附件3第{excel_row}行预报时刻顺序错误。")
        for node_index, value in enumerate(row[2:]):
            forecast[day_index, release_index, node_index] = q2.validate_nonnegative_number(
                value, f"附件3第{excel_row}行第{node_index + 3}列"
            )
    validate_nonnegative_array(forecast, "附件3光伏预报", (365, 4, 24))
    return forecast


def load_model_data() -> ModelData:
    """加载附件1至3，清洗实际光伏并预先形成四批因果预报特征。"""
    base = q2.load_model_data()
    if tuple(base.dates) != tuple(
        date.fromordinal(YEAR_START.toordinal() + i) for i in range(365)
    ):
        raise ValueError("附件2日期不构成2025年连续365天。")
    forecast_power = read_photovoltaic_forecasts()
    provisional = ModelData(
        base=base,
        photovoltaic_forecast_power=forecast_power,
        photovoltaic_feature_energy=np.empty((365, 4, N_PERIODS)),
        actual_net_load_energy=(
            np.asarray(base.actual_load_energy, dtype=float)
            - np.asarray(base.actual_photovoltaic_energy, dtype=float)
        ),
    )
    features = build_photovoltaic_features(provisional)
    return ModelData(
        base=base,
        photovoltaic_forecast_power=forecast_power,
        photovoltaic_feature_energy=features,
        actual_net_load_energy=provisional.actual_net_load_energy,
    )


def photovoltaic_start_node(
    data: ModelData, day_index: int, batch_index: int
) -> tuple[float, str]:
    """返回一批预报的补起点及严格早于发布时间的可追溯来源。"""
    if batch_index == 0:
        if day_index == 0:
            return 0.0, "预设：2025-01-01午夜0 kW"
        value = float(data.base.actual_photovoltaic_power[day_index - 1, -1])
        return value, f"{data.base.dates[day_index - 1].isoformat()} 23:50-24:00实测代理"
    cutoff = RELEASE_PERIODS[batch_index]
    value = float(data.base.actual_photovoltaic_power[day_index, cutoff - 1])
    return value, (
        f"{data.base.dates[day_index].isoformat()} "
        f"{q2.format_clock(cutoff * 10 - 10)}-{q2.format_clock(cutoff * 10)}实测代理"
    )


def integrate_photovoltaic_nodes(nodes: np.ndarray) -> np.ndarray:
    """把25个整点功率节点精确积分为发布后144个10分钟电量。"""
    validate_nonnegative_array(nodes, "光伏插值节点", (25,))
    boundary_x = np.arange(N_PERIODS + 1, dtype=float) / 6.0
    boundary_power = np.interp(boundary_x, np.arange(25, dtype=float), nodes)
    energy = DELTA_T * (boundary_power[:-1] + boundary_power[1:]) / 2.0
    # 每小时6段之和必须等于该线性段的梯形面积。
    hourly = energy.reshape(24, 6).sum(axis=1)
    expected = (nodes[:-1] + nodes[1:]) / 2.0
    if np.max(np.abs(hourly - expected)) > 1e-9:
        raise RuntimeError("光伏10分钟积分与整点线性段积分不一致。")
    return energy


def forecast_from_batch(
    data: ModelData,
    day_index: int,
    batch_index: int,
    previous: np.ndarray | None,
) -> tuple[np.ndarray, float, str]:
    """按目标绝对时刻，把指定已发布批次覆盖到当天剩余时段。"""
    if not 0 <= batch_index < 4:
        raise ValueError("光伏批次索引必须为0至3。")
    start_power, source = photovoltaic_start_node(data, day_index, batch_index)
    nodes = np.concatenate(
        ([start_power], data.photovoltaic_forecast_power[day_index, batch_index])
    )
    relative_energy = integrate_photovoltaic_nodes(nodes)
    release = RELEASE_PERIODS[batch_index]
    if previous is None:
        if batch_index != 0:
            raise ValueError("当天首个光伏预测版本必须来自0点批次。")
        full = relative_energy.copy()
    else:
        validate_nonnegative_array(previous, "上一光伏预测版本", (N_PERIODS,))
        full = previous.copy()
        full[release:] = relative_energy[: N_PERIODS - release]
    validate_nonnegative_array(full, "当日光伏电量预测", (N_PERIODS,))
    return full, start_power, source


def build_photovoltaic_features(data: ModelData) -> np.ndarray:
    """形成每个日期在四个发布边界真正可用的全天光伏预报特征版本。"""
    features = np.empty((len(data.base.dates), 4, N_PERIODS), dtype=float)
    for day_index in range(len(data.base.dates)):
        previous: np.ndarray | None = None
        for batch_index in range(4):
            current, _, _ = forecast_from_batch(
                data, day_index, batch_index, previous
            )
            features[day_index, batch_index] = current
            previous = current
    validate_nonnegative_array(
        features,
        "光伏预报特征",
        (len(data.base.dates), 4, N_PERIODS),
    )
    return features


def weighted_quantile_by_period(
    samples: np.ndarray, weights: np.ndarray, quantile: float
) -> np.ndarray:
    """逐时段计算带权分位数，仅用于净负荷预测诊断。"""
    if samples.ndim != 2 or weights.shape != (samples.shape[0],):
        raise ValueError("带权分位数的样本与权重形状不一致。")
    if not 0.0 <= quantile <= 1.0 or np.any(weights < 0) or not np.any(weights > 0):
        raise ValueError("带权分位数的分位点或权重无效。")
    result = np.empty(samples.shape[1], dtype=float)
    normalized = weights / np.sum(weights)
    for t in range(samples.shape[1]):
        order = np.argsort(samples[:, t], kind="stable")
        cumulative = np.cumsum(normalized[order])
        index = min(int(np.searchsorted(cumulative, quantile, side="left")), len(order) - 1)
        result[t] = samples[order[index], t]
    return result


def scenario_feature_distances(
    data: ModelData,
    day_index: int,
    decision_index: int,
    source_batch_index: int,
) -> tuple[np.ndarray, np.ndarray]:
    """只用当前边界已知信息，计算同日类型历史日期的标准化特征距离。"""
    cutoff = RELEASE_PERIODS[decision_index]
    day = data.base.dates[day_index]
    eligible = np.array(
        [
            j for j in range(day_index)
            if q2.is_high_load_day(data.base.dates[j]) == q2.is_high_load_day(day)
        ],
        dtype=int,
    )
    if eligible.size == 0:
        return eligible, np.empty(0)
    if source_batch_index > decision_index:
        raise RuntimeError("使用了尚未发布的光伏预报。")

    current_pv = data.photovoltaic_feature_energy[
        day_index, source_batch_index, cutoff:
    ]
    history_pv = data.photovoltaic_feature_energy[
        eligible, source_batch_index, cutoff:
    ]
    pv_scale = max(float(np.std(history_pv)), float(np.mean(np.abs(current_pv))) * 0.1, 1.0)
    pv_distance = np.sqrt(np.mean(((history_pv - current_pv) / pv_scale) ** 2, axis=1))

    if cutoff:
        recent_start = max(0, cutoff - PERIODS_PER_BLOCK)
        current_recent = data.actual_net_load_energy[
            day_index, recent_start:cutoff
        ]
        history_recent = data.actual_net_load_energy[
            eligible, recent_start:cutoff
        ]
        recent_scale = max(
            float(np.std(history_recent)),
            float(np.mean(np.abs(current_recent))) * 0.1,
            1.0,
        )
        recent_distance = np.sqrt(
            np.mean(((history_recent - current_recent) / recent_scale) ** 2, axis=1)
        )
    else:
        recent_distance = np.zeros(eligible.size)

    day_of_year = day.timetuple().tm_yday
    history_day_of_year = np.array(
        [data.base.dates[j].timetuple().tm_yday for j in eligible], dtype=float
    )
    season_distance = 2.0 * np.sin(
        np.pi * (day_of_year - history_day_of_year) / 365.0
    )
    distance = np.sqrt(
        pv_distance**2 + recent_distance**2 + season_distance**2
    )
    return eligible, distance


def select_scenario_indices(
    weights: np.ndarray, count: int, rule: str
) -> tuple[np.ndarray, np.ndarray]:
    """按候选规则确定保留的整日历史轨迹及最终概率。"""
    if weights.ndim != 1 or weights.size == 0 or np.any(weights < 0):
        raise ValueError("情景抽样权重无效。")
    probability = weights / np.sum(weights)
    count = min(count, weights.size)
    if rule == "top_weighted":
        selected = np.argsort(-probability, kind="stable")[:count]
        selected_probability = probability[selected]
        selected_probability /= np.sum(selected_probability)
        return selected, selected_probability
    if rule == "systematic":
        positions = (np.arange(count, dtype=float) + 0.5) / count
        draws = np.searchsorted(np.cumsum(probability), positions, side="left")
        selected, occurrences = np.unique(draws, return_counts=True)
        return selected, occurrences.astype(float) / count
    raise ValueError(f"未知情景重采样规则：{rule}")


def build_net_load_scenarios(
    data: ModelData,
    day_index: int,
    decision_index: int,
    source_batch_index: int,
    parameters: ScenarioParameters,
) -> NetLoadScenarios:
    """直接由历史实际净负荷整轨迹构造当前边界的条件情景。"""
    if not 0 <= source_batch_index <= decision_index < 4:
        raise ValueError("净负荷情景的决策批次或信息批次无效。")
    cutoff = RELEASE_PERIODS[decision_index]
    day = data.base.dates[day_index]
    eligible, distance = scenario_feature_distances(
        data, day_index, decision_index, source_batch_index
    )
    pv_feature = data.photovoltaic_feature_energy[
        day_index, source_batch_index
    ].copy()
    start_power, start_source = photovoltaic_start_node(
        data, day_index, source_batch_index
    )

    if eligible.size == 0:
        prior = np.asarray(data.base.cold_start_load_energy, dtype=float) - pv_feature
        paths = prior[None, :]
        probabilities = np.ones(1)
        origins: tuple[date | None, ...] = (None,)
        diagnostic = prior.copy()
        latest_training = None
        used_prior = True
    else:
        ages = np.array(
            [(day - data.base.dates[j]).days for j in eligible], dtype=float
        )
        raw_weights = np.exp(-ages / parameters.decay_days) * np.exp(
            -0.5 * (distance / parameters.similarity_scale) ** 2
        )
        if not np.isfinite(raw_weights).all() or float(np.sum(raw_weights)) <= 1e-300:
            raw_weights = np.exp(-ages / parameters.decay_days)
        selected_local, probabilities = select_scenario_indices(
            raw_weights, parameters.scenario_count, parameters.sampling_rule
        )
        selected_days = eligible[selected_local]
        paths = data.actual_net_load_energy[selected_days].copy()
        origins = tuple(data.base.dates[j] for j in selected_days)
        all_history_paths = data.actual_net_load_energy[eligible]
        diagnostic = weighted_quantile_by_period(
            all_history_paths, raw_weights, DIAGNOSTIC_QUANTILE
        )
        latest_training = max(data.base.dates[j] for j in eligible)
        used_prior = False

    validate_finite_array(paths, "净负荷情景", (len(probabilities), N_PERIODS))
    validate_finite_array(diagnostic, "净负荷诊断曲线", (N_PERIODS,))
    if not np.isclose(np.sum(probabilities), 1.0, atol=1e-12) or np.any(probabilities <= 0):
        raise RuntimeError("净负荷情景概率未正确归一化。")
    if any(origin is not None and origin >= day for origin in origins):
        raise RuntimeError("净负荷情景使用了当日或未来实际轨迹。")
    return NetLoadScenarios(
        decision_index=decision_index,
        source_batch_index=source_batch_index,
        net_load_energy=paths,
        weights=np.asarray(probabilities, dtype=float),
        origin_dates=origins,
        diagnostic_quantile_energy=diagnostic,
        photovoltaic_feature_energy=pv_feature,
        photovoltaic_start_power=start_power,
        photovoltaic_start_source=start_source,
        latest_training_date=latest_training,
        used_prior=used_prior,
        parameters=parameters,
    )


def solve_purchase_decision(
    day: date,
    decision_index: int,
    price: np.ndarray,
    scenarios: NetLoadScenarios,
    initial_energy: float,
    terminal_penalty: float,
    initial_plan: np.ndarray | None,
    hard_terminal: bool,
    force_milp: bool = False,
) -> PurchaseDecision:
    """求解共享购电承诺与情景补救变量组成的模式松弛、固定模式LP或回退MILP。"""
    start = RELEASE_PERIODS[decision_index]
    n = N_PERIODS - start
    if scenarios.decision_index != decision_index:
        raise ValueError("购电决策与净负荷情景的发布时间不一致。")
    scenario_count = len(scenarios.weights)
    price_h = np.asarray(price[start:], dtype=float)
    net_h = np.asarray(scenarios.net_load_energy[:, start:], dtype=float)
    validate_nonnegative_array(price_h, "购电优化电价", (n,))
    validate_finite_array(net_h, "购电优化净负荷情景", (scenario_count, n))
    if np.any(price_h <= 0):
        raise ValueError("购电优化要求电价严格为正。")
    if not E_MIN - BOUND_TOLERANCE <= initial_energy <= E_MAX + BOUND_TOLERANCE:
        raise ValueError("购电优化初始储电量超出设备范围。")
    if not isfinite(terminal_penalty) or terminal_penalty < 0:
        raise ValueError("日末软惩罚必须为有限非负数。")
    adjusted = decision_index > 0
    throughput_epsilon = float(1e-4 * np.min(price_h))
    if adjusted:
        if initial_plan is None:
            raise ValueError("日内调整必须提供0点初始计划作为结算基准。")
        validate_nonnegative_array(initial_plan, "0点初始计划", (N_PERIODS,))
        initial_h = np.asarray(initial_plan[start:], dtype=float)
    elif initial_plan is not None:
        raise ValueError("0点初始计划不应再传入结算基准。")

    scenario_periods = scenario_count * n
    offset_grid = 0
    offset_charge = n
    offset_discharge = offset_charge + scenario_periods
    offset_emergency = offset_discharge + scenario_periods
    offset_residual = offset_emergency + scenario_periods
    offset_mode = offset_residual + scenario_periods
    offset_energy = offset_mode + scenario_periods
    energy_variables = scenario_count * (n + 1)
    offset_terminal = offset_energy + energy_variables
    offset_absolute = offset_terminal + scenario_count
    variable_count = offset_absolute + (n if adjusted else 0)

    def scenario_slice(offset: int, scenario_index: int, width: int = n) -> slice:
        return slice(offset + scenario_index * width, offset + (scenario_index + 1) * width)

    def energy_slice(scenario_index: int) -> slice:
        width = n + 1
        return slice(offset_energy + scenario_index * width, offset_energy + (scenario_index + 1) * width)

    objective = np.zeros(variable_count)
    objective[offset_grid : offset_grid + n] = price_h
    if adjusted:
        objective[offset_absolute : offset_absolute + n] = 0.5 * price_h
    for s, probability in enumerate(scenarios.weights):
        objective[scenario_slice(offset_charge, s)] = probability * throughput_epsilon
        objective[scenario_slice(offset_discharge, s)] = probability * throughput_epsilon
        objective[scenario_slice(offset_emergency, s)] = (
            probability * EMERGENCY_PRICE_MULTIPLIER * price_h
        )
        if not hard_terminal:
            objective[offset_terminal + s] = probability * terminal_penalty

    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    grid_upper = np.maximum(np.max(net_h, axis=0), 0.0) + CHARGE_LIMIT
    upper[offset_grid : offset_grid + n] = grid_upper
    upper[offset_charge : offset_charge + scenario_periods] = CHARGE_LIMIT
    upper[offset_discharge : offset_discharge + scenario_periods] = DISCHARGE_LIMIT
    upper[offset_mode : offset_mode + scenario_periods] = 1.0
    for s in range(scenario_count):
        upper[scenario_slice(offset_emergency, s)] = np.maximum(net_h[s], 0.0)
        upper[scenario_slice(offset_residual, s)] = np.maximum(
            grid_upper - net_h[s], 0.0
        )
        e_slice = energy_slice(s)
        lower[e_slice] = E_MIN
        upper[e_slice] = E_MAX
        lower[e_slice.start] = initial_energy
        upper[e_slice.start] = initial_energy
        if hard_terminal:
            lower[e_slice.stop - 1] = E_TERMINAL_MIN
            upper[e_slice.stop - 1] = E_TERMINAL_MAX
            upper[offset_terminal + s] = 0.0

    eq_rows: list[int] = []
    eq_cols: list[int] = []
    eq_data: list[float] = []
    eq_rhs: list[float] = []
    row = 0
    for s in range(scenario_count):
        e_slice = energy_slice(s)
        for t in range(n):
            for column, value in (
                (offset_grid + t, 1.0),
                (scenario_slice(offset_charge, s).start + t, -1.0),
                (scenario_slice(offset_discharge, s).start + t, 1.0),
                (scenario_slice(offset_emergency, s).start + t, 1.0),
                (scenario_slice(offset_residual, s).start + t, -1.0),
            ):
                eq_rows.append(row); eq_cols.append(column); eq_data.append(value)
            eq_rhs.append(float(net_h[s, t])); row += 1
        for t in range(n):
            for column, value in (
                (e_slice.start + t, -1.0),
                (e_slice.start + t + 1, 1.0),
                (scenario_slice(offset_charge, s).start + t, -ETA_CHARGE),
                (scenario_slice(offset_discharge, s).start + t, 1.0 / ETA_DISCHARGE),
            ):
                eq_rows.append(row); eq_cols.append(column); eq_data.append(value)
            eq_rhs.append(0.0); row += 1
    equality_count = row
    a_eq = coo_matrix(
        (eq_data, (eq_rows, eq_cols)), shape=(equality_count, variable_count)
    ).tocsr()
    b_eq = np.asarray(eq_rhs, dtype=float)

    ub_rows: list[int] = []
    ub_cols: list[int] = []
    ub_data: list[float] = []
    ub_rhs: list[float] = []
    row = 0
    big_m = np.maximum(
        CHARGE_LIMIT + np.maximum(net_h, 0.0),
        grid_upper[None, :] + np.maximum(-net_h, 0.0),
    ) + BOUND_TOLERANCE
    for s in range(scenario_count):
        c_start = scenario_slice(offset_charge, s).start
        d_start = scenario_slice(offset_discharge, s).start
        u_start = scenario_slice(offset_mode, s).start
        for t in range(n):
            # c <= L*u；d <= L*(1-u)
            for column, value in ((c_start + t, 1.0), (u_start + t, -CHARGE_LIMIT)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(0.0); row += 1
            for column, value in ((d_start + t, 1.0), (u_start + t, DISCHARGE_LIMIT)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(DISCHARGE_LIMIT); row += 1
            # u=1时 c<=g-n；u=0时 d<=n-g，保证紧急购电不用于充电。
            for column, value in (
                (c_start + t, 1.0), (offset_grid + t, -1.0), (u_start + t, float(big_m[s, t]))
            ):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(float(big_m[s, t] - net_h[s, t])); row += 1
            for column, value in (
                (d_start + t, 1.0), (offset_grid + t, 1.0), (u_start + t, -float(big_m[s, t]))
            ):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(float(net_h[s, t])); row += 1
        if not hard_terminal:
            e_end = energy_slice(s).stop - 1
            z = offset_terminal + s
            for column, value in ((e_end, 1.0), (z, -1.0)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(E_REFERENCE); row += 1
            for column, value in ((e_end, -1.0), (z, -1.0)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(-E_REFERENCE); row += 1
    if adjusted:
        for t in range(n):
            # g - xi <= g0; -g - xi <= -g0
            for column, value in ((offset_grid + t, 1.0), (offset_absolute + t, -1.0)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(float(initial_h[t])); row += 1
            for column, value in ((offset_grid + t, -1.0), (offset_absolute + t, -1.0)):
                ub_rows.append(row); ub_cols.append(column); ub_data.append(value)
            ub_rhs.append(float(-initial_h[t])); row += 1
    a_ub = coo_matrix(
        (ub_data, (ub_rows, ub_cols)), shape=(row, variable_count)
    ).tocsr()
    b_ub = np.asarray(ub_rhs, dtype=float)

    result = None
    solver_kind = "MILP回退"
    if not force_milp:
        relaxed = linprog(
            c=objective,
            A_eq=a_eq,
            b_eq=b_eq,
            A_ub=a_ub,
            b_ub=b_ub,
            bounds=np.column_stack((lower, upper)),
            method="highs-ds",
            options={
                "presolve": True,
                "primal_feasibility_tolerance": 1e-8,
                "dual_feasibility_tolerance": 1e-8,
            },
        )
        if relaxed.success and relaxed.x is not None:
            relaxed_solution = np.asarray(relaxed.x, dtype=float)
            grid = relaxed_solution[offset_grid : offset_grid + n]
            inferred_mode = np.empty((scenario_count, n))
            for s in range(scenario_count):
                inferred_mode[s] = (grid - net_h[s] >= 0.0).astype(float)
            fixed_lower = lower.copy()
            fixed_upper = upper.copy()
            fixed_lower[offset_mode : offset_mode + scenario_periods] = inferred_mode.ravel()
            fixed_upper[offset_mode : offset_mode + scenario_periods] = inferred_mode.ravel()
            fixed = linprog(
                c=objective,
                A_eq=a_eq,
                b_eq=b_eq,
                A_ub=a_ub,
                b_ub=b_ub,
                bounds=np.column_stack((fixed_lower, fixed_upper)),
                method="highs-ds",
                options={
                    "presolve": True,
                    "primal_feasibility_tolerance": 1e-8,
                    "dual_feasibility_tolerance": 1e-8,
                },
            )
            if fixed.success and fixed.x is not None:
                fixed_solution = np.asarray(fixed.x, dtype=float)
                if (
                    np.isfinite(fixed_solution).all()
                    and np.max(np.abs(a_eq @ fixed_solution - b_eq)) <= BALANCE_TOLERANCE
                    and np.max(a_ub @ fixed_solution - b_ub) <= BALANCE_TOLERANCE
                    and np.all(fixed_solution >= fixed_lower - BOUND_TOLERANCE)
                    and np.all(fixed_solution <= fixed_upper + BOUND_TOLERANCE)
                ):
                    result = fixed
                    solver_kind = "松弛定模式后固定模式LP"
    if result is None:
        integrality = np.zeros(variable_count, dtype=int)
        integrality[offset_mode : offset_mode + scenario_periods] = 1
        matrix = vstack((a_eq, a_ub), format="csr")
        constraint_lower = np.concatenate((b_eq, np.full(row, -np.inf)))
        constraint_upper = np.concatenate((b_eq, b_ub))
        result = milp(
            c=objective,
            integrality=integrality,
            bounds=Bounds(lower, upper),
            constraints=LinearConstraint(matrix, constraint_lower, constraint_upper),
            options={"disp": False, "presolve": True, "mip_rel_gap": MIP_REL_GAP},
        )

    if not result.success or result.x is None:
        raise RuntimeError(
            f"{day} {RELEASE_HOURS[decision_index]}点购电优化失败，"
            f"状态码{result.status}：{result.message}"
        )
    solution = np.asarray(result.x, dtype=float).copy()
    # HiGHS 可能在非负下界返回约 -1e-12 的零值噪声。只清理项目容差内的
    # 负零，并在下方重新复算全部约束；更大的负值仍由独立校验拒绝。
    solution[(solution < 0.0) & (solution >= -BOUND_TOLERANCE)] = 0.0
    gap = float(getattr(result, "mip_gap", 0.0) or 0.0)
    absolute = (
        solution[offset_absolute : offset_absolute + n].copy()
        if adjusted else np.zeros(n)
    )
    charge = solution[offset_charge : offset_charge + scenario_periods].reshape(scenario_count, n)
    discharge = solution[offset_discharge : offset_discharge + scenario_periods].reshape(scenario_count, n)
    emergency = solution[offset_emergency : offset_emergency + scenario_periods].reshape(scenario_count, n)
    residual = solution[offset_residual : offset_residual + scenario_periods].reshape(scenario_count, n)
    mode = solution[offset_mode : offset_mode + scenario_periods].reshape(scenario_count, n)
    stored_energy = np.vstack([solution[energy_slice(s)] for s in range(scenario_count)])
    terminal_deviation = solution[offset_terminal : offset_terminal + scenario_count].copy()
    expected_emergency_cost = float(
        np.sum(
            scenarios.weights[:, None]
            * EMERGENCY_PRICE_MULTIPLIER
            * price_h[None, :]
            * emergency
        )
    )
    expected_terminal_penalty = float(
        0.0 if hard_terminal else terminal_penalty * scenarios.weights @ terminal_deviation
    )
    throughput_regularization = float(
        throughput_epsilon
        * np.sum(scenarios.weights[:, None] * (charge + discharge))
    )
    decision = PurchaseDecision(
        decision_index=decision_index,
        start_period=start,
        grid_purchase=solution[offset_grid : offset_grid + n].copy(),
        charge=charge.copy(),
        discharge=discharge.copy(),
        emergency_purchase=emergency.copy(),
        residual_energy=residual.copy(),
        mode=mode.copy(),
        stored_energy=stored_energy.copy(),
        absolute_adjustment=absolute,
        terminal_deviation=terminal_deviation,
        objective_value=float(result.fun),
        expected_emergency_cost=expected_emergency_cost,
        expected_terminal_penalty=expected_terminal_penalty,
        throughput_regularization=throughput_regularization,
        solver_message=str(result.message),
        solver_kind=solver_kind,
        mip_gap=gap,
    )
    validate_purchase_decision(
        day, decision, price, scenarios, initial_energy, terminal_penalty,
        initial_plan, hard_terminal
    )
    return decision


def validate_purchase_decision(
    day: date,
    decision: PurchaseDecision,
    price: np.ndarray,
    scenarios: NetLoadScenarios,
    initial_energy: float,
    terminal_penalty: float,
    initial_plan: np.ndarray | None,
    hard_terminal: bool,
) -> None:
    """独立复算情景购电解的物理约束、风险费用和目标值。"""
    start = decision.start_period
    n = N_PERIODS - start
    scenario_count = len(scenarios.weights)
    if start != RELEASE_PERIODS[decision.decision_index]:
        raise RuntimeError(f"{day}购电决策起点与发布时间不一致。")
    if decision.grid_purchase.shape != (n,) or not np.isfinite(decision.grid_purchase).all():
        raise RuntimeError(f"{day}购电解grid_purchase形状或数值无效。")
    for name in ("charge", "discharge", "emergency_purchase", "residual_energy", "mode"):
        values = getattr(decision, name)
        if values.shape != (scenario_count, n) or not np.isfinite(values).all():
            raise RuntimeError(f"{day}购电解{name}形状或数值无效。")
    if decision.absolute_adjustment.shape != (n,) or not np.isfinite(decision.absolute_adjustment).all():
        raise RuntimeError(f"{day}购电解absolute_adjustment形状或数值无效。")
    if decision.stored_energy.shape != (scenario_count, n + 1) or not np.isfinite(decision.stored_energy).all():
        raise RuntimeError(f"{day}购电解储电量轨迹无效。")
    net = scenarios.net_load_energy[:, start:]
    balance = (
        decision.grid_purchase[None, :] + decision.discharge
        + decision.emergency_purchase - decision.residual_energy
        - net - decision.charge
    )
    state = (
        decision.stored_energy[:, 1:] - decision.stored_energy[:, :-1]
        - ETA_CHARGE * decision.charge + decision.discharge / ETA_DISCHARGE
    )
    failures: list[str] = []
    if np.max(np.abs(balance)) > BALANCE_TOLERANCE:
        failures.append("预测电量平衡残差超限")
    if np.max(np.abs(state)) > BALANCE_TOLERANCE:
        failures.append("预测状态转移残差超限")
    if np.max(np.abs(decision.stored_energy[:, 0] - initial_energy)) > BOUND_TOLERANCE:
        failures.append("预测初态不等于决策时实际状态")
    if min(
        np.min(decision.grid_purchase), np.min(decision.charge),
        np.min(decision.discharge), np.min(decision.emergency_purchase),
        np.min(decision.residual_energy),
    ) < -BOUND_TOLERANCE:
        failures.append("预测候选量出现负数")
    if np.min(decision.stored_energy) < E_MIN - BOUND_TOLERANCE or np.max(decision.stored_energy) > E_MAX + BOUND_TOLERANCE:
        failures.append("预测储电量超出安全范围")
    if np.max(decision.charge - CHARGE_LIMIT * decision.mode) > BOUND_TOLERANCE:
        failures.append("充电模式约束失败")
    if np.max(decision.discharge - DISCHARGE_LIMIT * (1.0 - decision.mode)) > BOUND_TOLERANCE:
        failures.append("放电模式约束失败")
    charge_capacity = np.maximum(decision.grid_purchase[None, :] - net, 0.0)
    discharge_capacity = np.maximum(net - decision.grid_purchase[None, :], 0.0)
    if np.max(decision.charge - charge_capacity) > BOUND_TOLERANCE:
        failures.append("候选充电使用了紧急购电")
    if np.max(decision.discharge - discharge_capacity) > BOUND_TOLERANCE:
        failures.append("候选放电超过净负荷缺口")
    if np.max(np.abs(decision.mode - np.rint(decision.mode))) > INTEGER_TOLERANCE:
        failures.append("模式变量不是二元值")
    if hard_terminal:
        if max(q2.terminal_interval_error(float(value)) for value in decision.stored_energy[:, -1]) > TERMINAL_TOLERANCE:
            failures.append("年末候选储电量不在附加区间")
    else:
        expected_deviation = np.abs(decision.stored_energy[:, -1] - E_REFERENCE)
        if np.max(np.abs(decision.terminal_deviation - expected_deviation)) > BALANCE_TOLERANCE:
            failures.append("日末绝对偏差线性化失败")
    expected_objective = float(price[start:] @ decision.grid_purchase)
    if decision.decision_index > 0:
        if initial_plan is None:
            failures.append("日内决策缺少0点结算基准")
        else:
            expected_absolute = np.abs(decision.grid_purchase - initial_plan[start:])
            if np.max(np.abs(decision.absolute_adjustment - expected_absolute)) > BALANCE_TOLERANCE:
                failures.append("调整绝对差额线性化失败")
            expected_objective += float(0.5 * price[start:] @ expected_absolute)
    elif np.max(np.abs(decision.absolute_adjustment)) > BOUND_TOLERANCE:
        failures.append("0点计划不应含调整差额")
    if not hard_terminal:
        expected_objective += terminal_penalty * float(
            scenarios.weights @ decision.terminal_deviation
        )
    throughput_epsilon = float(1e-4 * np.min(price[start:]))
    expected_objective += throughput_epsilon * float(
        np.sum(scenarios.weights[:, None] * (decision.charge + decision.discharge))
    )
    expected_objective += float(np.sum(
        scenarios.weights[:, None] * EMERGENCY_PRICE_MULTIPLIER
        * price[start:][None, :] * decision.emergency_purchase
    ))
    if abs(decision.expected_emergency_cost - float(np.sum(
        scenarios.weights[:, None] * EMERGENCY_PRICE_MULTIPLIER
        * price[start:][None, :] * decision.emergency_purchase
    ))) > BALANCE_TOLERANCE:
        failures.append("情景期望紧急购电费复算不一致")
    expected_throughput = float(
        throughput_epsilon
        * np.sum(scenarios.weights[:, None] * (decision.charge + decision.discharge))
    )
    if abs(decision.throughput_regularization - expected_throughput) > BALANCE_TOLERANCE:
        failures.append("规划吞吐正则项复算不一致")
    if abs(decision.objective_value - expected_objective) > 5 * BALANCE_TOLERANCE:
        failures.append("目标函数原始精度复算不一致")
    if decision.mip_gap > MIP_REL_GAP + 1e-12:
        failures.append("MIP相对间隙超限")
    if failures:
        raise RuntimeError(
            f"{day} {RELEASE_HOURS[decision.decision_index]}点购电模型校验失败："
            + "；".join(failures)
        )


def update_realtime_weights(
    scenarios: NetLoadScenarios,
    actual_net_load: np.ndarray,
    period_index: int,
) -> np.ndarray:
    """用边界后截至当前时段的新增实测更新已保留情景概率。"""
    start = RELEASE_PERIODS[scenarios.decision_index]
    if period_index < start:
        raise ValueError("实时情景权重不能使用边界之前的当前时段。")
    observed = actual_net_load[start : period_index + 1]
    candidates = scenarios.net_load_energy[:, start : period_index + 1]
    scale = max(float(np.std(candidates)), float(np.mean(np.abs(observed))) * 0.1, 1.0)
    distance = np.sqrt(np.mean(((candidates - observed) / scale) ** 2, axis=1))
    likelihood = np.exp(
        -0.5 * (distance / scenarios.parameters.similarity_scale) ** 2
    )
    posterior = scenarios.weights * likelihood
    if not np.isfinite(posterior).all() or float(np.sum(posterior)) <= 1e-300:
        posterior = scenarios.weights.copy()
    posterior /= np.sum(posterior)
    return posterior


def build_realtime_blocks(period_index: int) -> tuple[tuple[int, ...], ...]:
    """把当前时段至下一个6小时边界切成近细远粗的候选控制块。"""
    if not 0 <= period_index < N_PERIODS:
        raise ValueError("实时滚动时段索引超出范围。")
    boundary_candidates = RELEASE_PERIODS[1:] + (N_PERIODS,)
    boundary_period = next(
        boundary for boundary in boundary_candidates if boundary >= period_index + 1
    )
    fine_stop = min(period_index + REALTIME_FINE_PERIODS, boundary_period)
    blocks: list[tuple[int, ...]] = [
        (period,) for period in range(period_index, fine_stop)
    ]
    for block_start in range(fine_stop, boundary_period, REALTIME_AGGREGATION_PERIODS):
        blocks.append(tuple(
            range(block_start, min(block_start + REALTIME_AGGREGATION_PERIODS, boundary_period))
        ))
    if not blocks or blocks[0] != (period_index,):
        raise RuntimeError("实时块集合首块必须是当前单个10分钟时段。")
    flattened = tuple(period for block in blocks for period in block)
    expected = tuple(range(period_index, boundary_period))
    if flattened != expected or len(set(flattened)) != len(flattened):
        raise RuntimeError("实时块集合存在重叠、缺漏或越过信息边界。")
    if any(len(block) > REALTIME_AGGREGATION_PERIODS for block in blocks[REALTIME_FINE_PERIODS:]):
        raise RuntimeError("实时远端聚合块长度超过3个10分钟时段。")
    return tuple(blocks)


def solve_realtime_scenario_step(
    period_index: int,
    price: np.ndarray,
    purchase_version: np.ndarray,
    scenarios: NetLoadScenarios,
    scenario_weights: np.ndarray,
    actual_load_now: float,
    actual_photovoltaic_now: float,
    current_energy: float,
    boundary_target_energy: float,
) -> RealtimeStep:
    """固定购电后求解下一信息边界内的多时间尺度情景LP，只执行当前块动作。"""
    if not 0 <= period_index < N_PERIODS:
        raise ValueError("实时滚动时段索引超出范围。")
    scenario_count = len(scenario_weights)
    validate_nonnegative_array(purchase_version, "实时有效购电版本", (N_PERIODS,))
    validate_nonnegative_array(price, "实时电价", (N_PERIODS,))
    validate_finite_array(
        scenarios.net_load_energy,
        "实时净负荷情景",
        (scenario_count, N_PERIODS),
    )
    if scenario_weights.shape != (scenario_count,) or np.any(scenario_weights < 0):
        raise ValueError("实时情景权重无效。")
    if not np.isclose(np.sum(scenario_weights), 1.0, atol=1e-12):
        raise ValueError("实时情景权重未归一化。")
    if not E_MIN - BOUND_TOLERANCE <= current_energy <= E_MAX + BOUND_TOLERANCE:
        raise ValueError("实时滚动初始储电量超出安全范围。")
    if not isfinite(boundary_target_energy):
        raise ValueError("实时边界储电量目标必须有限。")
    boundary_target_energy = float(np.clip(boundary_target_energy, E_MIN, E_MAX))

    actual_load_now = q2.validate_nonnegative_number(actual_load_now, "当前实际负载")
    actual_photovoltaic_now = q2.validate_nonnegative_number(
        actual_photovoltaic_now, "当前实际光伏"
    )
    actual_net_now = actual_load_now - actual_photovoltaic_now
    blocks = build_realtime_blocks(period_index)
    boundary_period = blocks[-1][-1] + 1
    block_count = len(blocks)
    fine_block_count = min(REALTIME_FINE_PERIODS, boundary_period - period_index)
    aggregated_block_count = block_count - fine_block_count

    net_periods = scenarios.net_load_energy[:, period_index:boundary_period].copy()
    net_periods[:, 0] = actual_net_now
    grid_periods = purchase_version[period_index:boundary_period]
    price_periods = price[period_index:boundary_period]
    block_offsets = tuple(
        tuple(period - period_index for period in block) for block in blocks
    )
    net_blocks = np.asarray(
        [[np.sum(net_periods[s, list(offsets)]) for offsets in block_offsets] for s in range(scenario_count)],
        dtype=float,
    )
    grid_blocks = np.asarray([np.sum(grid_periods[list(offsets)]) for offsets in block_offsets], dtype=float)
    emergency_prices = np.asarray(
        [np.max(price_periods[list(offsets)]) for offsets in block_offsets], dtype=float
    )
    if (
        np.max(np.abs(np.sum(net_blocks, axis=1) - np.sum(net_periods, axis=1))) > BALANCE_TOLERANCE
        or abs(np.sum(grid_blocks) - np.sum(grid_periods)) > BALANCE_TOLERANCE
    ):
        raise RuntimeError("实时聚合块电量守恒校验失败。")
    durations = np.asarray([len(block) * DELTA_T for block in blocks], dtype=float)
    charge_capacity = np.minimum(
        (CHARGE_LIMIT / DELTA_T) * durations[None, :],
        np.maximum(grid_blocks[None, :] - net_blocks, 0.0),
    )
    discharge_capacity = np.minimum(
        (DISCHARGE_LIMIT / DELTA_T) * durations[None, :],
        np.maximum(net_blocks - grid_blocks[None, :], 0.0),
    )
    emergency_capacity = np.maximum(net_blocks - grid_blocks[None, :], 0.0)
    residual_capacity = np.maximum(grid_blocks[None, :] - net_blocks, 0.0)

    scenario_blocks = scenario_count * block_count
    offset_charge = 0
    offset_discharge = offset_charge + scenario_blocks
    offset_emergency = offset_discharge + scenario_blocks
    offset_residual = offset_emergency + scenario_blocks
    offset_energy = offset_residual + scenario_blocks
    offset_positive = offset_energy + scenario_count * (block_count + 1)
    offset_negative = offset_positive + scenario_count
    variable_count = offset_negative + scenario_count

    def action_slice(offset: int, scenario_index: int) -> slice:
        return slice(
            offset + scenario_index * block_count,
            offset + (scenario_index + 1) * block_count,
        )

    def energy_slice(scenario_index: int) -> slice:
        width = block_count + 1
        return slice(
            offset_energy + scenario_index * width,
            offset_energy + (scenario_index + 1) * width,
        )

    objective = np.zeros(variable_count)
    for s, probability in enumerate(scenario_weights):
        objective[action_slice(offset_emergency, s)] = (
            probability * EMERGENCY_PRICE_MULTIPLIER * emergency_prices
        )
        objective[offset_positive + s] = probability * np.max(price)
        objective[offset_negative + s] = probability * np.max(price)

    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    upper[offset_charge : offset_charge + scenario_blocks] = charge_capacity.ravel()
    upper[offset_discharge : offset_discharge + scenario_blocks] = discharge_capacity.ravel()
    upper[offset_emergency : offset_emergency + scenario_blocks] = emergency_capacity.ravel()
    upper[offset_residual : offset_residual + scenario_blocks] = residual_capacity.ravel()
    for s in range(scenario_count):
        e_slice = energy_slice(s)
        lower[e_slice] = E_MIN
        upper[e_slice] = E_MAX
        lower[e_slice.start] = current_energy
        upper[e_slice.start] = current_energy

    rows: list[int] = []
    cols: list[int] = []
    values: list[float] = []
    rhs: list[float] = []
    row = 0
    for s in range(scenario_count):
        c_start = action_slice(offset_charge, s).start
        d_start = action_slice(offset_discharge, s).start
        h_start = action_slice(offset_emergency, s).start
        y_start = action_slice(offset_residual, s).start
        e_start = energy_slice(s).start
        for block_index in range(block_count):
            for column, value in (
                (c_start + block_index, -1.0),
                (d_start + block_index, 1.0),
                (h_start + block_index, 1.0),
                (y_start + block_index, -1.0),
            ):
                rows.append(row); cols.append(column); values.append(value)
            rhs.append(float(net_blocks[s, block_index] - grid_blocks[block_index]))
            row += 1
        for block_index in range(block_count):
            for column, value in (
                (e_start + block_index, -1.0),
                (e_start + block_index + 1, 1.0),
                (c_start + block_index, -ETA_CHARGE),
                (d_start + block_index, 1.0 / ETA_DISCHARGE),
            ):
                rows.append(row); cols.append(column); values.append(value)
            rhs.append(0.0)
            row += 1
        for column, value in (
            (e_start + block_count, 1.0),
            (offset_positive + s, -1.0),
            (offset_negative + s, 1.0),
        ):
            rows.append(row); cols.append(column); values.append(value)
        rhs.append(boundary_target_energy)
        row += 1

    # 当前块动作在全部情景间共享，远端块允许按情景补救。
    for s in range(1, scenario_count):
        for offset in (offset_charge, offset_discharge, offset_emergency):
            rows.extend((row, row))
            cols.extend((action_slice(offset, s).start, action_slice(offset, 0).start))
            values.extend((1.0, -1.0))
            rhs.append(0.0)
            row += 1

    matrix = coo_matrix((values, (rows, cols)), shape=(row, variable_count)).tocsr()
    right_hand_side = np.asarray(rhs, dtype=float)
    result = linprog(
        c=objective,
        A_eq=matrix,
        b_eq=right_hand_side,
        bounds=np.column_stack((lower, upper)),
        method="highs-ds",
        options={
            "presolve": True,
            "primal_feasibility_tolerance": 1e-9,
            "dual_feasibility_tolerance": 1e-9,
        },
    )
    if not result.success or result.x is None:
        raise RuntimeError(
            f"第{period_index + 1}时段情景实时优化失败，"
            f"状态码{result.status}：{result.message}"
        )
    solution = np.asarray(result.x, dtype=float).copy()
    solution[(solution < 0.0) & (solution >= -BOUND_TOLERANCE)] = 0.0
    if (
        not np.isfinite(solution).all()
        or np.max(np.abs(matrix @ solution - right_hand_side)) > BALANCE_TOLERANCE
        or np.any(solution < lower - BOUND_TOLERANCE)
        or np.any(solution > upper + BOUND_TOLERANCE)
    ):
        raise RuntimeError(f"第{period_index + 1}时段情景实时解复算失败。")

    charge = float(solution[action_slice(offset_charge, 0).start])
    discharge = float(solution[action_slice(offset_discharge, 0).start])
    emergency = float(solution[action_slice(offset_emergency, 0).start])
    actual_surplus = (
        purchase_version[period_index] + actual_photovoltaic_now + discharge
        + emergency - actual_load_now - charge
    )
    if actual_surplus < -BALANCE_TOLERANCE:
        raise RuntimeError("实时当前动作不能满足实际电量平衡。")
    actual_surplus = max(actual_surplus, 0.0)
    curtailment = min(actual_photovoltaic_now, actual_surplus)
    unused = actual_surplus - curtailment
    predicted_boundary = float(sum(
        scenario_weights[s] * solution[energy_slice(s).stop - 1]
        for s in range(scenario_count)
    ))
    return RealtimeStep(
        period_index=period_index,
        horizon=boundary_period - period_index,
        boundary_period=boundary_period,
        fine_block_count=fine_block_count,
        aggregated_block_count=aggregated_block_count,
        charge=charge,
        discharge=discharge,
        emergency_purchase=emergency,
        curtailment=curtailment,
        unused_planned_purchase=unused,
        predicted_boundary_energy=predicted_boundary,
        charge_capacity=float(charge_capacity[0, 0]),
        discharge_capacity=float(discharge_capacity[0, 0]),
        solver_message=str(result.message),
    )





def choose_cost_parameters(
    cost_history: dict[str, list[tuple[date, float]]],
) -> tuple[ScenarioParameters, date | None, int, float | None]:
    """只用已结束日期的候选策略实际账单选择当天参数。"""
    completed = min(len(cost_history[item.key]) for item in PARAMETER_CANDIDATES)
    if completed < MIN_TUNING_DAYS:
        latest = (
            cost_history[DEFAULT_PARAMETERS.key][-1][0]
            if cost_history[DEFAULT_PARAMETERS.key] else None
        )
        return DEFAULT_PARAMETERS, latest, completed, None
    scores: list[float] = []
    counts: list[int] = []
    latest_dates: list[date] = []
    for candidate in PARAMETER_CANDIDATES:
        records = cost_history[candidate.key][-candidate.validation_window :]
        scores.append(float(sum(value for _, value in records)))
        counts.append(len(records))
        latest_dates.append(records[-1][0])
    best = min(range(len(PARAMETER_CANDIDATES)), key=lambda index: (scores[index], index))
    return PARAMETER_CANDIDATES[best], latest_dates[best], counts[best], scores[best]


def simulate_day(
    data: ModelData,
    day_index: int,
    initial_energy: float,
    parameters: ScenarioParameters,
    spec: StrategySpec,
    parameter_latest_validation_date: date | None,
    parameter_validation_days: int,
    parameter_validation_cost: float | None,
) -> DailyResult:
    """按给定参数完整重放一天的四次购电和144次实时控制。"""
    day = data.base.dates[day_index]
    # 12月31日也采用普通日期的日末软惩罚；年末区间只在实际执行结束后验收。
    hard_terminal = False
    terminal_penalty = float(TERMINAL_PENALTY_MULTIPLIER * np.max(data.base.price))
    actual_load = data.base.actual_load_energy[day_index]
    actual_pv = data.base.actual_photovoltaic_energy[day_index]
    actual_net = data.actual_net_load_energy[day_index]
    scenario_versions: list[NetLoadScenarios] = []
    decisions: list[PurchaseDecision | None] = []
    versions = np.empty((4, N_PERIODS), dtype=float)
    effective = np.empty(N_PERIODS, dtype=float)
    replay_charge = np.zeros(N_PERIODS)
    replay_discharge = np.zeros(N_PERIODS)
    replay_emergency = np.zeros(N_PERIODS)
    replay_curtailment = np.zeros(N_PERIODS)
    replay_unused = np.zeros(N_PERIODS)
    replay_steps: list[RealtimeStep] = []
    stored = np.empty(N_PERIODS + 1)
    stored[0] = initial_energy
    accepted_batch = 0

    scenarios0 = build_net_load_scenarios(data, day_index, 0, 0, parameters)
    decision0 = solve_purchase_decision(
        day, 0, data.base.price, scenarios0, initial_energy,
        terminal_penalty, None, hard_terminal
    )
    scenario_versions.append(scenarios0)
    decisions.append(decision0)
    versions[0] = decision0.grid_purchase
    initial_plan = versions[0].copy()
    latest_complete_decision = decision0
    latest_complete_scenarios = scenarios0

    for decision_index in range(4):
        start = RELEASE_PERIODS[decision_index]
        stop = start + PERIODS_PER_BLOCK
        if decision_index > 0:
            if decision_index in spec.forecast_updates:
                accepted_batch = decision_index
            scenarios = build_net_load_scenarios(
                data, day_index, decision_index, accepted_batch, parameters
            )
            scenario_versions.append(scenarios)
            versions[decision_index] = versions[decision_index - 1]
            if decision_index in spec.purchase_updates:
                decision = solve_purchase_decision(
                    day, decision_index, data.base.price, scenarios,
                    float(stored[start]),
                    terminal_penalty,
                    initial_plan, hard_terminal
                )
                versions[decision_index, start:] = decision.grid_purchase
                decisions.append(decision)
                latest_complete_decision = decision
                latest_complete_scenarios = scenarios
            else:
                decisions.append(None)
        scenarios = scenario_versions[decision_index]
        version = versions[decision_index]
        for period_index in range(start, stop):
            posterior = update_realtime_weights(
                scenarios, actual_net, period_index
            )
            boundary_period = build_realtime_blocks(period_index)[-1][-1] + 1
            target_offset = boundary_period - latest_complete_decision.start_period
            if not 0 <= target_offset < latest_complete_decision.stored_energy.shape[1]:
                raise RuntimeError(f"{day}第{period_index + 1}时段无法取得实时边界状态目标。")
            boundary_target = float(
                latest_complete_decision.stored_energy[:, target_offset]
                @ latest_complete_scenarios.weights
            )
            # 购电已在边界提交后，当前实际净负荷才进入实时控制。
            step = solve_realtime_scenario_step(
                period_index=period_index,
                price=data.base.price,
                purchase_version=version,
                scenarios=scenarios,
                scenario_weights=posterior,
                actual_load_now=float(actual_load[period_index]),
                actual_photovoltaic_now=float(actual_pv[period_index]),
                current_energy=float(stored[period_index]),
                boundary_target_energy=boundary_target,
            )
            replay_steps.append(step)
            effective[period_index] = version[period_index]
            replay_charge[period_index] = step.charge
            replay_discharge[period_index] = step.discharge
            replay_emergency[period_index] = step.emergency_purchase
            replay_curtailment[period_index] = step.curtailment
            replay_unused[period_index] = step.unused_planned_purchase
            stored[period_index + 1] = (
                stored[period_index] + ETA_CHARGE * step.charge
                - step.discharge / ETA_DISCHARGE
            )

    replay = ReplayResult(
        charge=replay_charge,
        discharge=replay_discharge,
        emergency_purchase=replay_emergency,
        curtailment=replay_curtailment,
        unused_planned_purchase=replay_unused,
        stored_energy=stored,
        steps=tuple(replay_steps),
    )
    increase = np.maximum(effective - initial_plan, 0.0)
    decrease = np.maximum(initial_plan - effective, 0.0)
    initial_cost = float(data.base.price @ initial_plan)
    increase_cost = float(1.5 * data.base.price @ increase)
    decrease_refund = float(data.base.price @ decrease)
    decrease_penalty = float(0.5 * data.base.price @ decrease)
    emergency_cost = float(
        EMERGENCY_PRICE_MULTIPLIER * data.base.price @ replay_emergency
    )
    total_cost = (
        initial_cost + increase_cost - decrease_refund
        + decrease_penalty + emergency_cost
    )
    daily = DailyResult(
        day=day,
        parameters=parameters,
        parameter_latest_validation_date=parameter_latest_validation_date,
        parameter_validation_days=parameter_validation_days,
        parameter_validation_cost=parameter_validation_cost,
        scenarios=tuple(scenario_versions),
        decisions=tuple(decisions),
        purchase_versions=versions,
        effective_purchase=effective,
        actual_load_energy=actual_load.copy(),
        actual_photovoltaic_energy=actual_pv.copy(),
        replay=replay,
        initial_plan_cost=initial_cost,
        increase_purchase_cost=increase_cost,
        decrease_refund=decrease_refund,
        decrease_penalty=decrease_penalty,
        emergency_purchase_cost=emergency_cost,
        realized_total_cost=total_cost,
    )
    validate_daily_result(data, daily, float(initial_energy), spec)
    return daily


def run_strategy(
    data: ModelData,
    spec: StrategySpec,
    day_limit: int | None = None,
    show_progress: bool = False,
) -> StrategyResult:
    """连续运行策略，并在每日结束后用实际总费用更新候选参数评分。"""
    count = len(data.base.dates) if day_limit is None else day_limit
    if not 1 <= count <= len(data.base.dates):
        raise ValueError("运行天数必须在1至365之间。")
    results: list[DailyResult] = []
    initial_energy = E_INITIAL
    cost_history = {candidate.key: [] for candidate in PARAMETER_CANDIDATES}
    tuning_replay_count = 0
    tuning_records: list[TuningRecord] = []

    for day_index in range(count):
        selected, latest, validation_days, validation_cost = choose_cost_parameters(
            cost_history
        )
        daily = simulate_day(
            data, day_index, initial_energy, selected, spec,
            latest, validation_days, validation_cost
        )
        results.append(daily)

        # 日期结束后才能看到完整实际轨迹和账单；候选回放从同一实际日初状态出发。
        for candidate in PARAMETER_CANDIDATES:
            if candidate == selected:
                replayed = daily
            else:
                replayed = simulate_day(
                    data, day_index, initial_energy, candidate, spec,
                    latest, validation_days, validation_cost
                )
                tuning_replay_count += 1
            cost_history[candidate.key].append((replayed.day, replayed.realized_total_cost))
            tuning_records.append(TuningRecord(
                validation_day=replayed.day,
                parameters=candidate,
                realized_total_cost=replayed.realized_total_cost,
            ))

        initial_energy = float(daily.replay.stored_energy[-1])
        if show_progress and ((day_index + 1) % 10 == 0 or day_index + 1 == count):
            print(
                f"[{spec.name}] 已完成 {day_index + 1}/{count} 天；"
                f"参数={selected.key}；末态={initial_energy:.3f} kWh",
                flush=True,
            )

    warmup_count = min(JANUARY_DAYS, count)
    return StrategyResult(
        spec=spec,
        warmup_days=tuple(results[:warmup_count]),
        official_days=tuple(results[warmup_count:]),
        tuning_replay_count=tuning_replay_count,
        tuning_records=tuple(tuning_records),
    )


def validate_daily_result(
    data: ModelData,
    result: DailyResult,
    expected_initial_energy: float,
    spec: StrategySpec,
) -> None:
    """检查购电版本不可回写、实际物理约束及费用拆账。"""
    if len(result.scenarios) != 4 or len(result.decisions) != 4:
        raise RuntimeError(f"{result.day}缺少四个边界情景/决策记录。")
    if result.purchase_versions.shape != (4, N_PERIODS):
        raise RuntimeError(f"{result.day}购电版本矩阵形状错误。")
    validate_nonnegative_array(result.purchase_versions, "购电版本", (4, N_PERIODS))
    validate_nonnegative_array(result.effective_purchase, "有效购电", (N_PERIODS,))
    for k in range(1, 4):
        cutoff = RELEASE_PERIODS[k]
        if np.max(np.abs(
            result.purchase_versions[k, :cutoff]
            - result.purchase_versions[k - 1, :cutoff]
        )) > BALANCE_TOLERANCE:
            raise RuntimeError(f"{result.day}第{k}次版本回写了已执行时段。")
        if (k in spec.purchase_updates) != (result.decisions[k] is not None):
            raise RuntimeError(f"{result.day}购电更新记录与策略定义不一致。")
        expected_batch = max((item for item in (0,) + spec.forecast_updates if item <= k), default=0)
        if result.scenarios[k].source_batch_index != expected_batch:
            raise RuntimeError(f"{result.day}光伏批次接受记录与策略定义不一致。")
    for scenarios in result.scenarios:
        if scenarios.parameters != result.parameters:
            raise RuntimeError(f"{result.day}情景参数版本不统一。")
        if scenarios.latest_training_date is not None and scenarios.latest_training_date >= result.day:
            raise RuntimeError(f"{result.day}净负荷情景使用了前视历史。")
    expected_effective = np.concatenate(
        [result.purchase_versions[k, RELEASE_PERIODS[k] : RELEASE_PERIODS[k] + PERIODS_PER_BLOCK] for k in range(4)]
    )
    if np.max(np.abs(expected_effective - result.effective_purchase)) > BALANCE_TOLERANCE:
        raise RuntimeError(f"{result.day}最终有效购电不能由版本矩阵还原。")
    if abs(result.replay.stored_energy[0] - expected_initial_energy) > BOUND_TOLERANCE:
        raise RuntimeError(f"{result.day}跨日实际储电量不连续。")
    if len(result.replay.steps) != N_PERIODS:
        raise RuntimeError(f"{result.day}实时求解记录数量不是{N_PERIODS}。")
    for period_index, step in enumerate(result.replay.steps):
        if step.period_index != period_index:
            raise RuntimeError(f"{result.day}实时求解记录时段顺序错误。")
        expected_blocks = build_realtime_blocks(period_index)
        expected_boundary = expected_blocks[-1][-1] + 1
        if (
            step.boundary_period != expected_boundary
            or step.horizon != expected_boundary - period_index
            or step.fine_block_count != min(REALTIME_FINE_PERIODS, expected_boundary - period_index)
            or step.aggregated_block_count != len(expected_blocks) - step.fine_block_count
        ):
            raise RuntimeError(f"{result.day}第{period_index + 1}时段实时块诊断不一致。")
        if expected_blocks[0] != (period_index,) or any(
            len(block) > REALTIME_AGGREGATION_PERIODS for block in expected_blocks
        ):
            raise RuntimeError(f"{result.day}第{period_index + 1}时段实时块集合非法。")
    q2.validate_replay(
        result.day,
        SimpleNamespace(grid_purchase=result.effective_purchase),
        result.replay,
        result.actual_load_energy,
        result.actual_photovoltaic_energy,
    )
    g0 = result.purchase_versions[0]
    effective = result.effective_purchase
    increase = np.maximum(effective - g0, 0.0)
    decrease = np.maximum(g0 - effective, 0.0)
    non_emergency_split = (
        data.base.price @ g0 + 1.5 * data.base.price @ increase
        - data.base.price @ decrease + 0.5 * data.base.price @ decrease
    )
    non_emergency_absolute = (
        data.base.price @ effective
        + 0.5 * data.base.price @ np.abs(effective - g0)
    )
    if abs(float(non_emergency_split - non_emergency_absolute)) > BALANCE_TOLERANCE:
        raise RuntimeError(f"{result.day}调整费用两种等价口径不一致。")
    expected_total = float(
        non_emergency_split
        + EMERGENCY_PRICE_MULTIPLIER * data.base.price @ result.replay.emergency_purchase
    )
    fields = (
        result.initial_plan_cost,
        result.increase_purchase_cost,
        result.decrease_refund,
        result.decrease_penalty,
        result.emergency_purchase_cost,
        result.realized_total_cost,
    )
    if not np.isfinite(fields).all() or abs(expected_total - result.realized_total_cost) > BALANCE_TOLERANCE:
        raise RuntimeError(f"{result.day}实际账单原始精度复算失败。")


def validate_strategy(data: ModelData, result: StrategyResult, complete_year: bool) -> None:
    """检查日期顺序、跨日连续性和实际年末附加条件。"""
    days = result.all_days
    expected = tuple(data.base.dates[: len(days)])
    if tuple(item.day for item in days) != expected:
        raise RuntimeError(f"{result.spec.name}日期顺序或数量错误。")
    previous = E_INITIAL
    for item in days:
        if abs(item.replay.stored_energy[0] - previous) > BOUND_TOLERANCE:
            raise RuntimeError(f"{result.spec.name}在{item.day}跨日状态不连续。")
        if (
            item.parameter_latest_validation_date is not None
            and item.parameter_latest_validation_date >= item.day
        ):
            raise RuntimeError(f"{item.day}费用调参出现前视验证。")
        previous = float(item.replay.stored_energy[-1])
    if complete_year:
        if len(result.warmup_days) != JANUARY_DAYS or len(result.official_days) != OFFICIAL_DAYS:
            raise RuntimeError(f"{result.spec.name}初始化期或正式期天数错误。")
        final = days[-1]
        actual_error = q2.terminal_interval_error(float(final.replay.stored_energy[-1]))
        if actual_error > TERMINAL_TOLERANCE:
            raise RuntimeError(
                f"{result.spec.name}实际年末储电量{final.replay.stored_energy[-1]:.6f} kWh"
                f"不在[{E_TERMINAL_MIN:g},{E_TERMINAL_MAX:g}] kWh，越界{actual_error:.6f} kWh。"
            )


def display_number(value: float) -> str:
    """六位小数展示，并消除负零。"""
    if abs(value) < DISPLAY_ZERO_TOLERANCE:
        value = 0.0
    return f"{float(value):.{DISPLAY_DECIMALS}f}"


def excel_number(value: float) -> float:
    """仅在写出Excel时取六位小数。"""
    return float(display_number(float(value)))


def write_csv(path: Path, header: Sequence[str], rows: Iterable[Sequence[object]]) -> Path:
    """以Excel可识别的UTF-8 BOM写出CSV。"""
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    return path


def internal_interval_label(period_index: int) -> str:
    """内部真实10分钟区间标签。"""
    return (
        f"{q2.format_clock(period_index * 10)}-"
        f"{q2.format_clock((period_index + 1) * 10)}"
    )


def template_headers() -> tuple[str, ...]:
    """只读获取result3计划购电表的144个显示标签。"""
    workbook = load_workbook(RESULT_TEMPLATE_FILE, read_only=True, data_only=False)
    try:
        sheet = workbook["计划购电量"]
        return tuple(sheet.cell(1, column).value for column in range(2, N_PERIODS + 2))
    finally:
        workbook.close()


def write_mapping_csv() -> Path:
    """写出内部区间到整体后移10分钟模板列的一一映射。"""
    headers = template_headers()
    rows = (
        (
            t + 1,
            internal_interval_label(t),
            data_label,
            get_column_letter(t + 2),
            t + 2,
            headers[t],
            "数值不循环移动；末段仍归本日期",
        )
        for t, data_label in enumerate(q2.load_model_data().interval_labels)
    )
    return write_csv(
        MAPPING_FILE,
        ("内部时段序号", "内部真实区间", "附件1/2结束时刻口径区间", "模板列", "模板列号", "模板显示区间", "映射说明"),
        rows,
    )


def write_decision_csv(results: Sequence[DailyResult]) -> Path:
    """逐情景写出主方案购电决策、来源、概率、风险费用与候选轨迹。"""
    header = (
        "日期", "阶段", "决策时刻", "可用实测截止", "光伏批次来源", "光伏补起点来源",
        "净负荷训练截止", "费用验证截止", "费用验证天数", "费用验证总成本",
        "参数版本", "衰减天数", "相似度尺度", "目标情景数", "重采样规则",
        "实际保留情景数", "情景序号", "情景来源日期", "情景权重", "决策时储电量",
        "内部时段", "内部真实区间", "光伏预报特征电量", "净负荷诊断0.8分位数",
        "情景净负荷电量", "0点初始计划",
        "上版购电量", "本次提交购电量", "相对0点绝对差额", "候选充电量", "候选放电量",
        "候选紧急购电量", "候选剩余电量", "候选时段初储电量", "候选时段末储电量",
        "目标函数值", "期望紧急购电费", "期望末态惩罚", "吞吐正则项",
        "求解路径", "求解状态", "MIP相对间隙",
    )

    def rows() -> Iterable[Sequence[object]]:
        for result in results:
            for k, decision in enumerate(result.decisions):
                if decision is None:
                    raise RuntimeError("主方案决策版本CSV要求四次购电决策均存在。")
                scenarios = result.scenarios[k]
                start = RELEASE_PERIODS[k]
                previous = result.purchase_versions[max(k - 1, 0)]
                cutoff_label = (
                    f"{result.day.isoformat()} 00:00以前"
                    if k == 0 else f"{result.day.isoformat()} {RELEASE_HOURS[k]:02d}:00"
                )
                for s, (origin, probability) in enumerate(
                    zip(scenarios.origin_dates, scenarios.weights)
                ):
                    for local_index, t in enumerate(range(start, N_PERIODS)):
                        yield (
                            result.day.isoformat(),
                            "初始化" if result.day < OFFICIAL_START else "正式期",
                            f"{RELEASE_HOURS[k]:02d}:00",
                            cutoff_label,
                            f"{RELEASE_HOURS[scenarios.source_batch_index]:02d}:00批次",
                            scenarios.photovoltaic_start_source,
                            scenarios.latest_training_date or "附件1冷启动",
                            result.parameter_latest_validation_date or "预设参数",
                            result.parameter_validation_days,
                            "" if result.parameter_validation_cost is None else display_number(result.parameter_validation_cost),
                            result.parameters.key,
                            display_number(result.parameters.decay_days),
                            display_number(result.parameters.similarity_scale),
                            result.parameters.scenario_count,
                            result.parameters.sampling_rule,
                            len(scenarios.weights),
                            s + 1,
                            origin or "附件1减光伏预报冷启动",
                            display_number(probability),
                            display_number(decision.stored_energy[s, 0]),
                            t + 1,
                            internal_interval_label(t),
                            display_number(scenarios.photovoltaic_feature_energy[t]),
                            display_number(scenarios.diagnostic_quantile_energy[t]),
                            display_number(scenarios.net_load_energy[s, t]),
                            display_number(result.purchase_versions[0, t]),
                            display_number(previous[t]),
                            display_number(decision.grid_purchase[local_index]),
                            display_number(decision.absolute_adjustment[local_index]),
                            display_number(decision.charge[s, local_index]),
                            display_number(decision.discharge[s, local_index]),
                            display_number(decision.emergency_purchase[s, local_index]),
                            display_number(decision.residual_energy[s, local_index]),
                            display_number(decision.stored_energy[s, local_index]),
                            display_number(decision.stored_energy[s, local_index + 1]),
                            display_number(decision.objective_value),
                            display_number(decision.expected_emergency_cost),
                            display_number(decision.expected_terminal_penalty),
                            display_number(decision.throughput_regularization),
                            decision.solver_kind,
                            decision.solver_message,
                            f"{decision.mip_gap:.12g}",
                        )
    return write_csv(DECISION_FILE, header, rows())


def write_detail_csv(data: ModelData, results: Sequence[DailyResult]) -> Path:
    """写出全年52560个实际时段的版本、物理量和分项费用。"""
    headers = template_headers()
    header = (
        "日期", "阶段", "内部时段", "内部真实区间", "模板显示区间", "生效购电版本",
        "实时截断边界", "近端10分钟块数", "远端30分钟块数", "光伏预报批次", "参数版本", "电价", "光伏预报特征电量",
        "净负荷诊断0.8分位数", "实际净负荷电量", "实际负载电量", "实际光伏电量",
        "0点初始计划", "最终有效购电", "增购差额", "减购差额",
        "实际充电量", "实际放电量", "紧急购电量", "弃光量", "未利用有效购电量",
        "时段初储电量", "时段末储电量", "初始计划费", "增购费", "减购退费",
        "减购违约费", "紧急购电费", "时段实际总费",
    )

    def rows() -> Iterable[Sequence[object]]:
        for result in results:
            g0 = result.purchase_versions[0]
            for t in range(N_PERIODS):
                k = t // PERIODS_PER_BLOCK
                scenarios = result.scenarios[k]
                effective = result.effective_purchase[t]
                increase = max(effective - g0[t], 0.0)
                decrease = max(g0[t] - effective, 0.0)
                initial_fee = data.base.price[t] * g0[t]
                increase_fee = 1.5 * data.base.price[t] * increase
                refund = data.base.price[t] * decrease
                penalty = 0.5 * data.base.price[t] * decrease
                emergency_fee = EMERGENCY_PRICE_MULTIPLIER * data.base.price[t] * result.replay.emergency_purchase[t]
                step = result.replay.steps[t]
                yield (
                    result.day.isoformat(),
                    "初始化" if result.day < OFFICIAL_START else "正式期",
                    t + 1,
                    internal_interval_label(t),
                    headers[t],
                    f"{RELEASE_HOURS[k]:02d}:00版本",
                    q2.format_clock(step.boundary_period * 10),
                    step.fine_block_count,
                    step.aggregated_block_count,
                    f"{RELEASE_HOURS[scenarios.source_batch_index]:02d}:00批次",
                    result.parameters.key,
                    display_number(data.base.price[t]),
                    display_number(scenarios.photovoltaic_feature_energy[t]),
                    display_number(scenarios.diagnostic_quantile_energy[t]),
                    display_number(result.actual_load_energy[t] - result.actual_photovoltaic_energy[t]),
                    display_number(result.actual_load_energy[t]),
                    display_number(result.actual_photovoltaic_energy[t]),
                    display_number(g0[t]),
                    display_number(effective),
                    display_number(increase),
                    display_number(decrease),
                    display_number(result.replay.charge[t]),
                    display_number(result.replay.discharge[t]),
                    display_number(result.replay.emergency_purchase[t]),
                    display_number(result.replay.curtailment[t]),
                    display_number(result.replay.unused_planned_purchase[t]),
                    display_number(result.replay.stored_energy[t]),
                    display_number(result.replay.stored_energy[t + 1]),
                    display_number(initial_fee),
                    display_number(increase_fee),
                    display_number(refund),
                    display_number(penalty),
                    display_number(emergency_fee),
                    display_number(initial_fee + increase_fee - refund + penalty + emergency_fee),
                )
    return write_csv(DETAIL_FILE, header, rows())


def executed_diagnostic_net_load(results: Sequence[DailyResult]) -> np.ndarray:
    """拼接各执行块边界形成的净负荷0.8分位数诊断曲线。"""
    curves: list[np.ndarray] = []
    for result in results:
        curves.append(np.concatenate([
            result.scenarios[k].diagnostic_quantile_energy[
                RELEASE_PERIODS[k] : RELEASE_PERIODS[k] + PERIODS_PER_BLOCK
            ]
            for k in range(4)
        ]))
    return np.concatenate(curves)


def executed_scenario_envelope(
    results: Sequence[DailyResult],
) -> tuple[np.ndarray, np.ndarray]:
    """拼接各执行块所用整轨迹情景的逐时段最小值和最大值。"""
    lower: list[np.ndarray] = []
    upper: list[np.ndarray] = []
    for result in results:
        for k in range(4):
            start = RELEASE_PERIODS[k]
            stop = start + PERIODS_PER_BLOCK
            paths = result.scenarios[k].net_load_energy[:, start:stop]
            lower.append(np.min(paths, axis=0))
            upper.append(np.max(paths, axis=0))
    return np.concatenate(lower), np.concatenate(upper)


def strategy_metrics(result: StrategyResult) -> dict[str, float]:
    """按正式期同口径汇总一种对照策略。"""
    days = result.official_days
    diagnostic_net = executed_diagnostic_net_load(days)
    actual_net = np.concatenate([
        item.actual_load_energy - item.actual_photovoltaic_energy for item in days
    ])
    scenario_lower, scenario_upper = executed_scenario_envelope(days)
    return {
        "initial_plan_cost": sum(item.initial_plan_cost for item in days),
        "increase_cost": sum(item.increase_purchase_cost for item in days),
        "decrease_refund": sum(item.decrease_refund for item in days),
        "decrease_penalty": sum(item.decrease_penalty for item in days),
        "emergency_cost": sum(item.emergency_purchase_cost for item in days),
        "total_cost": sum(item.realized_total_cost for item in days),
        "emergency_energy": sum(float(np.sum(item.replay.emergency_purchase)) for item in days),
        "unused_purchase": sum(float(np.sum(item.replay.unused_planned_purchase)) for item in days),
        "curtailment": sum(float(np.sum(item.replay.curtailment)) for item in days),
        "net_mae": float(np.mean(np.abs(actual_net - diagnostic_net))),
        "net_coverage": float(np.mean(actual_net <= diagnostic_net)),
        "scenario_coverage": float(np.mean(
            (actual_net >= scenario_lower) & (actual_net <= scenario_upper)
        )),
        "official_initial_energy": float(days[0].replay.stored_energy[0]),
        "year_end_energy": float(days[-1].replay.stored_energy[-1]),
    }


def write_comparison_csv(results: Sequence[StrategyResult]) -> Path:
    """写出主方案与五种因果预报更新对照。"""
    rows = []
    for result in results:
        metrics = strategy_metrics(result)
        rows.append((
            result.spec.name,
            "0,6,12,18" if result.spec.forecast_updates == (1, 2, 3) else "0" + "".join(f",{RELEASE_HOURS[k]}" for k in result.spec.forecast_updates),
            "0,6,12,18" if result.spec.purchase_updates == (1, 2, 3) else "0" + "".join(f",{RELEASE_HOURS[k]}" for k in result.spec.purchase_updates),
            *[display_number(metrics[key]) for key in (
                "initial_plan_cost", "increase_cost", "decrease_refund", "decrease_penalty",
                "emergency_cost", "total_cost", "emergency_energy", "unused_purchase",
                "curtailment", "net_mae", "net_coverage", "scenario_coverage",
                "official_initial_energy", "year_end_energy",
            )],
        ))
    return write_csv(
        COMPARISON_FILE,
        (
            "策略", "接受光伏预报时刻", "购电决策时刻", "初始计划费", "增购费", "减购退费",
            "减购违约费", "紧急购电费", "正式期实际总费", "紧急购电量", "未利用有效购电量",
            "弃光量", "执行时净负荷诊断MAE", "净负荷0.8分位数覆盖率",
            "净负荷情景包络覆盖率", "2月1日初始储电量", "12月31日末储电量",
        ),
        rows,
    )


def build_check_rows(data: ModelData, result: StrategyResult) -> list[Sequence[object]]:
    """构造主方案预测、约束、费用、状态和输出规模校验。"""
    all_days = result.all_days
    official = result.official_days
    diagnostic_net = executed_diagnostic_net_load(official)
    actual_net = np.concatenate([
        item.actual_load_energy - item.actual_photovoltaic_energy for item in official
    ])
    scenario_lower, scenario_upper = executed_scenario_envelope(official)
    actual_balance: list[np.ndarray] = []
    state_balance: list[np.ndarray] = []
    fee_equivalence: list[float] = []
    for item in all_days:
        actual_balance.append(
            item.effective_purchase + item.actual_photovoltaic_energy
            + item.replay.discharge + item.replay.emergency_purchase
            - item.replay.curtailment - item.replay.unused_planned_purchase
            - item.actual_load_energy - item.replay.charge
        )
        state_balance.append(
            item.replay.stored_energy[1:] - item.replay.stored_energy[:-1]
            - ETA_CHARGE * item.replay.charge + item.replay.discharge / ETA_DISCHARGE
        )
        g0 = item.purchase_versions[0]
        fee_equivalence.append(float(
            data.base.price @ item.effective_purchase
            + 0.5 * data.base.price @ np.abs(item.effective_purchase - g0)
            - item.initial_plan_cost - item.increase_purchase_cost
            + item.decrease_refund - item.decrease_penalty
        ))
    official_metrics = strategy_metrics(result)
    january_cost = sum(item.realized_total_cost for item in result.warmup_days)
    yearly_cost = january_cost + official_metrics["total_cost"]
    final = all_days[-1]
    plan_count = sum(decision is not None for item in all_days for decision in item.decisions)
    planning_decisions = [
        decision for item in all_days for decision in item.decisions if decision is not None
    ]
    realtime_steps = [step for item in all_days for step in item.replay.steps]
    fixed_mode_count = sum(
        decision.solver_kind == "松弛定模式后固定模式LP" for decision in planning_decisions
    )
    milp_fallback_count = sum(
        decision.solver_kind == "MILP回退" for decision in planning_decisions
    )
    rows: list[Sequence[object]] = [
        ("输入", "附件3预报批次数", 365 * 4, "批", "通过"),
        ("输入", "每批原始预报节点数", 24, "个", "通过"),
        ("参数", "候选参数组数", len(PARAMETER_CANDIDATES), "组", "通过"),
        ("参数", "费用调参额外完整日回放次数", result.tuning_replay_count, "日", "参考"),
        ("参数", "费用调参额外购电求解次数", result.tuning_replay_count * (1 + len(result.spec.purchase_updates)), "次", "参考"),
        ("参数", "费用调参额外实时求解次数", result.tuning_replay_count * N_PERIODS, "次", "参考"),
        ("因果性", "主方案购电求解次数", plan_count, "次", "通过" if plan_count == 1460 else "失败"),
        ("求解", "松弛后固定模式LP次数", fixed_mode_count, "次", "通过"),
        ("求解", "购电MILP回退次数", milp_fallback_count, "次", "参考"),
        ("因果性", "主方案实时储能求解次数", len(all_days) * N_PERIODS, "次", "通过"),
        ("实时", "实时块集合完整性校验次数", len(realtime_steps), "次", "通过"),
        ("实时", "最大实时截短原始时段数", max(step.horizon for step in realtime_steps), "个10分钟时段", "通过"),
        ("实时", "最大近端10分钟块数", max(step.fine_block_count for step in realtime_steps), "块", "通过"),
        ("实时", "最大远端30分钟块数", max(step.aggregated_block_count for step in realtime_steps), "块", "参考"),
        ("输出", "正式期日期数", len(official), "天", "通过" if len(official) == OFFICIAL_DAYS else "失败"),
        ("输出", "正式期逐时段数", len(official) * N_PERIODS, "个", "通过" if len(official) * N_PERIODS == 48096 else "失败"),
        ("预测", "执行时净负荷诊断MAE", display_number(float(np.mean(np.abs(actual_net - diagnostic_net)))), "kWh/时段", "参考"),
        ("预测", "净负荷0.8分位数覆盖率", display_number(float(np.mean(actual_net <= diagnostic_net))), "比例", "参考"),
        ("情景", "净负荷情景包络覆盖率", display_number(float(np.mean((actual_net >= scenario_lower) & (actual_net <= scenario_upper)))), "比例", "参考"),
        ("约束", "最大实际电量平衡残差", display_number(float(np.max(np.abs(np.concatenate(actual_balance))))), "kWh", "通过"),
        ("约束", "最大实际状态转移残差", display_number(float(np.max(np.abs(np.concatenate(state_balance))))), "kWh", "通过"),
        ("费用", "费用拆账与绝对差额式最大残差", display_number(float(np.max(np.abs(fee_equivalence)))), "元", "通过"),
        ("费用", "1月初始化期实际总费", display_number(january_cost), "元", "参考"),
        ("费用", "2月1日至12月31日实际总费", display_number(official_metrics["total_cost"]), "元", "参考"),
        ("费用", "2025全年实际总费", display_number(yearly_cost), "元", "参考"),
        ("费用", "正式期初始计划费", display_number(official_metrics["initial_plan_cost"]), "元", "参考"),
        ("费用", "正式期增购费", display_number(official_metrics["increase_cost"]), "元", "参考"),
        ("费用", "正式期减购退费", display_number(official_metrics["decrease_refund"]), "元", "参考"),
        ("费用", "正式期减购违约费", display_number(official_metrics["decrease_penalty"]), "元", "参考"),
        ("费用", "正式期紧急购电费", display_number(official_metrics["emergency_cost"]), "元", "参考"),
        ("状态", "2025-01-01初始储电量", display_number(all_days[0].replay.stored_energy[0]), "kWh", "通过"),
        ("状态", "2025-02-01初始储电量", display_number(official[0].replay.stored_energy[0]), "kWh", "参考"),
        ("状态", "12月31日0点计划情景末态范围（诊断）", f"{np.min(final.decisions[0].stored_energy[:, -1]):.6f}~{np.max(final.decisions[0].stored_energy[:, -1]):.6f}" if final.decisions[0] else "", "kWh", "参考"),
        ("状态", "12月31日6点调整情景末态范围（诊断）", f"{np.min(final.decisions[1].stored_energy[:, -1]):.6f}~{np.max(final.decisions[1].stored_energy[:, -1]):.6f}" if final.decisions[1] else "", "kWh", "参考"),
        ("状态", "12月31日12点调整情景末态范围（诊断）", f"{np.min(final.decisions[2].stored_energy[:, -1]):.6f}~{np.max(final.decisions[2].stored_energy[:, -1]):.6f}" if final.decisions[2] else "", "kWh", "参考"),
        ("状态", "12月31日18点调整情景末态范围（诊断）", f"{np.min(final.decisions[3].stored_energy[:, -1]):.6f}~{np.max(final.decisions[3].stored_energy[:, -1]):.6f}" if final.decisions[3] else "", "kWh", "参考"),
        ("状态", "12月31日实际末态", display_number(final.replay.stored_energy[-1]), "kWh", "通过"),
        ("状态", "实际年末区间越界量", display_number(q2.terminal_interval_error(float(final.replay.stored_energy[-1]))), "kWh", "通过"),
    ]
    for item in all_days:
        rows.append((
            "参数选择",
            f"{item.day.isoformat()}当日参数",
            item.parameters.key,
            f"历史验证{item.parameter_validation_days}天",
            "预设" if item.parameter_validation_cost is None else f"滚动实际总费={display_number(item.parameter_validation_cost)}元",
        ))
    for record in result.tuning_records:
        rows.append((
            "费用调参",
            f"{record.validation_day.isoformat()}/{record.parameters.key}",
            display_number(record.realized_total_cost),
            "元",
            "完整日因果回放；次日起可用",
        ))
    return rows


def write_check_csv(data: ModelData, result: StrategyResult) -> Path:
    """写出主方案的模型校验。"""
    return write_csv(CHECK_FILE, ("类别", "校验项", "数值", "单位", "结论"), build_check_rows(data, result))


def result_by_date(results: Sequence[DailyResult]) -> dict[date, DailyResult]:
    mapping = {item.day: item for item in results}
    if len(mapping) != len(results):
        raise RuntimeError("正式结果中出现重复日期。")
    return mapping


def write_purchase_summary_csv(data: ModelData, results: Sequence[DailyResult]) -> Path:
    """按论文表1口径汇总四个指定日期的最终有效非紧急购电。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    for day in SPECIFIED_DATES:
        item = mapping[day]
        for hour in SPECIFIED_HOURS:
            t = hour * 6
            rows.append((day.isoformat(), "指定时段最终有效非紧急购电量", internal_interval_label(t), display_number(item.effective_purchase[t]), "kWh"))
        rows.extend((
            (day.isoformat(), "全天最终有效非紧急购电量", "00:00-24:00", display_number(float(np.sum(item.effective_purchase))), "kWh"),
            (day.isoformat(), "全天实际总购电费（含调整与紧急购电）", "00:00-24:00", display_number(item.realized_total_cost), "元"),
        ))
    return write_csv(PURCHASE_SUMMARY_FILE, ("日期", "指标", "真实时间段", "数值", "单位"), rows)


def write_storage_summary_csv(results: Sequence[DailyResult]) -> Path:
    """按论文表2口径汇总实际4小时充放电量及0/24点储电量。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    for day in SPECIFIED_DATES:
        item = mapping[day]
        for block_index, label in enumerate(STORAGE_BLOCK_LABELS):
            start = block_index * 24
            stop = start + 24
            rows.append((
                day.isoformat(), label,
                display_number(float(np.sum(item.replay.charge[start:stop]))),
                display_number(float(np.sum(item.replay.discharge[start:stop]))), "", "",
            ))
        rows.extend((
            (day.isoformat(), "", "", "", "00:00", display_number(item.replay.stored_energy[0])),
            (day.isoformat(), "", "", "", "24:00", display_number(item.replay.stored_energy[-1])),
        ))
    return write_csv(
        STORAGE_SUMMARY_FILE,
        ("日期", "真实时间段", "实际充电量（kWh）", "实际放电量（kWh）", "时刻", "实际储电量（kWh）"),
        rows,
    )


def write_emergency_summary_csv(results: Sequence[DailyResult]) -> Path:
    """按论文表3口径合并四个指定日期相邻的非零紧急购电时段。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    for day in SPECIFIED_DATES:
        intervals = q2.group_emergency_intervals(mapping[day].replay.emergency_purchase)
        if not intervals:
            rows.append((day.isoformat(), "无", display_number(0.0)))
        else:
            rows.extend((day.isoformat(), item.label, display_number(item.energy)) for item in intervals)
    return write_csv(EMERGENCY_SUMMARY_FILE, ("日期", "真实紧急购电时间段", "紧急购电量（kWh）"), rows)


def validate_template(workbook: object, official_results: Sequence[DailyResult]) -> None:
    """核对result3四张工作表、后移表头、日期及汇总列。"""
    if tuple(workbook.sheetnames) != RESULT_SHEET_NAMES:
        raise ValueError(f"result3.xlsx工作表应依次为：{RESULT_SHEET_NAMES}")
    # 复用Q2已验证的逐列表头解析，同时核对六个储能汇总段和紧急购电表。
    q2.validate_result_template(workbook, official_results)
    expected_dates = tuple(item.day for item in official_results)
    plan_headers = tuple(
        workbook["计划购电量"].cell(1, column).value
        for column in range(1, N_PERIODS + 4)
    )
    for sheet_name in ("计划购电量", "调整购电量"):
        sheet = workbook[sheet_name]
        if sheet.max_column != N_PERIODS + 3 or sheet.cell(1, 1).value != "日期\\时间":
            raise ValueError(f"{sheet_name}的列数或A1表头错误。")
        headers = tuple(sheet.cell(1, column).value for column in range(2, N_PERIODS + 2))
        if len(headers) != N_PERIODS or len(set(headers)) != N_PERIODS:
            raise ValueError(f"{sheet_name}必须含144个互不重复的显示区间。")
        if headers[0] != "0:10-0:20" or headers[-1] != "0:00-0:10+1":
            raise ValueError(f"{sheet_name}未保留整体后移10分钟的首末表头。")
        if tuple(
            sheet.cell(1, column).value for column in range(1, N_PERIODS + 4)
        ) != plan_headers:
            raise ValueError("计划购电量与调整购电量的全部表头必须完全一致。")
        if sheet.cell(1, N_PERIODS + 2).value != "全天购电量" or sheet.cell(1, N_PERIODS + 3).value != "全天购电费":
            raise ValueError(f"{sheet_name}缺少全天购电量或全天购电费列。")
        dates = tuple(
            q2.normalize_excel_date(sheet.cell(row, 1).value, f"{sheet_name}A{row}")
            for row in range(2, OFFICIAL_DAYS + 2)
        )
        if dates != expected_dates:
            raise ValueError(f"{sheet_name}日期必须从2025-02-01连续至2025-12-31。")
    if tuple(
        workbook["充放电量"].cell(1, column).value for column in range(1, 7)
    ) != ("日期", "时间段", "充电量", "放电量", "时刻", "储电量"):
        raise ValueError("充放电量工作表表头错误。")
    if tuple(
        workbook["紧急购电量"].cell(1, column).value for column in range(1, 4)
    ) != ("日期", "购电时间段", "购电量"):
        raise ValueError("紧急购电量工作表表头错误。")


def fill_purchase_sheet(worksheet: object, results: Sequence[DailyResult], adjusted: bool) -> None:
    """分别写入0点初始计划或最终有效非紧急购电。"""
    for row, item in enumerate(results, start=2):
        values = item.effective_purchase if adjusted else item.purchase_versions[0]
        for t, value in enumerate(values, start=2):
            cell = worksheet.cell(row, t)
            cell.value = excel_number(value)
            cell.number_format = "0.000000"
        total_energy = worksheet.cell(row, N_PERIODS + 2)
        total_cost = worksheet.cell(row, N_PERIODS + 3)
        total_energy.value = excel_number(float(np.sum(values)))
        total_cost.value = excel_number(item.realized_total_cost if adjusted else item.initial_plan_cost)
        total_energy.number_format = "0.000000"
        total_cost.number_format = "0.000000"


def write_result_workbook(result: StrategyResult) -> Path:
    """从只读模板生成临时文件，回读核对后原子发布result3.xlsx。"""
    official = result.official_days
    if len(official) != OFFICIAL_DAYS:
        raise RuntimeError("正式期不足334天，禁止发布result3.xlsx。")
    final_error = q2.terminal_interval_error(float(official[-1].replay.stored_energy[-1]))
    if final_error > TERMINAL_TOLERANCE:
        raise RuntimeError("实际年末状态未通过附加区间，禁止发布result3.xlsx。")
    workbook = load_workbook(RESULT_TEMPLATE_FILE)
    temporary: Path | None = None
    try:
        validate_template(workbook, official)
        fill_purchase_sheet(workbook["计划购电量"], official, adjusted=False)
        fill_purchase_sheet(workbook["调整购电量"], official, adjusted=True)
        q2.fill_storage_sheet(workbook["充放电量"], official)
        q2.fill_emergency_sheet(workbook["紧急购电量"], official)
        with tempfile.NamedTemporaryFile(dir=OUTPUT_RESULT_DIR, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        written = load_workbook(temporary, read_only=True, data_only=True)
        try:
            for sheet_name, adjusted in (("计划购电量", False), ("调整购电量", True)):
                sheet = written[sheet_name]
                if sheet.max_row != OFFICIAL_DAYS + 1:
                    raise RuntimeError(f"写出后{sheet_name}行数错误。")
                for values, item in zip(sheet.iter_rows(min_row=2, values_only=True), official):
                    array = item.effective_purchase if adjusted else item.purchase_versions[0]
                    expected = [excel_number(value) for value in array]
                    expected.extend((
                        excel_number(float(np.sum(array))),
                        excel_number(item.realized_total_cost if adjusted else item.initial_plan_cost),
                    ))
                    if q2.normalize_excel_date(values[0], "输出日期") != item.day or not np.allclose(
                        np.asarray(values[1:], dtype=float), expected, atol=1e-9, rtol=0.0
                    ):
                        raise RuntimeError(f"{item.day}的{sheet_name}回读校验失败。")
            if written["充放电量"].max_row != 1 + OFFICIAL_DAYS * 6:
                raise RuntimeError("写出后充放电量工作表行数错误。")
            emergency_count = sum(
                len(q2.group_emergency_intervals(item.replay.emergency_purchase))
                for item in official
            )
            if written["紧急购电量"].max_row != 1 + emergency_count:
                raise RuntimeError("写出后紧急购电量工作表行数错误。")
        finally:
            written.close()
        os.replace(temporary, RESULT_FILE)
    finally:
        workbook.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return RESULT_FILE


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="执行问题三净负荷情景购电/储能滚动调度。")
    parser.add_argument(
        "--smoke-days", type=int, metavar="N",
        help="仅连续运行前N天并做内存校验，不写正式结果文件。",
    )
    parser.add_argument(
        "--skip-comparisons", action="store_true",
        help="正式运行时只算主方案；策略对照CSV仅包含主方案（用于缩短调试时间）。",
    )
    return parser.parse_args()


def main() -> None:
    """完成数据校验、连续求解、输出校验及原子发布。"""
    if Path(sys.prefix).name != "2026C" or not (Path(sys.prefix) / "conda-meta").is_dir():
        raise RuntimeError("项目要求使用conda环境2026C。")
    args = parse_args()
    q2.validate_terminal_bounds()
    data = load_model_data()

    template = load_workbook(RESULT_TEMPLATE_FILE, read_only=True, data_only=False)
    try:
        placeholders = [SimpleNamespace(day=day) for day in data.base.dates[JANUARY_DAYS:]]
        validate_template(template, placeholders)
    finally:
        template.close()

    if args.smoke_days is not None:
        smoke = run_strategy(
            data, MAIN_STRATEGY,
            day_limit=args.smoke_days, show_progress=True
        )
        validate_strategy(data, smoke, complete_year=False)
        last = smoke.all_days[-1]
        print(
            f"冒烟测试通过：连续{len(smoke.all_days)}天，"
            f"共{len(smoke.all_days) * 4}次购电求解、"
            f"{len(smoke.all_days) * N_PERIODS}次实时储能求解；"
            f"{smoke.tuning_replay_count}次候选参数完整日回放；"
            f"末态{last.replay.stored_energy[-1]:.6f} kWh。"
        )
        return

    main_result = run_strategy(data, MAIN_STRATEGY, show_progress=True)
    validate_strategy(data, main_result, complete_year=True)
    strategy_results = [main_result]
    if not args.skip_comparisons:
        for spec in COMPARISON_STRATEGIES:
            comparison = run_strategy(data, spec, show_progress=True)
            validate_strategy(data, comparison, complete_year=True)
            strategy_results.append(comparison)

    OUTPUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    output_files = [
        write_decision_csv(main_result.all_days),
        write_detail_csv(data, main_result.all_days),
        write_check_csv(data, main_result),
        write_mapping_csv(),
        write_comparison_csv(strategy_results),
        write_purchase_summary_csv(data, main_result.official_days),
        write_storage_summary_csv(main_result.official_days),
        write_emergency_summary_csv(main_result.official_days),
        write_result_workbook(main_result),
    ]
    for path in output_files:
        print(f"已保存：{path.relative_to(PROJECT_DIR)}")


if __name__ == "__main__":
    main()
