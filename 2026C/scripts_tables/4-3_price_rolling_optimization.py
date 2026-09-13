"""问题四（问题三部分）：波动电价下的四次购电与储能滚动调度。

在 conda 环境 2026C 中运行：
    python scripts_tables/4-3_price_rolling_optimization.py

脚本复用问题三的净负荷情景、四次购电、实时储能控制和结果模板逻辑，
只增加问题四所需的因果价格预测、当期实际价格控制和实际价格结算。
默认运行主方案、仅0点和仅18点三种策略；冒烟测试可运行：
    python scripts_tables/4-3_price_rolling_optimization.py --smoke-days 2
"""

from __future__ import annotations

import argparse
import importlib
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

q3 = importlib.import_module("3_rolling_optimization")
q42 = importlib.import_module("4-2_price_rolling_optimization")

ATTACHMENT_DIR = PROJECT_DIR / "CUMCM 2026 C题" / "附件"
RESULT_TEMPLATE_FILE = ATTACHMENT_DIR / "附件5" / "result4-3.xlsx"
OUTPUT_TABLE_DIR = PROJECT_DIR / "outputs" / "tables"
OUTPUT_RESULT_DIR = PROJECT_DIR / "outputs" / "results"
RESULT_FILE = OUTPUT_RESULT_DIR / "result4-3.xlsx"

DETAIL_FILE = OUTPUT_TABLE_DIR / "4-3_实际调度与费用明细.csv"
PRICE_FORECAST_FILE = OUTPUT_TABLE_DIR / "4-3_电价预测与误差.csv"
CHECK_FILE = OUTPUT_TABLE_DIR / "4-3_模型校验.csv"
MAPPING_FILE = OUTPUT_TABLE_DIR / "4-3_输出时段映射.csv"
COMPARISON_FILE = OUTPUT_TABLE_DIR / "4-3_预报更新策略对照.csv"
PURCHASE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "4-3_指定日期购电量及全天结果.csv"
STORAGE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "4-3_指定日期充放电量及储电量.csv"
EMERGENCY_SUMMARY_FILE = OUTPUT_TABLE_DIR / "4-3_指定日期紧急购电量.csv"

N_PERIODS = q3.N_PERIODS
JANUARY_DAYS = q3.JANUARY_DAYS
OFFICIAL_DAYS = q3.OFFICIAL_DAYS
E_INITIAL = q3.E_INITIAL
E_MIN = q3.E_MIN
E_MAX = q3.E_MAX
E_TERMINAL_MIN = q3.E_TERMINAL_MIN
E_TERMINAL_MAX = q3.E_TERMINAL_MAX
ETA_CHARGE = q3.ETA_CHARGE
ETA_DISCHARGE = q3.ETA_DISCHARGE
EMERGENCY_PRICE_MULTIPLIER = q3.EMERGENCY_PRICE_MULTIPLIER
BALANCE_TOLERANCE = q3.BALANCE_TOLERANCE
BOUND_TOLERANCE = q3.BOUND_TOLERANCE
TERMINAL_TOLERANCE = q3.TERMINAL_TOLERANCE
RELEASE_PERIODS = q3.RELEASE_PERIODS
PERIODS_PER_BLOCK = q3.PERIODS_PER_BLOCK


@dataclass(frozen=True)
class PriceForecastAudit:
    """某日0点价格预测的因果来源和误差审计。"""

    day: date
    candidate_index: int
    candidate_label: str
    latest_validation_date: date | None
    validation_days: int
    validation_score: float | None
    predicted_price: np.ndarray
    actual_price: np.ndarray


@dataclass(frozen=True)
class ModelData:
    """问题三输入与问题四实际/预测价格的对齐结果。"""

    base: object
    photovoltaic_forecast_power: np.ndarray
    photovoltaic_feature_energy: np.ndarray
    actual_net_load_energy: np.ndarray
    actual_price: np.ndarray
    predicted_price: np.ndarray
    price_audits: tuple[PriceForecastAudit, ...]


def build_price_forecasts(price_data: object) -> tuple[np.ndarray, tuple[PriceForecastAudit, ...]]:
    """完全复用Q4-2口径，只用计划日前的实际价格预测当天144个时段。"""
    losses: list[np.ndarray] = []
    forecasts: list[np.ndarray] = []
    audits: list[PriceForecastAudit] = []
    labels = q42.price_candidate_labels()
    for day_index, day in enumerate(price_data.dates):
        scores = q42.historical_scores(
            losses,
            day_index,
            q42.PRICE_CANDIDATE_COUNT,
            price_data.dates,
        )
        candidate_index = q42.choose_candidate(
            scores, q42.DEFAULT_PRICE_CANDIDATE_INDEX
        )
        candidates = q42.price_candidates(
            day,
            price_data.dates[:day_index],
            price_data.actual_price[:day_index],
            price_data.reference_price,
        )
        selected = np.asarray(candidates[candidate_index], dtype=float).copy()
        validation_indices = q42.validation_indices(
            losses, day_index, price_data.dates
        )
        latest_validation_date = (
            price_data.dates[validation_indices[-1]] if validation_indices else None
        )
        validation_score = (
            float(scores[candidate_index]) if validation_indices else None
        )
        actual = np.asarray(price_data.actual_price[day_index], dtype=float)
        audits.append(
            PriceForecastAudit(
                day=day,
                candidate_index=candidate_index,
                candidate_label=labels[candidate_index],
                latest_validation_date=latest_validation_date,
                validation_days=len(validation_indices),
                validation_score=validation_score,
                predicted_price=selected,
                actual_price=actual.copy(),
            )
        )
        forecasts.append(selected)
        losses.append(np.mean(np.abs(candidates - actual[None, :]), axis=1))
    predicted = np.stack(forecasts)
    if predicted.shape != price_data.actual_price.shape:
        raise RuntimeError("价格预测矩阵与附件4实际电价没有完全对齐。")
    if not np.isfinite(predicted).all() or np.any(predicted <= 0):
        raise RuntimeError("价格预测必须全部为有限正数。")
    if any(
        item.latest_validation_date is not None
        and item.latest_validation_date >= item.day
        for item in audits
    ):
        raise RuntimeError("价格预测参数选择使用了当日或未来数据。")
    return predicted, tuple(audits)


def load_model_data() -> ModelData:
    """读取并交叉核对附件1至4，建立Q4-3统一输入。"""
    q3_data = q3.load_model_data()
    price_data = q42.load_model_data()
    if tuple(q3_data.base.dates) != tuple(price_data.dates):
        raise ValueError("问题三与附件4的日期没有完全对齐。")
    for name, q3_values, q42_values in (
        ("实际负载", q3_data.base.actual_load_energy, price_data.actual_load_energy),
        (
            "实际光伏",
            q3_data.base.actual_photovoltaic_energy,
            price_data.actual_photovoltaic_energy,
        ),
    ):
        if not np.allclose(q3_values, q42_values, atol=0.0, rtol=0.0):
            raise ValueError(f"Q3与Q4-2读取的{name}数据不一致。")
    predicted_price, audits = build_price_forecasts(price_data)
    return ModelData(
        base=q3_data.base,
        photovoltaic_forecast_power=q3_data.photovoltaic_forecast_power,
        photovoltaic_feature_energy=q3_data.photovoltaic_feature_energy,
        actual_net_load_energy=q3_data.actual_net_load_energy,
        actual_price=np.asarray(price_data.actual_price, dtype=float).copy(),
        predicted_price=predicted_price,
        price_audits=audits,
    )


def simulate_day(
    data: ModelData,
    day_index: int,
    initial_energy: float,
    parameters: q3.ScenarioParameters,
    spec: q3.StrategySpec,
    parameter_latest_validation_date: date | None,
    parameter_validation_days: int,
    parameter_validation_cost: float | None,
) -> q3.DailyResult:
    """以0点冻结价格预测规划，以逐时段实际价格控制并结算一天。"""
    day = data.base.dates[day_index]
    predicted_price = np.asarray(data.predicted_price[day_index], dtype=float)
    actual_price = np.asarray(data.actual_price[day_index], dtype=float)
    terminal_penalty = float(np.max(predicted_price))
    actual_load = data.base.actual_load_energy[day_index]
    actual_pv = data.base.actual_photovoltaic_energy[day_index]
    actual_net = data.actual_net_load_energy[day_index]

    scenario_versions: list[q3.NetLoadScenarios] = []
    decisions: list[q3.PurchaseDecision | None] = []
    versions = np.empty((4, N_PERIODS), dtype=float)
    effective = np.empty(N_PERIODS, dtype=float)
    replay_charge = np.zeros(N_PERIODS)
    replay_discharge = np.zeros(N_PERIODS)
    replay_emergency = np.zeros(N_PERIODS)
    replay_curtailment = np.zeros(N_PERIODS)
    replay_unused = np.zeros(N_PERIODS)
    replay_steps: list[q3.RealtimeStep] = []
    stored = np.empty(N_PERIODS + 1)
    stored[0] = initial_energy
    accepted_batch = 0

    scenarios0 = q3.build_net_load_scenarios(data, day_index, 0, 0, parameters)
    decision0 = q3.solve_purchase_decision(
        day,
        0,
        predicted_price,
        scenarios0,
        initial_energy,
        terminal_penalty,
        None,
        False,
    )
    scenario_versions.append(scenarios0)
    decisions.append(decision0)
    versions[0] = decision0.grid_purchase
    initial_plan = versions[0].copy()
    latest_complete_decision = decision0
    latest_complete_scenarios = scenarios0

    for decision_index in range(4):
        start = RELEASE_PERIODS[decision_index]
        stop = start + PERIODS_PER_BLOCK
        if decision_index > 0:
            if decision_index in spec.forecast_updates:
                accepted_batch = decision_index
            scenarios = q3.build_net_load_scenarios(
                data, day_index, decision_index, accepted_batch, parameters
            )
            scenario_versions.append(scenarios)
            versions[decision_index] = versions[decision_index - 1]
            if decision_index in spec.purchase_updates:
                decision = q3.solve_purchase_decision(
                    day,
                    decision_index,
                    predicted_price,
                    scenarios,
                    float(stored[start]),
                    terminal_penalty,
                    initial_plan,
                    False,
                )
                versions[decision_index, start:] = decision.grid_purchase
                decisions.append(decision)
                latest_complete_decision = decision
                latest_complete_scenarios = scenarios
            else:
                decisions.append(None)
        scenarios = scenario_versions[decision_index]
        version = versions[decision_index]
        for period_index in range(start, stop):
            posterior = q3.update_realtime_weights(
                scenarios, actual_net, period_index
            )
            boundary_period = q3.build_realtime_blocks(period_index)[-1][-1] + 1
            target_offset = boundary_period - latest_complete_decision.start_period
            if not 0 <= target_offset < latest_complete_decision.stored_energy.shape[1]:
                raise RuntimeError(
                    f"{day}第{period_index + 1}时段无法取得实时边界状态目标。"
                )
            boundary_target = float(
                latest_complete_decision.stored_energy[:, target_offset]
                @ latest_complete_scenarios.weights
            )
            # 购电先提交；随后仅将当前时段价格替换为已揭晓实际值，未来仍用0点预测。
            realtime_price = predicted_price.copy()
            realtime_price[period_index] = actual_price[period_index]
            step = q3.solve_realtime_scenario_step(
                period_index=period_index,
                price=realtime_price,
                purchase_version=version,
                scenarios=scenarios,
                scenario_weights=posterior,
                actual_load_now=float(actual_load[period_index]),
                actual_photovoltaic_now=float(actual_pv[period_index]),
                current_energy=float(stored[period_index]),
                boundary_target_energy=boundary_target,
                terminal_penalty=terminal_penalty,
            )
            replay_steps.append(step)
            effective[period_index] = version[period_index]
            replay_charge[period_index] = step.charge
            replay_discharge[period_index] = step.discharge
            replay_emergency[period_index] = step.emergency_purchase
            replay_curtailment[period_index] = step.curtailment
            replay_unused[period_index] = step.unused_planned_purchase
            stored[period_index + 1] = (
                stored[period_index]
                + ETA_CHARGE * step.charge
                - step.discharge / ETA_DISCHARGE
            )

    replay = q3.ReplayResult(
        charge=replay_charge,
        discharge=replay_discharge,
        emergency_purchase=replay_emergency,
        curtailment=replay_curtailment,
        unused_planned_purchase=replay_unused,
        stored_energy=stored,
        steps=tuple(replay_steps),
    )
    increase = np.maximum(effective - initial_plan, 0.0)
    decrease = np.maximum(initial_plan - effective, 0.0)
    initial_cost = float(actual_price @ initial_plan)
    increase_cost = float(1.5 * actual_price @ increase)
    decrease_refund = float(actual_price @ decrease)
    decrease_penalty = float(0.5 * actual_price @ decrease)
    emergency_cost = float(
        EMERGENCY_PRICE_MULTIPLIER * actual_price @ replay_emergency
    )
    daily = q3.DailyResult(
        day=day,
        parameters=parameters,
        parameter_latest_validation_date=parameter_latest_validation_date,
        parameter_validation_days=parameter_validation_days,
        parameter_validation_cost=parameter_validation_cost,
        scenarios=tuple(scenario_versions),
        decisions=tuple(decisions),
        purchase_versions=versions,
        effective_purchase=effective,
        actual_load_energy=actual_load.copy(),
        actual_photovoltaic_energy=actual_pv.copy(),
        replay=replay,
        initial_plan_cost=initial_cost,
        increase_purchase_cost=increase_cost,
        decrease_refund=decrease_refund,
        decrease_penalty=decrease_penalty,
        emergency_purchase_cost=emergency_cost,
        realized_total_cost=(
            initial_cost
            + increase_cost
            - decrease_refund
            + decrease_penalty
            + emergency_cost
        ),
    )
    # Q3的日校验只通过data.base.price读取结算价格；这里按当天实际价注入。
    validation_data = SimpleNamespace(base=SimpleNamespace(price=actual_price))
    q3.validate_daily_result(
        validation_data, daily, float(initial_energy), spec
    )
    return daily


def write_detail_csv(data: ModelData, results: Sequence[q3.DailyResult]) -> Path:
    """写出足以复算实际物理量与波动价格账单的逐时段明细。"""
    display_headers = q3.template_headers()
    day_index = {day: index for index, day in enumerate(data.base.dates)}
    header = (
        "日期",
        "阶段",
        "内部时段",
        "内部真实区间",
        "模板显示区间",
        "生效购电版本",
        "0点预测电价",
        "实际结算电价",
        "实际负载电量",
        "实际光伏电量",
        "实际净负荷电量",
        "0点初始计划",
        "最终有效购电",
        "增购差额",
        "减购差额",
        "实际充电量",
        "实际放电量",
        "紧急购电量",
        "弃光量",
        "未利用有效购电量",
        "时段初储电量",
        "时段末储电量",
        "初始计划费",
        "增购费",
        "减购退费",
        "减购违约费",
        "紧急购电费",
        "时段实际总费",
    )

    def rows() -> Iterable[Sequence[object]]:
        for result in results:
            d = day_index[result.day]
            predicted = data.predicted_price[d]
            actual = data.actual_price[d]
            g0 = result.purchase_versions[0]
            for t in range(N_PERIODS):
                k = t // PERIODS_PER_BLOCK
                effective = result.effective_purchase[t]
                increase = max(effective - g0[t], 0.0)
                decrease = max(g0[t] - effective, 0.0)
                initial_fee = actual[t] * g0[t]
                increase_fee = 1.5 * actual[t] * increase
                refund = actual[t] * decrease
                penalty = 0.5 * actual[t] * decrease
                emergency_fee = (
                    EMERGENCY_PRICE_MULTIPLIER
                    * actual[t]
                    * result.replay.emergency_purchase[t]
                )
                yield (
                    result.day.isoformat(),
                    "初始化" if result.day < q3.OFFICIAL_START else "正式期",
                    t + 1,
                    q3.internal_interval_label(t),
                    display_headers[t],
                    f"{q3.RELEASE_HOURS[k]:02d}:00版本",
                    q3.display_number(predicted[t]),
                    q3.display_number(actual[t]),
                    q3.display_number(result.actual_load_energy[t]),
                    q3.display_number(result.actual_photovoltaic_energy[t]),
                    q3.display_number(
                        result.actual_load_energy[t]
                        - result.actual_photovoltaic_energy[t]
                    ),
                    q3.display_number(g0[t]),
                    q3.display_number(effective),
                    q3.display_number(increase),
                    q3.display_number(decrease),
                    q3.display_number(result.replay.charge[t]),
                    q3.display_number(result.replay.discharge[t]),
                    q3.display_number(result.replay.emergency_purchase[t]),
                    q3.display_number(result.replay.curtailment[t]),
                    q3.display_number(result.replay.unused_planned_purchase[t]),
                    q3.display_number(result.replay.stored_energy[t]),
                    q3.display_number(result.replay.stored_energy[t + 1]),
                    q3.display_number(initial_fee),
                    q3.display_number(increase_fee),
                    q3.display_number(refund),
                    q3.display_number(penalty),
                    q3.display_number(emergency_fee),
                    q3.display_number(
                        initial_fee
                        + increase_fee
                        - refund
                        + penalty
                        + emergency_fee
                    ),
                )
    return q3.write_csv(DETAIL_FILE, header, rows())


def write_price_forecast_csv(data: ModelData) -> Path:
    """按日写出价格预测参数来源和误差，避免重复保存52560个时段。"""
    rows = []
    for item in data.price_audits:
        error = item.predicted_price - item.actual_price
        rows.append(
            (
                item.day.isoformat(),
                "初始化" if item.day < q3.OFFICIAL_START else "正式期",
                item.candidate_label,
                item.latest_validation_date or "附件1冷启动",
                item.validation_days,
                ""
                if item.validation_score is None
                else q3.display_number(item.validation_score),
                q3.display_number(float(np.mean(item.predicted_price))),
                q3.display_number(float(np.mean(item.actual_price))),
                q3.display_number(float(np.mean(np.abs(error)))),
                q3.display_number(float(np.sqrt(np.mean(error**2)))),
                q3.display_number(float(np.max(item.predicted_price))),
                q3.display_number(float(np.max(item.actual_price))),
            )
        )
    return q3.write_csv(
        PRICE_FORECAST_FILE,
        (
            "日期",
            "阶段",
            "价格候选",
            "最晚验证日期",
            "验证天数",
            "历史加权MAE",
            "预测日均价",
            "实际日均价",
            "当日MAE",
            "当日RMSE",
            "预测最高价",
            "实际最高价",
        ),
        rows,
    )


def build_check_rows(
    data: ModelData, result: q3.StrategyResult
) -> list[Sequence[object]]:
    """构造价格因果性、实际物理约束、费用和终端状态的核心校验。"""
    all_days = result.all_days
    official = result.official_days
    day_index = {day: index for index, day in enumerate(data.base.dates)}
    balances = []
    state_balances = []
    fee_residuals = []
    for item in all_days:
        actual_price = data.actual_price[day_index[item.day]]
        balances.append(
            item.effective_purchase
            + item.actual_photovoltaic_energy
            + item.replay.discharge
            + item.replay.emergency_purchase
            - item.replay.curtailment
            - item.replay.unused_planned_purchase
            - item.actual_load_energy
            - item.replay.charge
        )
        state_balances.append(
            item.replay.stored_energy[1:]
            - item.replay.stored_energy[:-1]
            - ETA_CHARGE * item.replay.charge
            + item.replay.discharge / ETA_DISCHARGE
        )
        g0 = item.purchase_versions[0]
        absolute_cost = float(
            actual_price @ item.effective_purchase
            + 0.5 * actual_price @ np.abs(item.effective_purchase - g0)
            + EMERGENCY_PRICE_MULTIPLIER
            * actual_price
            @ item.replay.emergency_purchase
        )
        fee_residuals.append(absolute_cost - item.realized_total_cost)

    official_indices = [day_index[item.day] for item in official]
    price_error = (
        data.predicted_price[official_indices]
        - data.actual_price[official_indices]
    )
    metrics = q3.strategy_metrics(result)
    january_cost = sum(item.realized_total_cost for item in result.warmup_days)
    planning_decisions = [
        decision
        for item in all_days
        for decision in item.decisions
        if decision is not None
    ]
    realtime_steps = [step for item in all_days for step in item.replay.steps]
    final_energy = float(all_days[-1].replay.stored_energy[-1])
    terminal_error = q3.q2.terminal_interval_error(final_energy)
    price_latest_ok = all(
        item.latest_validation_date is None
        or item.latest_validation_date < item.day
        for item in data.price_audits
    )
    rows: list[Sequence[object]] = [
        ("输入", "附件4实际电价矩阵", f"{data.actual_price.shape[0]}×{data.actual_price.shape[1]}", "天×时段", "通过"),
        ("输入", "实际电价最小值", q3.display_number(float(np.min(data.actual_price))), "元/kWh", "通过"),
        ("因果性", "价格调参最晚日期早于计划日", int(price_latest_ok), "布尔值", "通过" if price_latest_ok else "失败"),
        ("因果性", "购电决策价格口径", "0点价格预测全天冻结", "", "通过"),
        ("因果性", "实时价格口径", "当前实际价＋未来0点预测价", "", "通过"),
        ("预测", "正式期价格MAE", q3.display_number(float(np.mean(np.abs(price_error)))), "元/kWh", "参考"),
        ("预测", "正式期价格RMSE", q3.display_number(float(np.sqrt(np.mean(price_error**2)))), "元/kWh", "参考"),
        ("求解", "购电求解次数", len(planning_decisions), "次", "通过" if len(planning_decisions) == 1460 else "失败"),
        ("求解", "松弛后固定模式LP次数", sum(item.solver_kind == "松弛定模式后固定模式LP" for item in planning_decisions), "次", "通过"),
        ("求解", "购电MILP回退次数", sum(item.solver_kind == "MILP回退" for item in planning_decisions), "次", "参考"),
        ("求解", "实时储能求解次数", len(realtime_steps), "次", "通过" if len(realtime_steps) == 365 * N_PERIODS else "失败"),
        ("约束", "最大实际电量平衡残差", q3.display_number(float(np.max(np.abs(np.concatenate(balances))))), "kWh", "通过"),
        ("约束", "最大实际状态转移残差", q3.display_number(float(np.max(np.abs(np.concatenate(state_balances))))), "kWh", "通过"),
        ("费用", "费用等价式最大残差", q3.display_number(float(np.max(np.abs(fee_residuals)))), "元", "通过"),
        ("输出", "正式期日期数", len(official), "天", "通过" if len(official) == OFFICIAL_DAYS else "失败"),
        ("输出", "正式期逐时段数", len(official) * N_PERIODS, "个", "通过" if len(official) * N_PERIODS == OFFICIAL_DAYS * N_PERIODS else "失败"),
        ("费用", "1月初始化期实际总费", q3.display_number(january_cost), "元", "参考"),
        ("费用", "正式期初始计划费", q3.display_number(metrics["initial_plan_cost"]), "元", "参考"),
        ("费用", "正式期增购费", q3.display_number(metrics["increase_cost"]), "元", "参考"),
        ("费用", "正式期减购退费", q3.display_number(metrics["decrease_refund"]), "元", "参考"),
        ("费用", "正式期减购违约费", q3.display_number(metrics["decrease_penalty"]), "元", "参考"),
        ("费用", "正式期紧急购电费", q3.display_number(metrics["emergency_cost"]), "元", "参考"),
        ("费用", "2月1日至12月31日实际总费", q3.display_number(metrics["total_cost"]), "元", "参考"),
        ("状态", "2025-01-01初始储电量", q3.display_number(all_days[0].replay.stored_energy[0]), "kWh", "通过"),
        ("状态", "2025-02-01初始储电量", q3.display_number(official[0].replay.stored_energy[0]), "kWh", "参考"),
        ("状态", "12月31日实际末态", q3.display_number(final_energy), "kWh", "通过" if terminal_error <= TERMINAL_TOLERANCE else "失败"),
        ("状态", "实际年末区间越界量", q3.display_number(terminal_error), "kWh", "通过" if terminal_error <= TERMINAL_TOLERANCE else "失败"),
    ]
    for item in data.price_audits:
        rows.append(
            (
                "价格参数选择",
                f"{item.day.isoformat()}价格候选",
                item.candidate_label,
                f"历史验证{item.validation_days}天",
                "预设" if item.validation_score is None else f"加权MAE={q3.display_number(item.validation_score)}",
            )
        )
    return rows


def write_check_csv(data: ModelData, result: q3.StrategyResult) -> Path:
    return q3.write_csv(
        CHECK_FILE,
        ("类别", "校验项", "数值", "单位", "结论"),
        build_check_rows(data, result),
    )


def configure_q3_outputs() -> None:
    """把Q3通用写表器定向到Q4-3模板和文件名。"""
    q3.RESULT_TEMPLATE_FILE = RESULT_TEMPLATE_FILE
    q3.RESULT_FILE = RESULT_FILE
    q3.MAPPING_FILE = MAPPING_FILE
    q3.COMPARISON_FILE = COMPARISON_FILE
    q3.PURCHASE_SUMMARY_FILE = PURCHASE_SUMMARY_FILE
    q3.STORAGE_SUMMARY_FILE = STORAGE_SUMMARY_FILE
    q3.EMERGENCY_SUMMARY_FILE = EMERGENCY_SUMMARY_FILE


def selected_strategies() -> tuple[q3.StrategySpec, ...]:
    """只保留支撑最终结论的主方案、仅0点和仅18点。"""
    comparisons = {item.key: item for item in q3.COMPARISON_STRATEGIES}
    return (
        q3.MAIN_STRATEGY,
        comparisons["zero_only"],
        comparisons["single_18"],
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="执行问题四（问题三部分）波动电价滚动调度。"
    )
    parser.add_argument(
        "--smoke-days",
        type=int,
        metavar="N",
        help="仅连续运行前N天并做内存校验，不写正式结果文件。",
    )
    parser.add_argument(
        "--main-only",
        action="store_true",
        help="正式运行时只算主方案，跳过仅0点和仅18点对照。",
    )
    return parser.parse_args()


def main() -> None:
    """完成数据校验、三种策略连续求解、输出校验和原子发布。"""
    if Path(sys.prefix).name != "2026C" or not (Path(sys.prefix) / "conda-meta").is_dir():
        raise RuntimeError("项目要求使用conda环境2026C。")
    args = parse_args()
    configure_q3_outputs()
    q3.q2.validate_terminal_bounds()
    data = load_model_data()
    # q3.run_strategy在其模块内查找simulate_day；只在本进程中替换为Q4-3版本。
    q3.simulate_day = simulate_day

    # 模板只有数百行；普通模式可直接随机访问单元格，避免只读流式模式在
    # validate_template 中为每次 cell() 调用重复扫描整张工作表。
    template = load_workbook(RESULT_TEMPLATE_FILE, read_only=False, data_only=False)
    try:
        placeholders = [
            SimpleNamespace(day=day)
            for day in data.base.dates[JANUARY_DAYS:]
        ]
        q3.validate_template(template, placeholders)
    finally:
        template.close()

    if args.smoke_days is not None:
        if not 1 <= args.smoke_days <= len(data.base.dates):
            raise ValueError("冒烟测试天数必须在1至365之间。")
        smoke_days: list[q3.DailyResult] = []
        initial_energy = E_INITIAL
        for day_index in range(args.smoke_days):
            daily = simulate_day(
                data,
                day_index,
                initial_energy,
                q3.DEFAULT_PARAMETERS,
                q3.MAIN_STRATEGY,
                None,
                0,
                None,
            )
            smoke_days.append(daily)
            initial_energy = float(daily.replay.stored_energy[-1])
            print(
                f"[冒烟测试] 已完成{day_index + 1}/{args.smoke_days}天；"
                f"末态={initial_energy:.3f} kWh",
                flush=True,
            )
        smoke = q3.StrategyResult(
            spec=q3.MAIN_STRATEGY,
            warmup_days=tuple(smoke_days),
            official_days=(),
            tuning_replay_count=0,
            tuning_records=(),
        )
        q3.validate_strategy(data, smoke, complete_year=False)
        print(
            f"冒烟测试通过：连续{len(smoke.all_days)}天；"
            f"末态{smoke.all_days[-1].replay.stored_energy[-1]:.6f} kWh。"
        )
        return

    specs = selected_strategies()[:1] if args.main_only else selected_strategies()
    strategy_results = []
    for spec in specs:
        strategy_result = q3.run_strategy(data, spec, show_progress=True)
        q3.validate_strategy(data, strategy_result, complete_year=True)
        strategy_results.append(strategy_result)
    main_result = strategy_results[0]

    OUTPUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    output_files = [
        write_detail_csv(data, main_result.all_days),
        write_price_forecast_csv(data),
        write_check_csv(data, main_result),
        q3.write_mapping_csv(),
        q3.write_comparison_csv(strategy_results),
        q3.write_purchase_summary_csv(data, main_result.official_days),
        q3.write_storage_summary_csv(main_result.official_days),
        q3.write_emergency_summary_csv(main_result.official_days),
        q3.write_result_workbook(main_result),
    ]
    for path in output_files:
        print(f"已保存：{path.relative_to(PROJECT_DIR)}")


if __name__ == "__main__":
    main()
