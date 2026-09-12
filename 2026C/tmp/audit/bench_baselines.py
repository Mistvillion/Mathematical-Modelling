"""Baselines for judging the Q2 solution quality.

A) no-storage rolling policy (same causal forecast, battery switched off).
B) full-period perfect-foresight LP lower bound (Feb 1 - Dec 31, actual data,
   same initial SOC and same year-end band).
C) energy-balance accounting for the delivered plan.
"""

from __future__ import annotations

import csv

import numpy as np
import openpyxl
from scipy.optimize import Bounds, LinearConstraint, linprog
from scipy.sparse import lil_matrix

DT = 1 / 6
E_MIN, E_MAX = 1200.0, 10800.0
ETA_C = ETA_D = 0.9
LIM = 5000 / 6
N = 144
PEN = 5.0


def read_actual(sheet):
    wb = openpyxl.load_workbook(
        "CUMCM 2026 C题/附件/附件2.xlsx", read_only=True, data_only=True
    )
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    dates = [r[0].date().isoformat() for r in rows[1:]]
    return dates, np.array([[float(v) for v in r[1:]] for r in rows[1:]]) * DT


dates, load = read_actual("小区负载")
_, pv = read_actual("光伏发电实际功率")
pv[pv < DT] = 0.0
wb = openpyxl.load_workbook(
    "CUMCM 2026 C题/附件/附件1.xlsx", read_only=True, data_only=True
)
price = np.array([float(r[1]) for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:]])
wb.close()

OFF = dates.index("2025-02-01")
with open("outputs/tables/2_滚动预测与调度明细.csv", encoding="utf-8-sig") as fh:
    det = list(csv.DictReader(fh))
day_seq = [det[k * N]["日期"] for k in range(334)]
sel = np.array([dates.index(d) for d in day_seq])
load_o = np.concatenate([load[i] for i in sel])
pv_o = np.concatenate([pv[i] for i in sel])
p_o = np.tile(price, 334)
print("official-period load =", load_o.sum(), "kWh")
print("official-period PV   =", pv_o.sum(), "kWh")
print("official-period net  =", (load_o - pv_o).sum(), "kWh")

def col(name):
    return np.array([float(r[name]) for r in det])

g = col("计划购电量（kWh）")
c = col("实际充电量（kWh）")
d = col("实际放电量（kWh）")
h = col("紧急购电量（kWh）")
r_ = col("实际弃光量（kWh）")
w = col("未利用计划购电量（kWh）")
e0 = col("时段初储电量（kWh）")
e1 = col("时段末储电量（kWh）")
lp = col("预测负载电量（kWh）")
vp = col("预测光伏电量（kWh）")
print("\n--- energy accounting for delivered plan (official period) ---")
losses = (1 - ETA_C) * c.sum() + (1 / ETA_D - 1) * d.sum()
dsoc = e1[-1] - e0[0]
print(f"purchase g      = {g.sum():15.3f}")
print(f"actual pv       = {pv_o.sum():15.3f}")
print(f"emergency h     = {h.sum():15.3f}")
print(f"curtail r       = {r_.sum():15.3f}")
print(f"unused w        = {w.sum():15.3f}")
print(f"actual load     = {load_o.sum():15.3f}")
print(f"storage losses  = {losses:15.3f}")
print(f"delta SOC       = {dsoc:15.3f}")
print("residual g+pv+h-r-w-load-losses-dsoc =",
      g.sum() + pv_o.sum() + h.sum() - r_.sum() - w_.sum() if (w_ := w.sum()) else 0.0,
      "-", load_o.sum() + losses + dsoc)

# ---------------- A) no-storage rolling policy ----------------
# day-ahead plan from the same published forecast, g = max(0, fhat_load - fhat_pv)
gp = np.maximum(lp - vp, 0.0)
h2 = np.maximum(load_o - gp - pv_o, 0.0)
cost_a = float(p_o @ gp) + PEN * float(p_o @ h2)
print("\n--- A) no-storage policy, same causal forecast ---")
print("plan cost      =", float(p_o @ gp))
print("emergency kWh  =", h2.sum(), "| cost =", PEN * float(p_o @ h2))
print("total 334d     =", cost_a)
print("vs delivered   =", cost_a - 14272696.316983)

# ---------------- B) perfect-foresight full-period LP ----------------
nD, nT = 334, N
T = nD * nT
# vars: g(T) c(T) d(T) r(T) e(T+1)
o = dict(g=0, c=T, d=2 * T, r=3 * T, e=4 * T)
nvar = 5 * T + 1
cvec = np.zeros(nvar)
cvec[o["g"]:o["g"] + T] = p_o
lb = np.zeros(nvar)
ub = np.full(nvar, np.inf)
ub[o["c"]:o["c"] + T] = LIM
ub[o["d"]:o["d"] + T] = LIM
ub[o["r"]:o["r"] + T] = pv_o
lb[o["e"]:o["e"] + T + 1] = E_MIN
ub[o["e"]:o["e"] + T + 1] = E_MAX
lb[o["e"]] = ub[o["e"]] = 6666.471759  # same Feb-1 initial SOC as delivered run
lb[o["e"] + T] = 5400.0
ub[o["e"] + T] = 6600.0

rows = 2 * T
A = lil_matrix((rows, nvar))
lo = np.full(rows, -np.inf)
hi = np.full(rows, np.inf)
for t in range(T):
    A[t, o["g"] + t] = 1
    A[t, o["c"] + t] = -1
    A[t, o["d"] + t] = 1
    A[t, o["r"] + t] = -1
    lo[t] = hi[t] = load_o[t] - pv_o[t]
    row = T + t
    A[row, o["c"] + t] = -ETA_C
    A[row, o["d"] + t] = 1 / ETA_D
    A[row, o["e"] + t] = -1
    A[row, o["e"] + t + 1] = 1
    lo[row] = hi[row] = 0.0
res = linprog(cvec, A_eq=A.tocsr(), b_eq=lo, bounds=np.column_stack((lb, ub)),
              method="highs")
if not res.success:
    raise SystemExit(f"LP failed: {res.message}")
print("\n--- B) perfect-foresight LP lower bound (Feb-Dec, same endpoints) ---")
print(res.message)
print("objective =", res.fun, " 元")
x = res.x
gg = x[o["g"]:o["g"] + T]
cc = x[o["c"]:o["c"] + T]
dd = x[o["d"]:o["d"] + T]
rr = x[o["r"]:o["r"] + T]
print("purchase kWh     =", gg.sum())
print("curtailment kWh  =", rr.sum())
print("simultaneous ch/dis max =", float(np.minimum(cc, dd).max()))
print("end SOC          =", x[o["e"] + T], "| check sum p*g =", float(p_o @ gg))
print("\nDelivered total cost 14272696.32 vs perfect-foresight", res.fun,
      f"-> gap = {14272696.316983 - res.fun:.2f} 元 "
      f"({100 * (14272696.316983 - res.fun) / res.fun:.2f}%)")
