"""Variant D: purchase plan frozen exactly as delivered, but the battery is
dispatched in real time (it may charge/discharge within physical limits to
avoid the 5x emergency purchase).  Measures how much the "frozen storage
instruction" convention costs.
"""

from __future__ import annotations

import csv

import numpy as np
import openpyxl
from scipy.optimize import Bounds, LinearConstraint, linprog
from scipy.sparse import lil_matrix

DT = 1 / 6
E_MIN, E_MAX = 1200.0, 10800.0
E_REF = 6000.0
ETA_C = ETA_D = 0.9
LIM = 5000 / 6
N = 144
PEN = 5.0
LAMBDA_E = 1.3952


def read_actual(sheet):
    wb = openpyxl.load_workbook(
        "CUMCM 2026 C题/附件/附件2.xlsx", read_only=True, data_only=True
    )
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    dates = [r[0].date().isoformat() for r in rows[1:]]
    return dates, np.array([[float(v) for v in r[1:]] for r in rows[1:]]) * DT


dates, load = read_actual("小区负载")
_, pv = read_actual("光伏发电实际功率")
pv[pv < DT] = 0.0
wb = openpyxl.load_workbook(
    "CUMCM 2026 C题/附件/附件1.xlsx", read_only=True, data_only=True
)
price = np.array([float(r[1]) for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:]])
wb.close()

det = list(csv.DictReader(open("outputs/tables/2_滚动预测与调度明细.csv", encoding="utf-8-sig")))
assert len(det) == 334 * N
days = [det[k * N]["日期"] for k in range(334)]
g_all = np.array([float(r["计划购电量（kWh）"]) for r in det]).reshape(334, N)
best = np.array([float(r["计划购电量（kWh）"]) for r in det]).reshape(334, N)

# variables per day: h, c, d, r, w, e(145)  -> 5N + N + 1
nv = 5 * N + N + 1
o = dict(h=0, c=N, d=2 * N, r=3 * N, w=4 * N, e=5 * N)
rows = 2 * N + 1
A = lil_matrix((rows, nv))
lo = np.full(rows, -np.inf)
hi = np.full(rows, np.inf)
for t in range(N):
    A[t, o["h"] + t] = 1
    A[t, o["c"] + t] = -1
    A[t, o["d"] + t] = 1
    A[t, o["r"] + t] = -1
    A[t, o["w"] + t] = -1
    row = N + t
    A[row, o["c"] + t] = -ETA_C
    A[row, o["d"] + t] = 1 / ETA_D
    A[row, o["e"] + t] = -1
    A[row, o["e"] + t + 1] = 1
    lo[row] = hi[row] = 0.0
A[N + 0, o["c"]] = A[N + 0, o["c"]]  # noop
# terminal soft penalty rows
A[2 * N, o["e"] + N] = 1.0
lo[2 * N] = hi[2 * N] = E_REF

cvec = np.zeros(nv)
cvec[o["h"]:o["h"] + N] = PEN * price
# |e_N - E_REF| handled by splitting into two extra variables -> simpler: add
# ep, em as variables
nv2 = nv + 2
rows_A = 2 * N + 1
A2 = lil_matrix((rows_A, nv2))
A2[:2 * N, :nv] = A[:2 * N]
A2[2 * N, o["e"] + N] = 1.0
A2[2 * N, nv] = -1.0
A2[2 * N, nv + 1] = 1.0
cvec2 = np.zeros(nv2)
cvec2[:nv] = cvec
cvec2[nv] = cvec2[nv + 1] = LAMBDA_E
lo2 = np.concatenate([lo[:2 * N], [E_REF]])
hi2 = np.concatenate([hi[:2 * N], [E_REF]])

lb = np.zeros(nv2)
ub = np.full(nv2, np.inf)
ub[o["c"]:o["c"] + N] = LIM
ub[o["d"]:o["d"] + N] = LIM
lb[o["e"]:o["e"] + N + 1] = E_MIN
ub[o["e"]:o["e"] + N + 1] = E_MAX
constraint_lower = lo2.copy()
constraint_upper = hi2.copy()

tot_h = 0.0
tot_h_cost = 0.0
soc = None
soc_min, soc_max = 1e18, -1e18
curt = 0.0
waste = 0.0
for k, d in enumerate(days):
    i = dates.index(d)
    la, va = load[i], pv[i]
    # balance rhs: g + v + d + h - r - w = load + c
    for t in range(N):
        constraint_lower[t] = constraint_upper[t] = la[t] - va[t] - g_all[k, t]
    lb[o["w"]:o["w"] + N] = 0.0
    ub[o["w"]:o["w"] + N] = g_all[k]
    ub[o["r"]:o["r"] + N] = va
    if soc is None:
        lb[o["e"]] = ub[o["e"]] = 6666.471759
    else:
        lb[o["e"]] = ub[o["e"]] = soc
    res = linprog(cvec2, A_eq=A2.tocsr(), b_eq=constraint_lower,
                  bounds=np.column_stack((lb, ub)), method="highs")
    if not res.success:
        raise SystemExit(f"{d}: {res.message}")
    x = res.x
    tot_h += x[o["h"]:o["h"] + N].sum()
    tot_h_cost += PEN * float(price @ x[o["h"]:o["h"] + N])
    curt += x[o["r"]:o["r"] + N].sum()
    waste += x[o["w"]:o["w"] + N].sum()
    soc = float(x[o["e"] + N])
    ee = x[o["e"]:o["e"] + N + 1]
    soc_min = min(soc_min, float(ee.min()))
    soc_max = max(soc_max, float(ee.max()))

plan_cost = float((np.tile(price, 334) * g_all.ravel()).sum())
print("\n=== D) frozen purchase + real-time battery dispatch ===")
print("plan purchase cost (unchanged) =", plan_cost)
print("emergency kWh                  =", tot_h)
print("emergency cost                 =", tot_h_cost)
print("total 334d cost                =", plan_cost + tot_h_cost)
print("delivered total                = 14272696.316983")
print("saving vs delivered            =", 14272696.316983 - (plan_cost + tot_h_cost))
print("curtailment =", curt, "| unused planned purchase =", waste)
print("SOC range =", soc_min, soc_max, "| year-end =", soc)
