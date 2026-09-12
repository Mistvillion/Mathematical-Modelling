"""问题二：按日滚动预测、制定购电计划并回放实际运行。

在 conda 环境 2026C 中运行：
    python scripts_tables/2_rolling_optimization.py

采用 docs/Q2.md 的负载上分位数、光伏下分位数预测，仅使用计划日前的数据。
年末计划和实际储电量均须位于 [5400, 6600] kWh。
CSV 写入 outputs/tables，正式工作簿写入 outputs/results/result2.xlsx。
"""

from __future__ import annotations

import csv
import os
import sys
import tempfile
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time
from functools import lru_cache
from math import isfinite
from pathlib import Path
from types import SimpleNamespace
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import lil_matrix


PROJECT_DIR = Path(__file__).resolve().parents[1]
ATTACHMENT_DIR = PROJECT_DIR / "CUMCM 2026 C题" / "附件"
PRICE_AND_PRIOR_FILE = ATTACHMENT_DIR / "附件1.xlsx"
ACTUAL_DATA_FILE = ATTACHMENT_DIR / "附件2.xlsx"
RESULT_TEMPLATE_FILE = ATTACHMENT_DIR / "附件5" / "result2.xlsx"
OUTPUT_TABLE_DIR = PROJECT_DIR / "outputs" / "tables"
OUTPUT_RESULT_DIR = PROJECT_DIR / "outputs" / "results"
RESULT_FILE = OUTPUT_RESULT_DIR / "result2.xlsx"

DETAIL_FILE = OUTPUT_TABLE_DIR / "2_滚动预测与调度明细.csv"
PURCHASE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_指定日期购电量及全天结果.csv"
STORAGE_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_指定日期充放电量及储电量.csv"
EMERGENCY_SUMMARY_FILE = OUTPUT_TABLE_DIR / "2_指定日期紧急购电量.csv"
CHECK_FILE = OUTPUT_TABLE_DIR / "2_模型校验.csv"

N_PERIODS = 144
DELTA_T = 1 / 6
YEAR_START = date(2025, 1, 1)
OFFICIAL_START = date(2025, 2, 1)
YEAR_END = date(2025, 12, 31)
JANUARY_DAYS = 31
OFFICIAL_DAYS = 334

E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
E_REFERENCE = 6000.0
E_TERMINAL_MIN = 5400.0
E_TERMINAL_MAX = 6600.0
P_CHARGE_MAX = 5000.0
P_DISCHARGE_MAX = 5000.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
CHARGE_LIMIT = P_CHARGE_MAX * DELTA_T
DISCHARGE_LIMIT = P_DISCHARGE_MAX * DELTA_T
EMERGENCY_PRICE_MULTIPLIER = 5.0

LOAD_QUANTILE = 0.8
PHOTOVOLTAIC_QUANTILE = 0.2
NET_LOAD_QUANTILE = 1.0 - 1.0 / EMERGENCY_PRICE_MULTIPLIER
PHOTOVOLTAIC_NOISE_THRESHOLD = 1.0
DECAY_DAY_CANDIDATES = (3.0, 5.0, 7.0, 10.0, 14.0, 21.0, 28.0, 42.0, 56.0, 84.0)
TERMINAL_PENALTY_MULTIPLIERS = (0.5, 1.0, 2.0, 5.0)
# 参数候选、验证窗口、评分方式及默认参数均在初始化前固定。验证记录从
# 第一个已经结束的日期开始积累；没有可用记录时由 choose_candidate 使用默认值。
TUNING_START_INDEX = 0
VALIDATION_WINDOW = 28
VALIDATION_DECAY_DAYS = 14.0
DEFAULT_DECAY_DAYS = 14.0
MIP_REL_GAP = 1e-9

BALANCE_TOLERANCE = 1e-5
BOUND_TOLERANCE = 1e-5
INTEGER_TOLERANCE = 1e-6
TERMINAL_TOLERANCE = 1e-4
EMERGENCY_TOLERANCE = 1e-7
DISPLAY_DECIMALS = 6
DISPLAY_ZERO_TOLERANCE = 0.5 * 10 ** (-DISPLAY_DECIMALS)

ATTACHMENT1_COLUMNS = ("时间", "电价", "小区负载", "光伏发电预测功率")
ACTUAL_SHEET_NAMES = ("小区负载", "光伏发电实际功率")
RESULT_SHEET_NAMES = ("计划购电量", "充放电量", "紧急购电量")
STORAGE_BLOCK_LABELS = (
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
)
SPECIFIED_DATES = (
    date(2025, 3, 20),
    date(2025, 6, 21),
    date(2025, 9, 23),
    date(2025, 12, 21),
)
SPECIFIED_HOURS = (10, 12, 14, 16, 18, 20)


@dataclass(frozen=True)
class ModelData:
    """完成校验、清洗和时间对齐后的全部模型输入。"""

    dates: tuple[date, ...]
    interval_labels: tuple[str, ...]
    price: np.ndarray
    cold_start_load_energy: np.ndarray
    cold_start_photovoltaic_energy: np.ndarray
    actual_load_power: np.ndarray
    actual_photovoltaic_power: np.ndarray
    actual_load_energy: np.ndarray
    actual_photovoltaic_energy: np.ndarray


@dataclass(frozen=True)
class Hyperparameters:
    """在当前计划日之前选定的模型超参数。"""

    load_decay_days: float
    photovoltaic_decay_days: float
    terminal_penalty: float


@dataclass(frozen=True)
class TuningSummary:
    """超参数候选值与对应因果验证得分。"""

    load_decay_scores: tuple[tuple[float, float], ...]
    photovoltaic_decay_scores: tuple[tuple[float, float], ...]
    terminal_penalty_scores: tuple[tuple[float, float], ...]


@dataclass(frozen=True)
class Forecast:
    """某日在 0:00 可获得的风险修正预测。"""

    load_energy: np.ndarray
    photovoltaic_energy: np.ndarray
    latest_training_date: date | None
    used_prior_load: bool = False
    used_prior_photovoltaic: bool = False


@dataclass(frozen=True)
class PlanResult:
    """某日计划阶段 MILP 的最优解。"""

    grid_purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    curtailment: np.ndarray
    mode: np.ndarray
    stored_energy: np.ndarray
    positive_terminal_deviation: float
    negative_terminal_deviation: float
    objective_value: float
    solver_message: str
    solver_kind: str = "MILP"
    mip_gap: float = 0.0


@dataclass(frozen=True)
class ReplayResult:
    """固定计划后按实际负载和光伏逐时段回放的结果。"""

    charge: np.ndarray
    discharge: np.ndarray
    emergency_purchase: np.ndarray
    curtailment: np.ndarray
    unused_planned_purchase: np.ndarray
    stored_energy: np.ndarray


@dataclass(frozen=True)
class DailyResult:
    """一个自然日的预测、计划和实际执行结果。"""

    day: date
    forecast: Forecast
    plan: PlanResult
    replay: ReplayResult
    actual_load_energy: np.ndarray
    actual_photovoltaic_energy: np.ndarray
    planned_purchase_cost: float
    emergency_purchase_cost: float
    realized_total_cost: float
    hyperparameters: Hyperparameters
    parameter_latest_date: date | None


@dataclass(frozen=True)
class RollingResult:
    """问题二完整滚动求解结果。"""

    hyperparameters: Hyperparameters
    tuning: TuningSummary
    warmup_days: tuple[DailyResult, ...]
    official_days: tuple[DailyResult, ...]


@dataclass(frozen=True)
class EmergencyInterval:
    """由相邻紧急购电时段合并而成的连续区间。"""

    start_index: int
    stop_index: int
    label: str
    energy: float


def normalize_excel_date(value: object, location: str) -> date:
    """把 Excel 日期单元格统一转换为不含时分秒的 date。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"{location}应为 Excel 日期，实际为：{value!r}")


def time_to_end_minutes(value: time | str) -> int:
    """把附件的区间结束时刻转换为当天累计分钟。"""
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"时间应精确到分钟：{value!r}")
        return value.hour * 60 + value.minute

    if isinstance(value, str):
        text_value = value.strip()
        if text_value == "0:00+1":
            return 24 * 60
        try:
            hour, minute = map(int, text_value.split(":"))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"无法识别附件中的时间：{value!r}") from exc
        if 0 <= hour < 24 and 0 <= minute < 60:
            return hour * 60 + minute

    raise ValueError(f"无法识别附件中的时间：{value!r}")


def format_clock(minutes: int) -> str:
    """将 0 至 1440 分钟格式化为明确的当天时刻。"""
    if not 0 <= minutes <= 24 * 60:
        raise ValueError(f"分钟数超出当天范围：{minutes}")
    if minutes == 24 * 60:
        return "24:00"
    hour, minute = divmod(minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def validate_nonnegative_number(
    value: object,
    location: str,
) -> float:
    """检查模型输入是否为有限非负实数。"""
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not isfinite(value)
        or value < 0
    ):
        raise ValueError(f"{location}应为有限非负数，实际为：{value!r}")
    return float(value)


def read_price_and_cold_start() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """只读加载附件 1 中的重复日内电价和冷启动曲线。"""
    workbook = load_workbook(PRICE_AND_PRIOR_FILE, read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("附件 1 缺少工作表 Sheet1。")
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows or tuple(rows[0]) != ATTACHMENT1_COLUMNS:
        raise ValueError(f"附件 1 的表头应为：{ATTACHMENT1_COLUMNS}")
    if len(rows) != N_PERIODS + 1:
        raise ValueError(
            f"附件 1 应包含 {N_PERIODS} 条数据，实际为 {len(rows) - 1} 条。"
        )

    end_minutes = [time_to_end_minutes(row[0]) for row in rows[1:]]
    expected_end_minutes = list(range(10, 24 * 60 + 1, 10))
    if end_minutes != expected_end_minutes:
        raise ValueError("附件 1 的时间应从 0:10 到 0:00+1，每 10 分钟连续排列。")

    price = np.array(
        [
            validate_nonnegative_number(row[1], f"附件 1 第 {row_index} 行电价")
            for row_index, row in enumerate(rows[1:], start=2)
        ],
        dtype=float,
    )
    load_power = np.array(
        [
            validate_nonnegative_number(
                row[2], f"附件 1 第 {row_index} 行小区负载"
            )
            for row_index, row in enumerate(rows[1:], start=2)
        ],
        dtype=float,
    )
    photovoltaic_power = np.array(
        [
            validate_nonnegative_number(
                row[3], f"附件 1 第 {row_index} 行光伏发电预测功率"
            )
            for row_index, row in enumerate(rows[1:], start=2)
        ],
        dtype=float,
    )
    photovoltaic_power[photovoltaic_power < PHOTOVOLTAIC_NOISE_THRESHOLD] = 0.0
    return price, load_power * DELTA_T, photovoltaic_power * DELTA_T


def read_actual_sheet(
    worksheet_name: str,
    rows: list[tuple[object, ...]],
) -> tuple[tuple[date, ...], tuple[int, ...], np.ndarray]:
    """校验并读取附件 2 的一个 365 天功率工作表。"""
    if not rows:
        raise ValueError(f"附件 2 的“{worksheet_name}”工作表为空。")
    if len(rows) != 366:
        raise ValueError(
            f"附件 2 的“{worksheet_name}”应包含表头和 365 天数据，"
            f"实际共有 {len(rows)} 行。"
        )
    if len(rows[0]) != N_PERIODS + 1 or rows[0][0] != "日期\\时间":
        raise ValueError(
            f"附件 2 的“{worksheet_name}”首列应为“日期\\时间”，"
            f"并包含 {N_PERIODS} 个时段。"
        )

    end_minutes = tuple(time_to_end_minutes(value) for value in rows[0][1:])
    expected_end_minutes = tuple(range(10, 24 * 60 + 1, 10))
    if end_minutes != expected_end_minutes:
        raise ValueError(
            f"附件 2 的“{worksheet_name}”时间应从 0:10 到 0:00+1 连续排列。"
        )

    dates = tuple(
        normalize_excel_date(row[0], f"“{worksheet_name}”第 {row_index} 行日期")
        for row_index, row in enumerate(rows[1:], start=2)
    )
    expected_dates = tuple(
        date.fromordinal(YEAR_START.toordinal() + day_offset)
        for day_offset in range(365)
    )
    if dates != expected_dates:
        raise ValueError(
            f"附件 2 的“{worksheet_name}”日期必须从 {YEAR_START} 连续到 {YEAR_END}。"
        )

    power = np.empty((365, N_PERIODS), dtype=float)
    for day_index, row in enumerate(rows[1:]):
        if len(row) != N_PERIODS + 1:
            raise ValueError(
                f"附件 2 的“{worksheet_name}”第 {day_index + 2} 行列数不正确。"
            )
        for period_index, value in enumerate(row[1:]):
            power[day_index, period_index] = validate_nonnegative_number(
                value,
                f"“{worksheet_name}”第 {day_index + 2} 行第 {period_index + 2} 列",
            )
    return dates, end_minutes, power


def load_model_data() -> ModelData:
    """加载附件 1 和附件 2，并建立统一的 10 分钟时间轴。"""
    price, cold_load_energy, cold_photovoltaic_energy = (
        read_price_and_cold_start()
    )

    workbook = load_workbook(ACTUAL_DATA_FILE, read_only=True, data_only=True)
    try:
        missing_sheets = set(ACTUAL_SHEET_NAMES).difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "附件 2 缺少工作表：" + "、".join(sorted(missing_sheets))
            )
        load_rows = list(workbook[ACTUAL_SHEET_NAMES[0]].iter_rows(values_only=True))
        photovoltaic_rows = list(
            workbook[ACTUAL_SHEET_NAMES[1]].iter_rows(values_only=True)
        )
    finally:
        workbook.close()

    load_dates, load_end_minutes, load_power = read_actual_sheet(
        ACTUAL_SHEET_NAMES[0], load_rows
    )
    photovoltaic_dates, photovoltaic_end_minutes, photovoltaic_power = (
        read_actual_sheet(ACTUAL_SHEET_NAMES[1], photovoltaic_rows)
    )
    if load_dates != photovoltaic_dates:
        raise ValueError("附件 2 的负载和光伏工作表日期不一致。")
    if load_end_minutes != photovoltaic_end_minutes:
        raise ValueError("附件 2 的负载和光伏工作表时间表头不一致。")

    photovoltaic_power = photovoltaic_power.copy()
    photovoltaic_power[
        photovoltaic_power < PHOTOVOLTAIC_NOISE_THRESHOLD
    ] = 0.0
    interval_labels = tuple(
        f"{format_clock(end_minute - 10)}-{format_clock(end_minute)}"
        for end_minute in load_end_minutes
    )
    return ModelData(
        dates=load_dates,
        interval_labels=interval_labels,
        price=price,
        cold_start_load_energy=cold_load_energy,
        cold_start_photovoltaic_energy=cold_photovoltaic_energy,
        actual_load_power=load_power,
        actual_photovoltaic_power=photovoltaic_power,
        actual_load_energy=load_power * DELTA_T,
        actual_photovoltaic_energy=photovoltaic_power * DELTA_T,
    )


def is_high_load_day(day: date) -> bool:
    """周日至周四为高负载类，周五、周六为低负载类。"""
    return day.weekday() <= 3 or day.weekday() == 6


class ForecastEngine:
    """批量计算历史加权分位数，输入限定为计划日前的数据。"""

    def quantile_rows(self, values: np.ndarray, ages: np.ndarray, q: float) -> np.ndarray:
        """返回分位数对应的原样本行号；净负荷允许为负，输入仍须有限。"""
        values = np.asarray(values, dtype=np.float64)
        ages = np.asarray(ages, dtype=np.float64)
        if ages.ndim != 1:
            raise ValueError("样本年龄必须为一维数组。")
        if values.ndim != 2 or values.shape != (len(ages), N_PERIODS):
            raise ValueError("分位数样本形状不正确。")
        if (
            not len(ages) or not np.isfinite(ages).all() or np.any(ages <= 0)
            or not np.isfinite(values).all() or not 0 < q < 1
        ):
            raise ValueError("预测样本必须为有限历史数据，样本年龄必须为正。")
        # 减去最近样本年龄只改变共同倍率，不改变分位数，并避免极老样本全下溢。
        weights = np.exp(-(ages[None, :] - np.min(ages)) / np.asarray(DECAY_DAY_CANDIDATES)[:, None])
        order = np.argsort(values, axis=0, kind="stable")
        cumulative = np.cumsum(weights[:, order], axis=1)
        indices = np.argmax(cumulative >= q * cumulative[:, -1:, :], axis=1)
        return order[indices, np.arange(N_PERIODS)[None, :]]

    def quantiles(self, values: np.ndarray, ages: np.ndarray, q: float) -> np.ndarray:
        """返回形状为 (衰减候选数, 144) 的分位数预测。"""
        values = np.asarray(values, dtype=np.float64)
        return values[self.quantile_rows(values, ages, q), np.arange(N_PERIODS)[None, :]]

    @staticmethod
    def validate_history(
        target: date, history_dates: Sequence[date],
        historical_load: np.ndarray, historical_pv: np.ndarray,
        prior_load: np.ndarray, prior_pv: np.ndarray,
    ) -> None:
        """所有预测方法共享同一因果边界与物理输入校验。"""
        if any(day >= target for day in history_dates):
            raise ValueError("预测接口收到当天或未来日期，拒绝前视数据。")
        if any(a >= b for a, b in zip(history_dates, history_dates[1:])):
            raise ValueError("历史日期必须严格递增且无重复。")
        if historical_load.shape != (len(history_dates), N_PERIODS):
            raise ValueError("历史负载样本数与日期数不一致。")
        if historical_pv.shape != historical_load.shape:
            raise ValueError("历史负载与光伏样本数不一致。")
        if prior_load.shape != (N_PERIODS,) or prior_pv.shape != (N_PERIODS,):
            raise ValueError("冷启动曲线必须各含144个时段。")
        for values in (historical_load, historical_pv, prior_load, prior_pv):
            if not np.isfinite(values).all() or np.any(values < 0):
                raise ValueError("负载、光伏必须为有限非负数。")

    @staticmethod
    def latest_training_date(
        target: date, history_dates: Sequence[date],
    ) -> date | None:
        """光伏预测使用全部历史，返回最晚实测日期。"""
        history_dates = tuple(history_dates)
        if any(day >= target for day in history_dates):
            raise ValueError("预测接口收到当天或未来日期，拒绝前视数据。")
        if any(a >= b for a, b in zip(history_dates, history_dates[1:])):
            raise ValueError("历史日期必须严格递增且无重复。")
        return history_dates[-1] if history_dates else None

    @staticmethod
    def prior_fallback_flags(
        target: date, history_dates: Sequence[date],
    ) -> tuple[bool, bool]:
        """返回负载、光伏是否回退到附件 1 冷启动先验。"""
        history_dates = tuple(history_dates)
        if any(day >= target for day in history_dates):
            raise ValueError("预测接口收到当天或未来日期，拒绝前视数据。")
        same_type_history = any(
            is_high_load_day(day) == is_high_load_day(target)
            for day in history_dates
        )
        return not same_type_history, not history_dates

    def candidates(
        self,
        target: date,
        history_dates: Sequence[date],
        historical_load: np.ndarray,
        historical_pv: np.ndarray,
        prior_load: np.ndarray,
        prior_pv: np.ndarray,
    ) -> tuple[np.ndarray, np.ndarray]:
        self.validate_history(target, history_dates, historical_load, historical_pv, prior_load, prior_pv)
        ages = np.array([(target - day).days for day in history_dates], dtype=float)
        same_type = np.array(
            [is_high_load_day(day) == is_high_load_day(target) for day in history_dates],
            dtype=bool,
        )
        # 有同类真实历史后完全退出附件 1 先验，避免长期掺入未分类伪样本。
        load_values = historical_load[same_type]
        load_ages = ages[same_type]
        if not len(load_values):
            load_values, load_ages = prior_load[None, :], np.ones(1)
        pv_values, pv_ages = historical_pv, ages
        if not len(pv_values):
            pv_values, pv_ages = prior_pv[None, :], np.ones(1)
        return (
            self.quantiles(load_values, load_ages, LOAD_QUANTILE),
            self.quantiles(pv_values, pv_ages, PHOTOVOLTAIC_QUANTILE),
        )

def pinball_loss(
    actual: np.ndarray,
    forecast: np.ndarray,
    quantile: float,
) -> float:
    """计算分位数预测的平均 pinball 损失。"""
    error = actual - forecast
    loss = np.maximum(quantile * error, (quantile - 1.0) * error)
    return float(np.mean(loss))


def validation_indices(
    losses: Sequence[np.ndarray],
    day_index: int,
    dates: Sequence[date],
    same_load_type: bool = False,
) -> tuple[int, ...]:
    """返回当前日期可用于调参的历史验证日索引。"""
    if len(losses) != day_index:
        raise RuntimeError("验证损失的截止日与计划日不一致。")
    if not 0 <= day_index < len(dates):
        raise RuntimeError("计划日索引超出日期范围。")
    return tuple(
        i for i in range(max(TUNING_START_INDEX, day_index - VALIDATION_WINDOW), day_index)
        if not same_load_type or is_high_load_day(dates[i]) == is_high_load_day(dates[day_index])
    )


def latest_validation_date(
    losses: Sequence[np.ndarray],
    day_index: int,
    dates: Sequence[date],
    same_load_type: bool = False,
) -> date | None:
    """返回当前参数选择实际使用的最晚已结束验证日。"""
    indices = validation_indices(losses, day_index, dates, same_load_type)
    return dates[indices[-1]] if indices else None


def parameter_validation_latest_date(
    sources: Sequence[tuple[Sequence[np.ndarray], bool]],
    day_index: int,
    dates: Sequence[date],
) -> date | None:
    """合并各参数来源，返回本日调参使用的最晚已结束验证日。"""
    latest_dates = [
        latest for losses, same_load_type in sources
        if (latest := latest_validation_date(
            losses, day_index, dates, same_load_type
        )) is not None
    ]
    return max(latest_dates) if latest_dates else None


def historical_scores(
    losses: Sequence[np.ndarray],
    day_index: int,
    count: int,
    dates: Sequence[date],
    same_load_type: bool = False,
) -> np.ndarray:
    """只汇总当前日以前的真实一步预测损失，最近 28 天指数加权。"""
    indices = validation_indices(losses, day_index, dates, same_load_type)
    if not indices:
        return np.zeros(count)
    weights = np.exp(-(day_index - np.array(indices)) / VALIDATION_DECAY_DAYS)
    return np.average(np.stack([losses[i] for i in indices]), axis=0, weights=weights)


def choose_candidate(scores: np.ndarray, default_index: int) -> int:
    """得分相同时优先预设参数，避免仅由候选排列决定冷启动策略。"""
    best = float(np.min(scores))
    tied = np.flatnonzero(np.isclose(scores, best, atol=1e-10, rtol=0.0))
    return default_index if default_index in tied else int(tied[0])


def validate_terminal_bounds() -> None:
    """年末上下限须有限、有序，且包含在设备物理安全范围内。"""
    if not all(isfinite(value) for value in (E_TERMINAL_MIN, E_TERMINAL_MAX)) or not (
        E_MIN <= E_TERMINAL_MIN <= E_TERMINAL_MAX <= E_MAX
    ):
        raise ValueError("年末区间必须满足 E_MIN <= E_TERMINAL_MIN <= E_TERMINAL_MAX <= E_MAX，且上下限均有限。")


def terminal_interval_error(energy: float) -> float:
    """距年末允许区间的越界量；区间内为 0，非有限末态不能通过校验。"""
    validate_terminal_bounds()
    if not isfinite(energy):
        return float("inf")
    return float(max(E_TERMINAL_MIN - energy, energy - E_TERMINAL_MAX, 0.0))


@lru_cache(maxsize=2)
def dispatch_matrix(hard_terminal: bool):
    """每个进程只构造一次固定稀疏系数矩阵；逐日只更新边界与右端项。"""
    n = N_PERIODS
    matrix = lil_matrix((4 * n + (not hard_terminal), 6 * n + 3), dtype=float)
    for t in range(n):
        matrix[t, t] = 1.0
        matrix[t, n + t] = -1.0
        matrix[t, 2 * n + t] = 1.0
        matrix[t, 3 * n + t] = -1.0
        matrix[n + t, n + t] = -ETA_CHARGE
        matrix[n + t, 2 * n + t] = 1.0 / ETA_DISCHARGE
        matrix[n + t, 5 * n + t] = -1.0
        matrix[n + t, 5 * n + t + 1] = 1.0
        matrix[2 * n + t, n + t] = 1.0
        matrix[2 * n + t, 4 * n + t] = -CHARGE_LIMIT
        matrix[3 * n + t, 2 * n + t] = 1.0
        matrix[3 * n + t, 4 * n + t] = DISCHARGE_LIMIT
    if not hard_terminal:
        matrix[4 * n, 6 * n] = 1.0
        matrix[4 * n, 6 * n + 1] = -1.0
        matrix[4 * n, 6 * n + 2] = 1.0
    csc = matrix.tocsc()
    equality_rows = list(range(2 * n)) + ([] if hard_terminal else [4 * n])
    return csc, csc[equality_rows, :], csc[2 * n : 4 * n, :]


def solve_daily_plan(
    price: np.ndarray,
    forecast: Forecast,
    initial_energy: float,
    terminal_penalty: float,
    require_terminal_interval: bool,
    force_milp: bool = False,
) -> PlanResult:
    """先用 LP 下界验证最优性；互斥不成立才求解完整 MILP。

    只重建无成本的模式 u，不对连续能量决策做取整。当 LP 最优解本身
    满足所有原 MILP 约束，其目标既是可行上界也是松弛下界，故可接受。
    禁止不加检查地删除二元变量或用启发式解冒充最优解。
    """
    validate_terminal_bounds()
    if not E_MIN - BOUND_TOLERANCE <= initial_energy <= E_MAX + BOUND_TOLERANCE:
        raise ValueError("初始储电量超出物理安全边界。")
    if not isfinite(terminal_penalty) or terminal_penalty < 0:
        raise ValueError("终端惩罚必须是有限非负数。")
    offset_grid = 0
    offset_charge = N_PERIODS
    offset_discharge = 2 * N_PERIODS
    offset_curtailment = 3 * N_PERIODS
    offset_mode = 4 * N_PERIODS
    offset_energy = 5 * N_PERIODS
    offset_positive_deviation = 6 * N_PERIODS + 1
    offset_negative_deviation = 6 * N_PERIODS + 2
    variable_count = 6 * N_PERIODS + 3

    objective = np.zeros(variable_count)
    objective[offset_grid : offset_grid + N_PERIODS] = price
    if not require_terminal_interval:
        objective[offset_positive_deviation] = terminal_penalty
        objective[offset_negative_deviation] = terminal_penalty

    lower_bounds = np.zeros(variable_count)
    upper_bounds = np.full(variable_count, np.inf)
    upper_bounds[offset_charge : offset_charge + N_PERIODS] = CHARGE_LIMIT
    upper_bounds[offset_discharge : offset_discharge + N_PERIODS] = (
        DISCHARGE_LIMIT
    )
    upper_bounds[offset_curtailment : offset_curtailment + N_PERIODS] = (
        forecast.photovoltaic_energy
    )
    upper_bounds[offset_mode : offset_mode + N_PERIODS] = 1.0
    lower_bounds[offset_energy : offset_energy + N_PERIODS + 1] = E_MIN
    upper_bounds[offset_energy : offset_energy + N_PERIODS + 1] = E_MAX
    lower_bounds[offset_energy] = initial_energy
    upper_bounds[offset_energy] = initial_energy

    if require_terminal_interval:
        lower_bounds[offset_energy + N_PERIODS] = E_TERMINAL_MIN
        upper_bounds[offset_energy + N_PERIODS] = E_TERMINAL_MAX
        upper_bounds[offset_positive_deviation] = 0.0
        upper_bounds[offset_negative_deviation] = 0.0

    integrality = np.zeros(variable_count, dtype=int)
    integrality[offset_mode : offset_mode + N_PERIODS] = 1

    terminal_constraint_count = 0 if require_terminal_interval else 1
    constraint_count = 4 * N_PERIODS + terminal_constraint_count
    matrix, equality_matrix, inequality_matrix = dispatch_matrix(require_terminal_interval)
    constraint_lower = np.full(constraint_count, -np.inf)
    constraint_upper = np.full(constraint_count, np.inf)

    for period_index in range(N_PERIODS):
        balance_row = period_index
        balance_rhs = (
            forecast.load_energy[period_index]
            - forecast.photovoltaic_energy[period_index]
        )
        constraint_lower[balance_row] = balance_rhs
        constraint_upper[balance_row] = balance_rhs

        state_row = N_PERIODS + period_index
        constraint_lower[state_row] = 0.0
        constraint_upper[state_row] = 0.0

        charge_mode_row = 2 * N_PERIODS + period_index
        constraint_upper[charge_mode_row] = 0.0

        discharge_mode_row = 3 * N_PERIODS + period_index
        constraint_upper[discharge_mode_row] = DISCHARGE_LIMIT

    if not require_terminal_interval:
        terminal_row = 4 * N_PERIODS
        constraint_lower[terminal_row] = E_REFERENCE
        constraint_upper[terminal_row] = E_REFERENCE

    solver_kind = "MILP"
    result = None
    if not force_milp:
        equality_rhs = np.concatenate([
            constraint_lower[:2 * N_PERIODS],
            np.array([] if require_terminal_interval else [E_REFERENCE]),
        ])
        relaxed = linprog(
            c=objective,
            A_eq=equality_matrix, b_eq=equality_rhs,
            A_ub=inequality_matrix,
            b_ub=constraint_upper[2 * N_PERIODS:4 * N_PERIODS],
            bounds=np.column_stack((lower_bounds, upper_bounds)),
            method="highs-ds",
            options={"presolve": True, "primal_feasibility_tolerance": 1e-8,
                     "dual_feasibility_tolerance": 1e-8},
        )
        if relaxed.success and relaxed.x is not None:
            candidate = relaxed.x.copy()
            c = candidate[offset_charge:offset_charge + N_PERIODS]
            d = candidate[offset_discharge:offset_discharge + N_PERIODS]
            candidate[offset_mode:offset_mode + N_PERIODS] = (c > d).astype(float)
            activity = matrix @ candidate
            if (
                np.isfinite(candidate).all()
                and np.all(candidate >= lower_bounds - BOUND_TOLERANCE)
                and np.all(candidate <= upper_bounds + BOUND_TOLERANCE)
                and np.all(activity >= constraint_lower - BALANCE_TOLERANCE)
                and np.all(activity <= constraint_upper + BALANCE_TOLERANCE)
            ):
                relaxed.x = candidate
                result, solver_kind = relaxed, "LP已验证互斥及全部原约束"
    if result is None:
        result = milp(
            c=objective,
            integrality=integrality,
            bounds=Bounds(lower_bounds, upper_bounds),
            constraints=LinearConstraint(matrix, constraint_lower, constraint_upper),
            options={"disp": False, "presolve": True, "mip_rel_gap": MIP_REL_GAP},
        )
    if not result.success or result.x is None:
        raise RuntimeError(
            f"MILP 求解失败，状态码 {result.status}：{result.message}；"
            f"当日初始储电量为 {initial_energy:.6f} kWh。"
        )

    solution = result.x
    if not np.isfinite(solution).all() or not isfinite(float(result.fun)):
        raise RuntimeError("求解器返回非有限值。")
    actual_gap = float(getattr(result, "mip_gap", 0.0) or 0.0)
    if actual_gap > MIP_REL_GAP + 1e-12:
        raise RuntimeError(f"求解间隙 {actual_gap} 超过要求 {MIP_REL_GAP}。")
    return PlanResult(
        grid_purchase=solution[offset_grid : offset_grid + N_PERIODS],
        charge=solution[offset_charge : offset_charge + N_PERIODS],
        discharge=solution[offset_discharge : offset_discharge + N_PERIODS],
        curtailment=solution[
            offset_curtailment : offset_curtailment + N_PERIODS
        ],
        mode=solution[offset_mode : offset_mode + N_PERIODS],
        stored_energy=solution[offset_energy : offset_energy + N_PERIODS + 1],
        positive_terminal_deviation=float(solution[offset_positive_deviation]),
        negative_terminal_deviation=float(solution[offset_negative_deviation]),
        objective_value=float(result.fun),
        solver_message=str(result.message),
        solver_kind=solver_kind,
        mip_gap=actual_gap,
    )


def replay_actual_day(
    plan: PlanResult,
    actual_load_energy: np.ndarray,
    actual_photovoltaic_energy: np.ndarray,
    initial_energy: float,
) -> ReplayResult:
    """固定日计划，按 docs/Q2.md 的优先级规则回放实际运行。"""
    charge = np.zeros(N_PERIODS)
    discharge = np.zeros(N_PERIODS)
    emergency_purchase = np.zeros(N_PERIODS)
    curtailment = np.zeros(N_PERIODS)
    unused_planned_purchase = np.zeros(N_PERIODS)
    stored_energy = np.empty(N_PERIODS + 1)
    stored_energy[0] = initial_energy

    for period_index in range(N_PERIODS):
        current_energy = stored_energy[period_index]
        planned_purchase = plan.grid_purchase[period_index]
        photovoltaic_energy = actual_photovoltaic_energy[period_index]
        load_energy = actual_load_energy[period_index]

        load_deficit_before_storage = max(
            load_energy - planned_purchase - photovoltaic_energy,
            0.0,
        )
        available_discharge = max(
            ETA_DISCHARGE * (current_energy - E_MIN),
            0.0,
        )
        discharge[period_index] = min(
            max(plan.discharge[period_index], 0.0),
            DISCHARGE_LIMIT,
            available_discharge,
            load_deficit_before_storage,
        )

        emergency_purchase[period_index] = max(
            load_energy
            - planned_purchase
            - photovoltaic_energy
            - discharge[period_index],
            0.0,
        )

        surplus_after_load = max(
            planned_purchase
            + photovoltaic_energy
            + discharge[period_index]
            - load_energy,
            0.0,
        )
        storage_headroom = max((E_MAX - current_energy) / ETA_CHARGE, 0.0)
        charge[period_index] = min(
            max(plan.charge[period_index], 0.0),
            CHARGE_LIMIT,
            storage_headroom,
            surplus_after_load,
        )

        stored_energy[period_index + 1] = (
            current_energy
            + ETA_CHARGE * charge[period_index]
            - discharge[period_index] / ETA_DISCHARGE
        )

        remaining_energy = max(
            planned_purchase
            + photovoltaic_energy
            + discharge[period_index]
            + emergency_purchase[period_index]
            - load_energy
            - charge[period_index],
            0.0,
        )
        curtailment[period_index] = min(
            photovoltaic_energy,
            remaining_energy,
        )
        unused_planned_purchase[period_index] = max(
            remaining_energy - curtailment[period_index],
            0.0,
        )

    return ReplayResult(
        charge=charge,
        discharge=discharge,
        emergency_purchase=emergency_purchase,
        curtailment=curtailment,
        unused_planned_purchase=unused_planned_purchase,
        stored_energy=stored_energy,
    )


def validate_plan(
    day: date,
    forecast: Forecast,
    plan: PlanResult,
    initial_energy: float,
    require_terminal_interval: bool,
) -> None:
    """独立检查计划阶段 MILP 的主要约束。"""
    for name in ("grid_purchase", "charge", "discharge", "curtailment", "mode", "stored_energy"):
        array = getattr(plan, name)
        size = N_PERIODS + (name == "stored_energy")
        if array.shape != (size,) or not np.isfinite(array).all():
            raise RuntimeError(f"{day} 计划 {name} 形状或数值无效。")
    for array in (forecast.load_energy, forecast.photovoltaic_energy):
        if array.shape != (N_PERIODS,) or not np.isfinite(array).all() or np.any(array < 0):
            raise RuntimeError(f"{day} 预测值形状或数值无效。")
    if not isinstance(forecast.used_prior_load, bool) or not isinstance(forecast.used_prior_photovoltaic, bool):
        raise RuntimeError(f"{day} 预测先验回退标记无效。")
    if not all(isfinite(value) for value in (
        plan.positive_terminal_deviation, plan.negative_terminal_deviation,
        plan.objective_value, plan.mip_gap,
    )):
        raise RuntimeError(f"{day} 计划目标值或偏差变量非有限。")
    balance_residual = (
        plan.grid_purchase
        + forecast.photovoltaic_energy
        + plan.discharge
        - plan.curtailment
        - forecast.load_energy
        - plan.charge
    )
    state_residual = (
        plan.stored_energy[1:]
        - plan.stored_energy[:-1]
        - ETA_CHARGE * plan.charge
        + plan.discharge / ETA_DISCHARGE
    )
    failures: list[str] = []
    if np.min(plan.mode) < -INTEGER_TOLERANCE or np.max(plan.mode) > 1 + INTEGER_TOLERANCE:
        failures.append("计划模式不是二元值")
    if np.max(plan.charge - CHARGE_LIMIT * plan.mode) > BOUND_TOLERANCE:
        failures.append("计划充电量违反模式约束")
    if np.max(plan.discharge - DISCHARGE_LIMIT * (1 - plan.mode)) > BOUND_TOLERANCE:
        failures.append("计划放电量违反模式约束")
    if min(plan.positive_terminal_deviation, plan.negative_terminal_deviation) < -BOUND_TOLERANCE:
        failures.append("终端偏差变量为负")
    if np.max(np.abs(balance_residual)) > BALANCE_TOLERANCE:
        failures.append("计划电量平衡残差超限")
    if np.max(np.abs(state_residual)) > BALANCE_TOLERANCE:
        failures.append("计划储电量状态转移残差超限")
    if abs(plan.stored_energy[0] - initial_energy) > BOUND_TOLERANCE:
        failures.append("计划初始储电量不等于实际日初状态")
    if np.min(plan.stored_energy) < E_MIN - BOUND_TOLERANCE:
        failures.append("计划储电量低于安全下限")
    if np.max(plan.stored_energy) > E_MAX + BOUND_TOLERANCE:
        failures.append("计划储电量高于安全上限")
    if np.min(plan.grid_purchase) < -BOUND_TOLERANCE:
        failures.append("计划购电量出现负值")
    if np.min(plan.charge) < -BOUND_TOLERANCE:
        failures.append("计划充电量出现负值")
    if np.min(plan.discharge) < -BOUND_TOLERANCE:
        failures.append("计划放电量出现负值")
    if np.max(plan.charge) > CHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("计划充电量超过时段功率上限")
    if np.max(plan.discharge) > DISCHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("计划放电量超过时段功率上限")
    if np.max(np.minimum(plan.charge, plan.discharge)) > BOUND_TOLERANCE:
        failures.append("计划中存在同时充放电")
    if np.max(np.abs(plan.mode - np.rint(plan.mode))) > INTEGER_TOLERANCE:
        failures.append("计划充放电模式不是整数")
    if np.min(plan.curtailment) < -BOUND_TOLERANCE:
        failures.append("计划弃光量出现负值")
    if (
        np.max(plan.curtailment - forecast.photovoltaic_energy)
        > BOUND_TOLERANCE
    ):
        failures.append("计划弃光量超过预测光伏电量")

    if require_terminal_interval:
        if terminal_interval_error(plan.stored_energy[-1]) > TERMINAL_TOLERANCE:
            failures.append("最终日计划储电量超出年末允许区间")
    else:
        terminal_deviation_residual = (
            plan.stored_energy[-1]
            - E_REFERENCE
            - plan.positive_terminal_deviation
            + plan.negative_terminal_deviation
        )
        if abs(terminal_deviation_residual) > BALANCE_TOLERANCE:
            failures.append("日末储电量偏差线性化残差超限")

    if failures:
        raise RuntimeError(f"{day} 计划模型校验失败：" + "；".join(failures))


def validate_replay(
    day: date,
    plan: PlanResult,
    replay: ReplayResult,
    actual_load_energy: np.ndarray,
    actual_photovoltaic_energy: np.ndarray,
) -> None:
    """检查实际回放的平衡、边界、优先级和紧急购电口径。"""
    for name in ("charge", "discharge", "emergency_purchase", "curtailment", "unused_planned_purchase", "stored_energy"):
        array = getattr(replay, name)
        size = N_PERIODS + (name == "stored_energy")
        if array.shape != (size,) or not np.isfinite(array).all() or np.any(array < -BOUND_TOLERANCE):
            raise RuntimeError(f"{day} 实际 {name} 形状或数值无效。")
    balance_residual = (
        plan.grid_purchase
        + actual_photovoltaic_energy
        + replay.discharge
        + replay.emergency_purchase
        - replay.curtailment
        - replay.unused_planned_purchase
        - actual_load_energy
        - replay.charge
    )
    state_residual = (
        replay.stored_energy[1:]
        - replay.stored_energy[:-1]
        - ETA_CHARGE * replay.charge
        + replay.discharge / ETA_DISCHARGE
    )
    expected_emergency_purchase = np.maximum(
        actual_load_energy
        - plan.grid_purchase
        - actual_photovoltaic_energy
        - replay.discharge,
        0.0,
    )
    failures: list[str] = []
    if np.max(np.abs(balance_residual)) > BALANCE_TOLERANCE:
        failures.append("实际电量平衡残差超限")
    if np.max(np.abs(state_residual)) > BALANCE_TOLERANCE:
        failures.append("实际储电量状态转移残差超限")
    if (
        np.max(np.abs(replay.emergency_purchase - expected_emergency_purchase))
        > BALANCE_TOLERANCE
    ):
        failures.append("紧急购电量不等于实际负载缺口")
    if np.max(replay.charge - plan.charge) > BOUND_TOLERANCE:
        failures.append("实际充电量超过计划充电指令")
    if np.max(replay.discharge - plan.discharge) > BOUND_TOLERANCE:
        failures.append("实际放电量超过计划放电指令")
    if np.max(replay.charge) > CHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("实际充电量超过时段功率上限")
    if np.max(replay.discharge) > DISCHARGE_LIMIT + BOUND_TOLERANCE:
        failures.append("实际放电量超过时段功率上限")
    if np.max(np.minimum(replay.charge, replay.discharge)) > BOUND_TOLERANCE:
        failures.append("实际运行中存在同时充放电")
    if np.any(
        (replay.emergency_purchase > EMERGENCY_TOLERANCE)
        & (replay.charge > BOUND_TOLERANCE)
    ):
        failures.append("紧急购电被用于储能充电")
    if np.min(replay.curtailment) < -BOUND_TOLERANCE:
        failures.append("实际弃光量出现负值")
    if (
        np.max(replay.curtailment - actual_photovoltaic_energy)
        > BOUND_TOLERANCE
    ):
        failures.append("实际弃光量超过可用光伏电量")
    if np.min(replay.unused_planned_purchase) < -BOUND_TOLERANCE:
        failures.append("未利用计划购电量出现负值")
    if (
        np.max(replay.unused_planned_purchase - plan.grid_purchase)
        > BOUND_TOLERANCE
    ):
        failures.append("未利用计划购电量超过计划购电量")
    if np.min(replay.stored_energy) < E_MIN - BOUND_TOLERANCE:
        failures.append("实际储电量低于安全下限")
    if np.max(replay.stored_energy) > E_MAX + BOUND_TOLERANCE:
        failures.append("实际储电量高于安全上限")

    if failures:
        raise RuntimeError(f"{day} 实际回放校验失败：" + "；".join(failures))


def solve_validated_plan(
    day: date, price: np.ndarray, forecast: Forecast, initial: float,
    penalty: float, hard_terminal: bool,
) -> PlanResult:
    """求解并校验日前计划；输入不含当天实际负载或光伏。"""
    try:
        plan = solve_daily_plan(price, forecast, initial, penalty, hard_terminal)
        validate_plan(day, forecast, plan, initial, hard_terminal)
        expected = float(price @ plan.grid_purchase)
        if not hard_terminal:
            expected += penalty * (plan.positive_terminal_deviation + plan.negative_terminal_deviation)
        if abs(plan.objective_value - expected) > BALANCE_TOLERANCE:
            raise RuntimeError("目标函数重算不一致。")
        return plan
    except Exception as exc:
        raise RuntimeError(f"{day} 求解失败，e0={initial:.6f}，lambda={penalty}：{exc}") from exc


def solve_rolling_model(data: ModelData, engine: ForecastEngine) -> RollingResult:
    """逐日冻结预测与计划，实际回放后才更新下一日的验证记录和储电量。"""
    penalties = np.max(data.price) * np.asarray(TERMINAL_PENALTY_MULTIPLIERS)
    default_penalty = TERMINAL_PENALTY_MULTIPLIERS.index(1.0)
    default_decay = DECAY_DAY_CANDIDATES.index(DEFAULT_DECAY_DAYS)
    terminal_value = float(np.max(data.price) / ETA_CHARGE)
    losses_load: list[np.ndarray] = []
    losses_pv: list[np.ndarray] = []
    losses_penalty: list[np.ndarray] = []
    results: list[DailyResult] = []
    initial_energy = E_INITIAL

    for day_index, day in enumerate(data.dates):
        hard_terminal = day == YEAR_END
        history_dates = data.dates[:day_index]
        load_scores = historical_scores(losses_load, day_index, len(DECAY_DAY_CANDIDATES), data.dates, True)
        pv_scores = historical_scores(losses_pv, day_index, len(DECAY_DAY_CANDIDATES), data.dates)
        penalty_scores = historical_scores(losses_penalty, day_index, len(penalties), data.dates)
        load_index = choose_candidate(load_scores, default_decay)
        pv_index = choose_candidate(pv_scores, default_decay)
        penalty_index = choose_candidate(penalty_scores, default_penalty)
        parameter_latest = parameter_validation_latest_date(
            ((losses_load, True), (losses_pv, False), (losses_penalty, False)),
            day_index, data.dates,
        )
        parameters = Hyperparameters(
            DECAY_DAY_CANDIDATES[load_index], DECAY_DAY_CANDIDATES[pv_index],
            float(penalties[penalty_index]),
        )
        load_candidates, pv_candidates = engine.candidates(
            day, history_dates, data.actual_load_energy[:day_index],
            data.actual_photovoltaic_energy[:day_index],
            data.cold_start_load_energy, data.cold_start_photovoltaic_energy,
        )
        prior_load, prior_pv = engine.prior_fallback_flags(day, history_dates)
        forecast = Forecast(
            load_candidates[load_index], pv_candidates[pv_index],
            engine.latest_training_date(day, history_dates), prior_load, prior_pv,
        )
        # 年末硬区间替代软惩罚，此时所有惩罚候选的计划相同，只需求解一次。
        if hard_terminal:
            plan = solve_validated_plan(day, data.price, forecast, initial_energy, 0.0, True)
            plans = [plan] * len(penalties)
        else:
            plans = [
                solve_validated_plan(day, data.price, forecast, initial_energy, float(penalty), False)
                for penalty in penalties
            ]
        selected_plan = plans[penalty_index]

        # 选定并冻结 selected_plan 后才读取当天实测；各候选共用相同实际日初状态。
        actual_load = data.actual_load_energy[day_index]
        actual_pv = data.actual_photovoltaic_energy[day_index]
        replays = [replay_actual_day(plan, actual_load, actual_pv, initial_energy) for plan in plans]
        for plan, replay in zip(plans, replays):
            validate_replay(day, plan, replay, actual_load, actual_pv)
        planned_costs = np.array([float(data.price @ plan.grid_purchase) for plan in plans])
        emergency_costs = np.array([
            float(EMERGENCY_PRICE_MULTIPLIER * (data.price @ replay.emergency_purchase))
            for replay in replays
        ])
        realized_costs = planned_costs + emergency_costs
        result = DailyResult(
            day=day, forecast=forecast, plan=selected_plan, replay=replays[penalty_index],
            actual_load_energy=actual_load, actual_photovoltaic_energy=actual_pv,
            planned_purchase_cost=float(planned_costs[penalty_index]),
            emergency_purchase_cost=float(emergency_costs[penalty_index]),
            realized_total_cost=float(realized_costs[penalty_index]),
            hyperparameters=parameters, parameter_latest_date=parameter_latest,
        )
        for loss, history in (
            (np.array([pinball_loss(actual_load, item, LOAD_QUANTILE) for item in load_candidates]), losses_load),
            (np.array([pinball_loss(actual_pv, item, PHOTOVOLTAIC_QUANTILE) for item in pv_candidates]), losses_pv),
            (realized_costs + terminal_value * np.array([
                abs(replay.stored_energy[-1] - E_REFERENCE) for replay in replays
            ]), losses_penalty),
        ):
            if not np.isfinite(loss).all() or np.any(loss < 0):
                raise RuntimeError(f"{day} 的候选验证损失无效。")
            history.append(loss)
        results.append(result)
        initial_energy = float(result.replay.stored_energy[-1])

    return RollingResult(
        hyperparameters=results[-1].hyperparameters,
        tuning=TuningSummary(
            load_decay_scores=tuple(zip(DECAY_DAY_CANDIDATES, load_scores)),
            photovoltaic_decay_scores=tuple(zip(DECAY_DAY_CANDIDATES, pv_scores)),
            terminal_penalty_scores=tuple(zip(map(float, penalties), map(float, penalty_scores))),
        ),
        warmup_days=tuple(results[:JANUARY_DAYS]),
        official_days=tuple(results[JANUARY_DAYS:]),
    )


def group_emergency_intervals(
    emergency_purchase: np.ndarray,
) -> tuple[EmergencyInterval, ...]:
    """把相邻的非零紧急购电 10 分钟时段合并为连续区间。"""
    active = emergency_purchase > EMERGENCY_TOLERANCE
    intervals: list[EmergencyInterval] = []
    period_index = 0
    while period_index < N_PERIODS:
        if not active[period_index]:
            period_index += 1
            continue
        start_index = period_index
        while period_index < N_PERIODS and active[period_index]:
            period_index += 1
        stop_index = period_index
        intervals.append(
            EmergencyInterval(
                start_index=start_index,
                stop_index=stop_index,
                label=(
                    f"{format_clock(start_index * 10)}-"
                    f"{format_clock(stop_index * 10)}"
                ),
                energy=float(np.sum(emergency_purchase[start_index:stop_index])),
            )
        )
    return tuple(intervals)


def display_number(value: float) -> str:
    """按展示精度格式化数值，并清除极小的正负误差。"""
    if abs(value) < DISPLAY_ZERO_TOLERANCE:
        value = 0.0
    return f"{value:.{DISPLAY_DECIMALS}f}"


def excel_number(value: float) -> float:
    """仅在写出 Excel 时取六位小数并消除负零。"""
    return float(display_number(float(value)))


def write_csv(
    path: Path,
    header: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> Path:
    """用 Excel 可直接识别的 UTF-8 BOM 编码写出 CSV。"""
    with path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.writer(csv_file, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    return path


def write_detail_csv(
    data: ModelData,
    results: Sequence[DailyResult],
) -> Path:
    """写出 334 天逐时段预测、计划、实际执行与费用明细。"""
    header = (
        "日期",
        "时间段",
        "训练样本最晚日期",
        "参数验证最晚日期",
        "负载预测来源",
        "光伏预测来源",
        "电价（元/kWh）",
        "预测负载电量（kWh）",
        "预测光伏电量（kWh）",
        "实际负载电量（kWh）",
        "实际光伏电量（kWh）",
        "计划购电量（kWh）",
        "计划充电量（kWh）",
        "计划放电量（kWh）",
        "实际充电量（kWh）",
        "实际放电量（kWh）",
        "紧急购电量（kWh）",
        "实际弃光量（kWh）",
        "未利用计划购电量（kWh）",
        "时段初储电量（kWh）",
        "时段末储电量（kWh）",
        "时段计划购电费（元）",
        "时段紧急购电费（元）",
        "时段总购电费（元）",
        "预测净负荷（kWh）",
        "实际净负荷（kWh）",
        "计划充电截断量（kWh）",
        "计划放电截断量（kWh）",
    )

    def rows() -> Iterable[Sequence[object]]:
        for result in results:
            latest_training_date = (
                result.forecast.latest_training_date.isoformat()
                if result.forecast.latest_training_date is not None
                else "附件1冷启动"
            )
            parameter_latest_date = (
                result.parameter_latest_date.isoformat()
                if result.parameter_latest_date is not None
                else "预设参数"
            )
            load_source = "附件1冷启动先验" if result.forecast.used_prior_load else "历史实测分位数"
            pv_source = "附件1冷启动先验" if result.forecast.used_prior_photovoltaic else "历史实测分位数"
            for period_index, interval_label in enumerate(data.interval_labels):
                planned_cost = (
                    data.price[period_index]
                    * result.plan.grid_purchase[period_index]
                )
                emergency_cost = (
                    EMERGENCY_PRICE_MULTIPLIER
                    * data.price[period_index]
                    * result.replay.emergency_purchase[period_index]
                )
                yield (
                    result.day.isoformat(),
                    interval_label,
                    latest_training_date,
                    parameter_latest_date,
                    load_source,
                    pv_source,
                    display_number(data.price[period_index]),
                    display_number(result.forecast.load_energy[period_index]),
                    display_number(
                        result.forecast.photovoltaic_energy[period_index]
                    ),
                    display_number(result.actual_load_energy[period_index]),
                    display_number(
                        result.actual_photovoltaic_energy[period_index]
                    ),
                    display_number(result.plan.grid_purchase[period_index]),
                    display_number(result.plan.charge[period_index]),
                    display_number(result.plan.discharge[period_index]),
                    display_number(result.replay.charge[period_index]),
                    display_number(result.replay.discharge[period_index]),
                    display_number(
                        result.replay.emergency_purchase[period_index]
                    ),
                    display_number(result.replay.curtailment[period_index]),
                    display_number(
                        result.replay.unused_planned_purchase[period_index]
                    ),
                    display_number(result.replay.stored_energy[period_index]),
                    display_number(
                        result.replay.stored_energy[period_index + 1]
                    ),
                    display_number(planned_cost),
                    display_number(emergency_cost),
                    display_number(planned_cost + emergency_cost),
                    display_number(result.forecast.load_energy[period_index] - result.forecast.photovoltaic_energy[period_index]),
                    display_number(result.actual_load_energy[period_index] - result.actual_photovoltaic_energy[period_index]),
                    display_number(result.plan.charge[period_index] - result.replay.charge[period_index]),
                    display_number(result.plan.discharge[period_index] - result.replay.discharge[period_index]),
                )

    return write_csv(DETAIL_FILE, header, rows())


def result_by_date(
    results: Sequence[DailyResult],
) -> dict[date, DailyResult]:
    """建立日期到日结果的索引，并检查日期是否唯一。"""
    mapping = {result.day: result for result in results}
    if len(mapping) != len(results):
        raise RuntimeError("滚动结果中存在重复日期。")
    return mapping


def write_purchase_summary_csv(
    data: ModelData,
    results: Sequence[DailyResult],
) -> Path:
    """按表 1 口径汇总四个指定日期的计划购电量和全天费用。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    for specified_day in SPECIFIED_DATES:
        if specified_day not in mapping:
            raise RuntimeError(f"缺少指定日期 {specified_day} 的结果。")
        result = mapping[specified_day]
        for hour in SPECIFIED_HOURS:
            period_index = hour * 6
            rows.append(
                (
                    specified_day.isoformat(),
                    "指定时段计划购电量",
                    data.interval_labels[period_index],
                    display_number(result.plan.grid_purchase[period_index]),
                    "kWh",
                )
            )
        rows.extend(
            (
                (
                    specified_day.isoformat(),
                    "全天的计划购电量",
                    "00:00-24:00",
                    display_number(float(np.sum(result.plan.grid_purchase))),
                    "kWh",
                ),
                (
                    specified_day.isoformat(),
                    "全天购电费（含紧急购电）",
                    "00:00-24:00",
                    display_number(result.realized_total_cost),
                    "元",
                ),
            )
        )
    return write_csv(
        PURCHASE_SUMMARY_FILE,
        ("日期", "指标", "时间段", "数值", "单位"),
        rows,
    )


def write_storage_summary_csv(
    results: Sequence[DailyResult],
) -> Path:
    """按表 2 口径汇总四个指定日期的实际充放电量和储电量。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    periods_per_block = 4 * 6
    for specified_day in SPECIFIED_DATES:
        if specified_day not in mapping:
            raise RuntimeError(f"缺少指定日期 {specified_day} 的结果。")
        result = mapping[specified_day]
        for block_index, block_label in enumerate(STORAGE_BLOCK_LABELS):
            start = block_index * periods_per_block
            stop = (block_index + 1) * periods_per_block
            rows.append(
                (
                    specified_day.isoformat(),
                    block_label,
                    display_number(float(np.sum(result.replay.charge[start:stop]))),
                    display_number(
                        float(np.sum(result.replay.discharge[start:stop]))
                    ),
                    "",
                    "",
                )
            )
        rows.extend(
            (
                (
                    specified_day.isoformat(),
                    "",
                    "",
                    "",
                    "00:00",
                    display_number(result.replay.stored_energy[0]),
                ),
                (
                    specified_day.isoformat(),
                    "",
                    "",
                    "",
                    "24:00",
                    display_number(result.replay.stored_energy[-1]),
                ),
            )
        )
    return write_csv(
        STORAGE_SUMMARY_FILE,
        (
            "日期",
            "时间段",
            "充电量（kWh）",
            "放电量（kWh）",
            "时刻",
            "储电量（kWh）",
        ),
        rows,
    )


def write_emergency_summary_csv(
    results: Sequence[DailyResult],
) -> Path:
    """按表 3 口径汇总四个指定日期的连续紧急购电区间。"""
    mapping = result_by_date(results)
    rows: list[Sequence[object]] = []
    for specified_day in SPECIFIED_DATES:
        if specified_day not in mapping:
            raise RuntimeError(f"缺少指定日期 {specified_day} 的结果。")
        intervals = group_emergency_intervals(
            mapping[specified_day].replay.emergency_purchase
        )
        if not intervals:
            rows.append((specified_day.isoformat(), "无", display_number(0.0)))
            continue
        for interval in intervals:
            rows.append(
                (
                    specified_day.isoformat(),
                    interval.label,
                    display_number(interval.energy),
                )
            )
    return write_csv(
        EMERGENCY_SUMMARY_FILE,
        ("日期", "紧急购电时间段", "紧急购电量（kWh）"),
        rows,
    )


def concatenate_result_array(
    results: Sequence[DailyResult],
    getter: str,
    nested_getter: str,
) -> np.ndarray:
    """拼接 DailyResult 中某个计划或回放数组。"""
    return np.concatenate(
        [getattr(getattr(result, getter), nested_getter) for result in results]
    )


def build_check_rows(
    rolling_result: RollingResult,
) -> list[Sequence[object]]:
    """汇总预测误差、约束残差、状态边界和费用统计。"""
    results = rolling_result.official_days
    forecast_load = concatenate_result_array(results, "forecast", "load_energy")
    forecast_photovoltaic = concatenate_result_array(
        results, "forecast", "photovoltaic_energy"
    )
    actual_load = np.concatenate(
        [result.actual_load_energy for result in results]
    )
    actual_photovoltaic = np.concatenate(
        [result.actual_photovoltaic_energy for result in results]
    )
    planned_purchase = concatenate_result_array(
        results, "plan", "grid_purchase"
    )
    planned_charge = concatenate_result_array(results, "plan", "charge")
    planned_discharge = concatenate_result_array(results, "plan", "discharge")
    actual_charge = concatenate_result_array(results, "replay", "charge")
    actual_discharge = concatenate_result_array(results, "replay", "discharge")
    emergency_purchase = concatenate_result_array(
        results, "replay", "emergency_purchase"
    )
    actual_curtailment = concatenate_result_array(
        results, "replay", "curtailment"
    )
    unused_planned_purchase = concatenate_result_array(
        results, "replay", "unused_planned_purchase"
    )
    all_stored_energy = np.concatenate(
        [result.replay.stored_energy for result in results]
    )

    plan_balance_residuals: list[np.ndarray] = []
    plan_state_residuals: list[np.ndarray] = []
    actual_balance_residuals: list[np.ndarray] = []
    actual_state_residuals: list[np.ndarray] = []
    for result in results:
        plan_balance_residuals.append(
            result.plan.grid_purchase
            + result.forecast.photovoltaic_energy
            + result.plan.discharge
            - result.plan.curtailment
            - result.forecast.load_energy
            - result.plan.charge
        )
        plan_state_residuals.append(
            result.plan.stored_energy[1:]
            - result.plan.stored_energy[:-1]
            - ETA_CHARGE * result.plan.charge
            + result.plan.discharge / ETA_DISCHARGE
        )
        actual_balance_residuals.append(
            result.plan.grid_purchase
            + result.actual_photovoltaic_energy
            + result.replay.discharge
            + result.replay.emergency_purchase
            - result.replay.curtailment
            - result.replay.unused_planned_purchase
            - result.actual_load_energy
            - result.replay.charge
        )
        actual_state_residuals.append(
            result.replay.stored_energy[1:]
            - result.replay.stored_energy[:-1]
            - ETA_CHARGE * result.replay.charge
            + result.replay.discharge / ETA_DISCHARGE
        )

    continuity_errors = [
        abs(
            results[result_index].replay.stored_energy[0]
            - results[result_index - 1].replay.stored_energy[-1]
        )
        for result_index in range(1, len(results))
    ]
    emergency_interval_count = sum(
        len(group_emergency_intervals(result.replay.emergency_purchase))
        for result in results
    )
    planned_cost = sum(result.planned_purchase_cost for result in results)
    emergency_cost = sum(result.emergency_purchase_cost for result in results)
    total_cost = planned_cost + emergency_cost

    rows: list[Sequence[object]] = [
        (
            "超参数",
            "最终日负载指数衰减天数",
            display_number(rolling_result.hyperparameters.load_decay_days),
            "天",
            "通过",
        ),
        (
            "超参数",
            "最终日光伏指数衰减天数",
            display_number(
                rolling_result.hyperparameters.photovoltaic_decay_days
            ),
            "天",
            "通过",
        ),
        (
            "超参数",
            "最终日选择的日末储电量惩罚系数（硬终端日不计罚）",
            display_number(rolling_result.hyperparameters.terminal_penalty),
            "元/kWh",
            "通过",
        ),
        (
            "预测",
            "负载 MAE",
            display_number(float(np.mean(np.abs(actual_load - forecast_load)))),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "负载 RMSE",
            display_number(
                float(np.sqrt(np.mean((actual_load - forecast_load) ** 2)))
            ),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "负载 0.8 分位数损失",
            display_number(
                pinball_loss(actual_load, forecast_load, LOAD_QUANTILE)
            ),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "光伏 MAE",
            display_number(
                float(np.mean(np.abs(actual_photovoltaic - forecast_photovoltaic)))
            ),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "光伏 RMSE",
            display_number(
                float(
                    np.sqrt(
                        np.mean(
                            (actual_photovoltaic - forecast_photovoltaic) ** 2
                        )
                    )
                )
            ),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "光伏 0.2 分位数损失",
            display_number(
                pinball_loss(
                    actual_photovoltaic,
                    forecast_photovoltaic,
                    PHOTOVOLTAIC_QUANTILE,
                )
            ),
            "kWh/时段",
            "参考",
        ),
        (
            "预测",
            "负载使用附件1冷启动先验的正式天数",
            str(sum(day.forecast.used_prior_load for day in results)),
            "天",
            "参考",
        ),
        (
            "预测",
            "光伏使用附件1冷启动先验的正式天数",
            str(sum(day.forecast.used_prior_photovoltaic for day in results)),
            "天",
            "参考",
        ),
        (
            "约束",
            "最大计划电量平衡残差",
            display_number(
                float(np.max(np.abs(np.concatenate(plan_balance_residuals))))
            ),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "最大计划状态转移残差",
            display_number(
                float(np.max(np.abs(np.concatenate(plan_state_residuals))))
            ),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "最大实际电量平衡残差",
            display_number(
                float(np.max(np.abs(np.concatenate(actual_balance_residuals))))
            ),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "最大实际状态转移残差",
            display_number(
                float(np.max(np.abs(np.concatenate(actual_state_residuals))))
            ),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "最大跨日储电量连续误差",
            display_number(float(max(continuity_errors, default=0.0))),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "实际最小储电量",
            display_number(float(np.min(all_stored_energy))),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "实际最大储电量",
            display_number(float(np.max(all_stored_energy))),
            "kWh",
            "通过",
        ),
        (
            "约束",
            "12 月 31 日 24:00 实际终端储电量区间越界量",
            display_number(terminal_interval_error(results[-1].replay.stored_energy[-1])),
            "kWh",
            "通过" if terminal_interval_error(results[-1].replay.stored_energy[-1]) <= TERMINAL_TOLERANCE else "失败：不得提交",
        ),
        (
            "数量",
            "计划购电总量",
            display_number(float(np.sum(planned_purchase))),
            "kWh",
            "参考",
        ),
        (
            "数量",
            "紧急购电总量",
            display_number(float(np.sum(emergency_purchase))),
            "kWh",
            "参考",
        ),
        (
            "数量",
            "紧急购电 10 分钟时段数",
            str(int(np.count_nonzero(emergency_purchase > EMERGENCY_TOLERANCE))),
            "个",
            "参考",
        ),
        (
            "数量",
            "连续紧急购电区间数",
            str(emergency_interval_count),
            "个",
            "参考",
        ),
        (
            "数量",
            "弃光总量",
            display_number(float(np.sum(actual_curtailment))),
            "kWh",
            "参考",
        ),
        (
            "数量",
            "未利用计划购电总量",
            display_number(float(np.sum(unused_planned_purchase))),
            "kWh",
            "参考",
        ),
        (
            "费用",
            "计划购电费",
            display_number(planned_cost),
            "元",
            "参考",
        ),
        (
            "费用",
            "紧急购电费",
            display_number(emergency_cost),
            "元",
            "参考",
        ),
        (
            "费用",
            "334天总费用",
            display_number(total_cost),
            "元",
            "参考",
        ),
        (
            "费用",
            "紧急购电费占比",
            display_number(emergency_cost / total_cost if total_cost else 0.0),
            "比例",
            "参考",
        ),
        (
            "边界",
            "最大实际充电量超过计划量",
            display_number(float(np.max(actual_charge - planned_charge))),
            "kWh",
            "通过",
        ),
        (
            "边界",
            "最大实际放电量超过计划量",
            display_number(float(np.max(actual_discharge - planned_discharge))),
            "kWh",
            "通过",
        ),
    ]

    actual_net = actual_load - actual_photovoltaic
    forecast_net = forecast_load - forecast_photovoltaic
    rows.extend((
        ("超参数", "年末储电量下限", display_number(E_TERMINAL_MIN), "kWh", "参考"),
        ("超参数", "年末储电量上限", display_number(E_TERMINAL_MAX), "kWh", "参考"),
        ("费用", "1月初始化费用", display_number(sum(day.realized_total_cost for day in rolling_result.warmup_days)), "元", "参考"),
        ("费用", "365天总费用", display_number(sum(day.realized_total_cost for day in rolling_result.warmup_days + results)), "元", "参考"),
        ("求解", "最大MIP相对间隙", f"{max(day.plan.mip_gap for day in rolling_result.warmup_days + results):.12g}", "比例", "通过"),
        ("边界", "年末计划储电量", display_number(results[-1].plan.stored_energy[-1]), "kWh", "参考"),
        ("边界", "年末实际储电量", display_number(results[-1].replay.stored_energy[-1]), "kWh", "参考"),
        ("预测", "净负荷 MAE", display_number(float(np.mean(np.abs(actual_net - forecast_net)))), "kWh/时段", "参考"),
        ("预测", "净负荷 0.8 分位数损失", display_number(pinball_loss(actual_net, forecast_net, NET_LOAD_QUANTILE)), "kWh/时段", "参考"),
        ("预测", "净负荷实际值不超过预测值的时段比例",
         display_number(float(np.mean(actual_net <= forecast_net))), "比例", "参考"),
        ("执行", "计划充电截断总量", display_number(float(np.sum(planned_charge - actual_charge))), "kWh", "参考"),
        ("执行", "计划放电截断总量", display_number(float(np.sum(planned_discharge - actual_discharge))), "kWh", "参考"),
        ("边界", "年末实际储电量减计划储电量",
         display_number(results[-1].replay.stored_energy[-1] - results[-1].plan.stored_energy[-1]), "kWh", "参考"),
    ))
    for decay_days, score in rolling_result.tuning.load_decay_scores:
        rows.append(
            (
                "超参数验证",
                f"12月31日前滚动验证：负载衰减 {decay_days:g} 天的 pinball 损失",
                display_number(score),
                "kWh/时段",
                "参考",
            )
        )
    for decay_days, score in rolling_result.tuning.photovoltaic_decay_scores:
        rows.append(
            (
                "超参数验证",
                f"12月31日前滚动验证：光伏衰减 {decay_days:g} 天的 pinball 损失",
                display_number(score),
                "kWh/时段",
                "参考",
            )
        )
    for penalty, score in rolling_result.tuning.terminal_penalty_scores:
        rows.append(
            (
                "超参数验证",
                f"12月31日前滚动验证：日末惩罚 {penalty:.6f} 元/kWh 的单步评分",
                display_number(score),
                "元",
                "参考",
            )
        )
    return rows


def write_check_csv(rolling_result: RollingResult) -> Path:
    """写出模型超参数、预测指标、约束残差和全年统计。"""
    return write_csv(
        CHECK_FILE,
        ("类别", "校验项", "数值", "单位", "结论"),
        build_check_rows(rolling_result),
    )


def capture_row_styles(
    worksheet: object,
    row_numbers: Sequence[int],
    maximum_column: int,
) -> tuple[tuple[object, ...], tuple[float | None, ...]]:
    """在扩展模板工作表前保存示例行的单元格样式和行高。"""
    styles = tuple(
        tuple(
            copy(worksheet.cell(row=row_number, column=column_index)._style)
            for column_index in range(1, maximum_column + 1)
        )
        for row_number in row_numbers
    )
    heights = tuple(
        worksheet.row_dimensions[row_number].height for row_number in row_numbers
    )
    return styles, heights


def apply_row_style(
    worksheet: object,
    row_number: int,
    styles: Sequence[object],
    row_height: float | None,
) -> None:
    """把模板示例行样式复制到新结果行。"""
    for column_index, style in enumerate(styles, start=1):
        worksheet.cell(row=row_number, column=column_index)._style = copy(style)
    worksheet.row_dimensions[row_number].height = row_height


def validate_result_template(
    workbook: object,
    official_results: Sequence[DailyResult],
) -> None:
    """在写入前核对 result2.xlsx 的工作表、日期和关键表头。"""
    missing_sheets = set(RESULT_SHEET_NAMES).difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(
            "result2.xlsx 模板缺少工作表：" + "、".join(sorted(missing_sheets))
        )

    purchase_sheet = workbook["计划购电量"]
    if purchase_sheet.cell(row=1, column=1).value != "日期\\时间":
        raise ValueError("计划购电量工作表 A1 应为“日期\\时间”。")
    interval_headers = [
        purchase_sheet.cell(row=1, column=column_index).value
        for column_index in range(2, N_PERIODS + 2)
    ]
    if len(interval_headers) != N_PERIODS or any(
        not isinstance(header, str) or not header.strip()
        for header in interval_headers
    ):
        raise ValueError("计划购电量工作表必须保留 144 个非空时间段表头。")
    if len(set(interval_headers)) != N_PERIODS:
        raise ValueError("计划购电量工作表的 144 个时间段表头必须互不重复。")
    # 原模板表头整体从 0:10 开始，比内部区间起点晚 10 分钟。
    # 按 Q2 的约定保留原标签并按序号映射，不能按标签再移动一次数据。
    for index, header in enumerate(interval_headers):
        start_label, stop_label = header.split("-")
        def minutes(label: str) -> int:
            clock_label, _, day_offset = label.partition("+")
            hour, minute = map(int, clock_label.split(":"))
            if not (0 <= hour <= 24 and 0 <= minute < 60):
                raise ValueError(f"模板时间标签无效：{header}")
            return hour * 60 + minute + 1440 * int(day_offset or 0)
        start, stop = minutes(start_label), minutes(stop_label)
        if index == N_PERIODS - 1 and start == 0:
            start = 1440  # 模板末列原文为 0:00-0:10+1。
        if (start, stop) != ((index + 1) * 10, (index + 2) * 10):
            raise ValueError(f"计划购电模板第 {index + 2} 列时间错位：{header}")
    if purchase_sheet.cell(row=1, column=N_PERIODS + 2).value != "全天购电量":
        raise ValueError("计划购电量工作表倒数第二列应为“全天购电量”。")
    if purchase_sheet.cell(row=1, column=N_PERIODS + 3).value != "全天购电费":
        raise ValueError("计划购电量工作表最后一列应为“全天购电费”。")

    expected_dates = tuple(result.day for result in official_results)
    template_dates = tuple(
        normalize_excel_date(
            purchase_sheet.cell(row=row_index, column=1).value,
            f"计划购电量工作表 A{row_index}",
        )
        for row_index in range(2, OFFICIAL_DAYS + 2)
    )
    if template_dates != expected_dates:
        raise ValueError("计划购电量工作表日期必须从 2025-02-01 连续到 2025-12-31。")

    storage_sheet = workbook["充放电量"]
    storage_headers = tuple(
        storage_sheet.cell(row=1, column=column_index).value
        for column_index in range(1, 7)
    )
    if storage_headers != (
        "日期",
        "时间段",
        "充电量",
        "放电量",
        "时刻",
        "储电量",
    ):
        raise ValueError("充放电量工作表表头与模板约定不一致。")
    template_block_labels = tuple(
        storage_sheet.cell(row=row_index, column=2).value
        for row_index in range(2, 8)
    )
    if template_block_labels != STORAGE_BLOCK_LABELS:
        raise ValueError("充放电量工作表的六个 4 小时时间段与题目不一致。")

    emergency_sheet = workbook["紧急购电量"]
    emergency_headers = tuple(
        emergency_sheet.cell(row=1, column=column_index).value
        for column_index in range(1, 4)
    )
    if emergency_headers != ("日期", "购电时间段", "购电量"):
        raise ValueError("紧急购电量工作表表头与模板约定不一致。")


def fill_purchase_sheet(
    worksheet: object,
    results: Sequence[DailyResult],
) -> None:
    """填写每天 144 个计划购电量、全天的计划购电量和实际总费用。"""
    for row_index, result in enumerate(results, start=2):
        for period_index, value in enumerate(result.plan.grid_purchase, start=2):
            cell = worksheet.cell(row=row_index, column=period_index)
            cell.value = excel_number(value)
            cell.number_format = "0.000000"
        total_purchase_cell = worksheet.cell(
            row=row_index,
            column=N_PERIODS + 2,
        )
        total_cost_cell = worksheet.cell(
            row=row_index,
            column=N_PERIODS + 3,
        )
        total_purchase_cell.value = excel_number(np.sum(result.plan.grid_purchase))
        total_cost_cell.value = excel_number(result.realized_total_cost)
        total_purchase_cell.number_format = "0.000000"
        total_cost_cell.number_format = "0.000000"


def fill_storage_sheet(
    worksheet: object,
    results: Sequence[DailyResult],
) -> None:
    """扩展模板并填写每个日期六个时段的实际充放电量和边界储电量。"""
    prototype_styles, prototype_heights = capture_row_styles(
        worksheet,
        row_numbers=tuple(range(2, 8)),
        maximum_column=6,
    )
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    periods_per_block = 4 * 6
    destination_row = 2
    for result in results:
        for block_index, block_label in enumerate(STORAGE_BLOCK_LABELS):
            apply_row_style(
                worksheet,
                destination_row,
                prototype_styles[block_index],
                prototype_heights[block_index],
            )
            if block_index == 0:
                worksheet.cell(row=destination_row, column=1).value = result.day
                worksheet.cell(row=destination_row, column=1).number_format = (
                    "yyyy/m/d"
                )
            worksheet.cell(row=destination_row, column=2).value = block_label
            start = block_index * periods_per_block
            stop = (block_index + 1) * periods_per_block
            charge_cell = worksheet.cell(row=destination_row, column=3)
            discharge_cell = worksheet.cell(row=destination_row, column=4)
            charge_cell.value = excel_number(np.sum(result.replay.charge[start:stop]))
            discharge_cell.value = excel_number(
                np.sum(result.replay.discharge[start:stop])
            )
            charge_cell.number_format = "0.000000"
            discharge_cell.number_format = "0.000000"

            if block_index == 0:
                worksheet.cell(row=destination_row, column=5).value = time(0, 0)
                worksheet.cell(row=destination_row, column=5).number_format = "h:mm"
                storage_cell = worksheet.cell(row=destination_row, column=6)
                storage_cell.value = excel_number(result.replay.stored_energy[0])
                storage_cell.number_format = "0.000000"
            elif block_index == 1:
                worksheet.cell(row=destination_row, column=5).value = "24:00"
                storage_cell = worksheet.cell(row=destination_row, column=6)
                storage_cell.value = excel_number(result.replay.stored_energy[-1])
                storage_cell.number_format = "0.000000"
            destination_row += 1


def fill_emergency_sheet(
    worksheet: object,
    results: Sequence[DailyResult],
) -> None:
    """扩展模板并写入所有非零连续紧急购电区间。"""
    prototype_styles, prototype_heights = capture_row_styles(
        worksheet,
        row_numbers=(2, 3, 4),
        maximum_column=3,
    )
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    destination_row = 2
    for result in results:
        intervals = group_emergency_intervals(result.replay.emergency_purchase)
        for interval_index, interval in enumerate(intervals):
            if interval_index == 0:
                prototype_index = 0
            elif interval_index == len(intervals) - 1:
                prototype_index = 2
            else:
                prototype_index = 1
            apply_row_style(
                worksheet,
                destination_row,
                prototype_styles[prototype_index],
                prototype_heights[prototype_index],
            )
            if interval_index == 0:
                worksheet.cell(row=destination_row, column=1).value = result.day
                worksheet.cell(row=destination_row, column=1).number_format = (
                    "yyyy/m/d"
                )
            worksheet.cell(row=destination_row, column=2).value = interval.label
            purchase_cell = worksheet.cell(row=destination_row, column=3)
            purchase_cell.value = excel_number(interval.energy)
            purchase_cell.number_format = "0.000000"
            destination_row += 1


def write_result_workbook(rolling_result: RollingResult) -> Path:
    """从原始模板生成临时工作簿，回读校验后原子替换正式文件。"""
    if len(rolling_result.official_days) != OFFICIAL_DAYS:
        raise RuntimeError("正式日期不足334天，禁止发布结果工作簿。")
    require_terminal_interval(rolling_result.official_days[-1])
    if not RESULT_TEMPLATE_FILE.exists():
        raise FileNotFoundError(f"未找到问题二结果模板：{RESULT_TEMPLATE_FILE}")
    OUTPUT_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(RESULT_TEMPLATE_FILE)
    temporary = None
    try:
        validate_result_template(workbook, rolling_result.official_days)
        fill_purchase_sheet(workbook["计划购电量"], rolling_result.official_days)
        fill_storage_sheet(workbook["充放电量"], rolling_result.official_days)
        fill_emergency_sheet(workbook["紧急购电量"], rolling_result.official_days)
        with tempfile.NamedTemporaryFile(dir=OUTPUT_RESULT_DIR, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        written = load_workbook(temporary, read_only=True, data_only=True)
        try:
            purchase = written["计划购电量"]
            if purchase.max_row != OFFICIAL_DAYS + 1:
                raise RuntimeError("写出后的计划购电表行数不正确。")
            for values, result in zip(purchase.iter_rows(min_row=2, values_only=True), rolling_result.official_days):
                expected = [excel_number(x) for x in result.plan.grid_purchase]
                expected += [excel_number(np.sum(result.plan.grid_purchase)), excel_number(result.realized_total_cost)]
                if normalize_excel_date(values[0], "输出日期") != result.day or not np.allclose(
                    np.asarray(values[1:], dtype=float), expected, atol=1e-9, rtol=0.0,
                ):
                    raise RuntimeError(f"{result.day} 的 Excel 购电数据回读校验失败。")
            if written["充放电量"].max_row != 1 + OFFICIAL_DAYS * 6:
                raise RuntimeError("写出后的储能表行数不正确。")
            expected_intervals = sum(len(group_emergency_intervals(x.replay.emergency_purchase)) for x in rolling_result.official_days)
            if written["紧急购电量"].max_row != 1 + expected_intervals:
                raise RuntimeError("写出后的紧急购电表行数不正确。")
        finally:
            written.close()
        os.replace(temporary, RESULT_FILE)
    finally:
        workbook.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return RESULT_FILE


def validate_rolling_result(data: ModelData, rolling: RollingResult) -> float:
    """检查全年因果性、费用、连续性与能量守恒，返回实际终端区间越界量。"""
    all_days = rolling.warmup_days + rolling.official_days
    if tuple(result.day for result in all_days) != data.dates:
        raise RuntimeError("全年结果缺少日期或顺序错误。")
    if len(rolling.warmup_days) != JANUARY_DAYS or len(rolling.official_days) != OFFICIAL_DAYS:
        raise RuntimeError("初始化天数或正式结果天数不正确。")
    previous_energy = E_INITIAL
    for result in all_days:
        for cutoff in (result.forecast.latest_training_date, result.parameter_latest_date):
            if cutoff is not None and cutoff >= result.day:
                raise RuntimeError(f"{result.day} 出现前视预测或前视调参。")
        if result.forecast.latest_training_date is None and not (
            result.forecast.used_prior_load and result.forecast.used_prior_photovoltaic
        ):
            raise RuntimeError(f"{result.day} 没有训练样本时的先验回退标记不一致。")
        validate_plan(result.day, result.forecast, result.plan, previous_energy, result.day == YEAR_END)
        validate_replay(result.day, result.plan, result.replay, result.actual_load_energy, result.actual_photovoltaic_energy)
        if abs(previous_energy - result.replay.stored_energy[0]) > BOUND_TOLERANCE:
            raise RuntimeError(f"{result.day} 实际储电量跨日不连续。")
        if abs(terminal_diagnostics(result)["identity_residual_kwh"]) > N_PERIODS * BALANCE_TOLERANCE:
            raise RuntimeError(f"{result.day} 计划与实际储电偏差不能由充放电截断解释。")
        cost = float(data.price @ result.plan.grid_purchase)
        emergency = float(EMERGENCY_PRICE_MULTIPLIER * (data.price @ result.replay.emergency_purchase))
        if not np.isfinite([result.planned_purchase_cost, result.emergency_purchase_cost, result.realized_total_cost]).all():
            raise RuntimeError(f"{result.day} 费用非有限。")
        if max(abs(cost - result.planned_purchase_cost), abs(emergency - result.emergency_purchase_cost), abs(cost + emergency - result.realized_total_cost)) > BALANCE_TOLERANCE:
            raise RuntimeError(f"{result.day} 原始精度费用重算失败。")
        previous_energy = float(result.replay.stored_energy[-1])
    total_residual = sum(float(np.sum(
        day.plan.grid_purchase + day.actual_photovoltaic_energy + day.replay.emergency_purchase
        - day.replay.curtailment - day.replay.unused_planned_purchase - day.actual_load_energy
        - (1 - ETA_CHARGE) * day.replay.charge
        - (1 / ETA_DISCHARGE - 1) * day.replay.discharge
    )) for day in all_days) - (previous_energy - E_INITIAL)
    if abs(total_residual) > len(all_days) * BALANCE_TOLERANCE:
        raise RuntimeError(f"全年能量守恒残差超限：{total_residual} kWh。")
    return terminal_interval_error(previous_energy)


def terminal_diagnostics(result: DailyResult) -> dict[str, float]:
    """从已冻结的计划与实际执行分解终端偏差，不修改任何调度或状态。

    实际末态-计划末态 = 初态差 - eta_c*充电截断 + 放电截断/eta_d。
    差值保留原始精度；正向放电截断项解释实际末态偏高的来源。
    """
    charge_shortfall = float(np.sum(result.plan.charge - result.replay.charge))
    discharge_shortfall = float(np.sum(result.plan.discharge - result.replay.discharge))
    initial_gap = float(result.replay.stored_energy[0] - result.plan.stored_energy[0])
    gap = float(result.replay.stored_energy[-1] - result.plan.stored_energy[-1])
    charge_effect = -ETA_CHARGE * charge_shortfall
    discharge_effect = discharge_shortfall / ETA_DISCHARGE
    return {
        "reference_kwh": E_REFERENCE,
        "lower_bound_kwh": E_TERMINAL_MIN,
        "upper_bound_kwh": E_TERMINAL_MAX,
        "planned_end_kwh": float(result.plan.stored_energy[-1]),
        "actual_end_kwh": float(result.replay.stored_energy[-1]),
        "planned_error_kwh": terminal_interval_error(result.plan.stored_energy[-1]),
        "actual_error_kwh": terminal_interval_error(result.replay.stored_energy[-1]),
        "actual_minus_planned_kwh": gap,
        "charge_shortfall_kwh": charge_shortfall,
        "discharge_shortfall_kwh": discharge_shortfall,
        "charge_shortfall_soc_effect_kwh": charge_effect,
        "discharge_shortfall_soc_effect_kwh": discharge_effect,
        "identity_residual_kwh": gap - initial_gap - charge_effect - discharge_effect,
        "tolerance_kwh": TERMINAL_TOLERANCE,
    }


def require_terminal_interval(result: DailyResult) -> None:
    """正式发布前同时检查计划和实际年末区间，自检调用同一函数。"""
    if result.day != YEAR_END:
        raise RuntimeError("终端门禁必须检查12月31日，不能用其他日期代替。")
    diagnostics = terminal_diagnostics(result)
    if not all(isfinite(value) for value in diagnostics.values()):
        raise RuntimeError("年末终端诊断包含非有限值，禁止发布。")
    if max(diagnostics["planned_error_kwh"], diagnostics["actual_error_kwh"]) > TERMINAL_TOLERANCE:
        raise RuntimeError(
            f"终端区间校验失败：计划末态={diagnostics['planned_end_kwh']:.6f}，"
            f"实际末态={diagnostics['actual_end_kwh']:.6f}，"
            f"要求均位于[{E_TERMINAL_MIN:g}, {E_TERMINAL_MAX:g}] kWh。"
            f"充电截断影响={diagnostics['charge_shortfall_soc_effect_kwh']:+.6f} kWh，"
            f"放电截断影响={diagnostics['discharge_shortfall_soc_effect_kwh']:+.6f} kWh。"
            f"诊断保留于 {OUTPUT_TABLE_DIR}，未发布正式 result2.xlsx。"
            "计划末态满足区间不能保证实际末态满足区间；禁止用当日标签反向修改计划。"
        )


def write_terminal_audit(result: DailyResult) -> Path:
    """保存12月31日全部144时段的计划、截断、末态偏差及恒等式残差。"""
    header = ("日期", "时间段", "计划充电量", "实际充电量", "计划放电量", "实际放电量",
              "充电截断量", "放电截断量", "计划时段末储电量", "实际时段末储电量",
              "实际减计划储电量", "累计截断解释的储电偏差", "偏差恒等式残差")
    rows = []
    explained = float(result.replay.stored_energy[0] - result.plan.stored_energy[0])
    for t in range(N_PERIODS):
        charge_shortfall = float(result.plan.charge[t] - result.replay.charge[t])
        discharge_shortfall = float(result.plan.discharge[t] - result.replay.discharge[t])
        explained += -ETA_CHARGE * charge_shortfall + discharge_shortfall / ETA_DISCHARGE
        actual_gap = float(result.replay.stored_energy[t + 1] - result.plan.stored_energy[t + 1])
        rows.append((result.day.isoformat(), f"{format_clock(t * 10)}-{format_clock((t + 1) * 10)}",
                     *[float(value) for value in (result.plan.charge[t], result.replay.charge[t],
                                                 result.plan.discharge[t], result.replay.discharge[t])],
                     charge_shortfall, discharge_shortfall,
                     float(result.plan.stored_energy[t + 1]), float(result.replay.stored_energy[t + 1]),
                     actual_gap, explained, actual_gap - explained))
    return write_csv(OUTPUT_TABLE_DIR / "2_年末终端偏差分解.csv", header, rows)


def write_daily_audit(data: ModelData, rolling: RollingResult) -> Path:
    """即使年末终端区间校验失败，也保留每天的参数、求解器精度和费用。"""
    header = ("日期", "阶段", "预测历史截止", "调参历史截止", "负载先验回退", "光伏先验回退",
              "负载衰减天数", "光伏衰减天数", "终端惩罚", "求解路径", "求解状态", "MIP相对间隙",
              "日初储电量", "计划日末储电量", "实际日末储电量",
              "计划购电费", "紧急购电费", "实际总购电费",
              "净负荷MAE", "净负荷0.8分位数损失",
              "净负荷覆盖率", "紧急购电费占比", "紧急购电量", "未利用计划购电量",
              "未利用计划购电对应费用（已含在计划费用中）", "实际弃光量",
              "计划充电截断量", "计划放电截断量", "实际减计划日末储电量", "截断偏差恒等式残差")
    rows = []
    for day in rolling.warmup_days + rolling.official_days:
        actual_net = day.actual_load_energy - day.actual_photovoltaic_energy
        forecast_net = day.forecast.load_energy - day.forecast.photovoltaic_energy
        diagnostic = terminal_diagnostics(day)
        rows.append((
            day.day.isoformat(), "初始化" if day.day < OFFICIAL_START else "正式期",
            day.forecast.latest_training_date or "附件1", day.parameter_latest_date or "预设参数",
            "是" if day.forecast.used_prior_load else "否",
            "是" if day.forecast.used_prior_photovoltaic else "否",
            day.hyperparameters.load_decay_days, day.hyperparameters.photovoltaic_decay_days,
            day.hyperparameters.terminal_penalty, day.plan.solver_kind, day.plan.solver_message,
            f"{day.plan.mip_gap:.12g}",
            display_number(day.replay.stored_energy[0]), display_number(day.plan.stored_energy[-1]),
            display_number(day.replay.stored_energy[-1]), display_number(day.planned_purchase_cost),
            display_number(day.emergency_purchase_cost), display_number(day.realized_total_cost),
            display_number(float(np.mean(np.abs(actual_net - forecast_net)))),
            display_number(pinball_loss(actual_net, forecast_net, NET_LOAD_QUANTILE)),
            display_number(float(np.mean(actual_net <= forecast_net))),
            display_number(day.emergency_purchase_cost / day.realized_total_cost if day.realized_total_cost else 0.0),
            display_number(float(np.sum(day.replay.emergency_purchase))),
            display_number(float(np.sum(day.replay.unused_planned_purchase))),
            display_number(float(data.price @ day.replay.unused_planned_purchase)),
            display_number(float(np.sum(day.replay.curtailment))),
            display_number(diagnostic["charge_shortfall_kwh"]), display_number(diagnostic["discharge_shortfall_kwh"]),
            display_number(diagnostic["actual_minus_planned_kwh"]), diagnostic["identity_residual_kwh"],
        ))
    return write_csv(OUTPUT_TABLE_DIR / "2_每日运行审计.csv", header, rows)


def main() -> None:
    """完成全年求解和校验，仅写出 CSV 与正式结果工作簿。"""
    if Path(sys.prefix).name != "2026C" or not (Path(sys.prefix) / "conda-meta").is_dir():
        raise RuntimeError("项目要求使用 conda 环境 2026C。")
    validate_terminal_bounds()
    data = load_model_data()
    template = load_workbook(RESULT_TEMPLATE_FILE)
    try:
        validate_result_template(template, [SimpleNamespace(day=day) for day in data.dates[JANUARY_DAYS:]])
    finally:
        template.close()
    rolling_result = solve_rolling_model(data, ForecastEngine())
    validate_rolling_result(data, rolling_result)
    OUTPUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    output_files = [
        write_daily_audit(data, rolling_result),
        write_check_csv(rolling_result),
        write_detail_csv(data, rolling_result.official_days),
        write_terminal_audit(rolling_result.official_days[-1]),
    ]
    require_terminal_interval(rolling_result.official_days[-1])
    output_files.extend((
        write_purchase_summary_csv(data, rolling_result.official_days),
        write_storage_summary_csv(rolling_result.official_days),
        write_emergency_summary_csv(rolling_result.official_days),
        write_result_workbook(rolling_result),
    ))
    for path in output_files:
        print(f"已保存：{path.relative_to(PROJECT_DIR)}")


if __name__ == "__main__":
    main()
